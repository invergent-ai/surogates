"""A small formula-driven model: an Inputs sheet, a Monthly sheet built in a
loop, and a Summary of model years.  After `recalc.py` it has zero errors.
Copy it, replace the INPUTS and the Monthly columns, keep the helpers.

    usage: python example_model.py [out.xlsx]
"""
import datetime
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")  # hardcoded inputs
BLACK = Font(name=FONT)  # formulas
BOLD = Font(name=FONT, bold=True)
NOTE = Font(name=FONT, italic=True, color="808080")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
USD = "$#,##0;($#,##0);-"
PCT = "0.0%"
INT = "#,##0;(#,##0);-"
MONTH = "mmm yyyy"

wb = Workbook()
ws_in = wb.active
ws_in.title = "Inputs"
ws_in["A1"] = "Inputs: edit the yellow cells"
ws_in["A1"].font = BOLD
ws_in.column_dimensions["A"].width = 34
ws_in.column_dimensions["B"].width = 14
_next_input_row = [3]


def inp(label, value, fmt, note=""):
    """Write a labelled input on Inputs; return its reference, e.g. 'Inputs'!$B$3.

    Put the returned string in formulas. Never type an input's address yourself.
    """
    r = _next_input_row[0]
    _next_input_row[0] += 1
    ws_in.cell(r, 1, label).font = BLACK
    cell = ws_in.cell(r, 2, value)
    cell.font, cell.fill, cell.number_format = BLUE, INPUT_FILL, fmt
    if note:
        ws_in.cell(r, 3, note).font = NOTE
    return f"'{ws_in.title}'!$B${r}"


def put(ws, ref, formula, fmt):
    """Write a formula. It must start with '=' or it is saved as text."""
    assert isinstance(formula, str) and formula.startswith("="), f"{ref}: not a formula: {formula!r}"
    ws[ref] = formula
    ws[ref].font, ws[ref].number_format = BLACK, fmt


def header(ws, labels):
    """Bold shaded header in row 1; returns {label: column letter}."""
    cols = {}
    for i, label in enumerate(labels, 1):
        cell = ws.cell(1, i, label)
        cell.font, cell.fill = BOLD, HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(label) + 2)
        cols[label] = get_column_letter(i)
    ws.freeze_panes = "B2"
    return cols


# ── INPUTS ─────────────────────────────────────────────────────────────
MONTHS = 36  # model length: structure, not an input
START = inp("First month", datetime.date(2026, 1, 1), MONTH)
GROWTH = inp("Monthly customer growth", 0.05, PCT, "Source: user brief")
CHURN = inp("Monthly churn", 0.02, PCT, "Source: user brief")
PRODUCTS = [  # (name, starting customers ref, price ref)
    (name, inp(f"{name}: starting customers", n, INT), inp(f"{name}: price per month", p, USD))
    for name, n, p in [("Basic", 100, 29), ("Pro", 40, 79)]
]

# ── MONTHLY: one row per month, every cell a formula ───────────────────
ws_m = wb.create_sheet("Monthly")
C = header(ws_m, ["Month"]
           + [f"{name} {kind}" for name, _, _ in PRODUCTS for kind in ("customers", "revenue")]
           + ["Total customers", "Total revenue"])
FIRST, LAST = 2, 2 + MONTHS - 1
for i in range(MONTHS):
    r = FIRST + i
    put(ws_m, f"{C['Month']}{r}", f"=EDATE({START},{i})", MONTH)
    for name, start_customers, price in PRODUCTS:
        cust, rev = C[f"{name} customers"], C[f"{name} revenue"]
        previous = start_customers if i == 0 else f"{cust}{r - 1}"
        put(ws_m, f"{cust}{r}", f"={previous}*(1+{GROWTH}-{CHURN})", INT)
        put(ws_m, f"{rev}{r}", f"={cust}{r}*{price}", USD)
    put(ws_m, f"{C['Total customers']}{r}", "=" + "+".join(f"{C[f'{n} customers']}{r}" for n, _, _ in PRODUCTS), INT)
    put(ws_m, f"{C['Total revenue']}{r}", "=" + "+".join(f"{C[f'{n} revenue']}{r}" for n, _, _ in PRODUCTS), USD)

# ── SUMMARY: one row per 12-month model year, then a total ─────────────
ws_s = wb.create_sheet("Summary")
S = header(ws_s, ["Model year", "Revenue", "Ending customers"])
years = MONTHS // 12
for y in range(years):
    r, a, b = 2 + y, FIRST + 12 * y, FIRST + 12 * y + 11  # b = the year's last Monthly row
    ws_s[f"{S['Model year']}{r}"] = f"Year {y + 1}"
    put(ws_s, f"{S['Revenue']}{r}", f"=SUM(Monthly!{C['Total revenue']}{a}:{C['Total revenue']}{b})", USD)
    put(ws_s, f"{S['Ending customers']}{r}", f"=Monthly!{C['Total customers']}{b}", INT)
total = 2 + years
ws_s[f"{S['Model year']}{total}"] = "Total"
ws_s[f"{S['Model year']}{total}"].font = BOLD
put(ws_s, f"{S['Revenue']}{total}", f"=SUM({S['Revenue']}2:{S['Revenue']}{total - 1})", USD)
ws_s[f"{S['Revenue']}{total}"].font = BOLD

out = sys.argv[1] if len(sys.argv) > 1 else "example_model.xlsx"
wb.save(out)
print(f"written {out}")
