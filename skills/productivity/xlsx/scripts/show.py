"""Print a sheet's cells: each formula next to the value it computed.

    usage: python show.py file.xlsx [Sheet] [A1:H20]

Run recalc.py first; until then every formula's value is None.
"""
import sys

from openpyxl import load_workbook

MAX_CELLS = 300

if len(sys.argv) < 2:
    sys.exit(__doc__)
path = sys.argv[1]
formulas = load_workbook(path)
values = load_workbook(path, data_only=True)
for name in [sys.argv[2]] if len(sys.argv) > 2 else formulas.sheetnames:
    ws, wv = formulas[name], values[name]
    rng = sys.argv[3] if len(sys.argv) > 3 else ws.dimensions
    if ":" not in rng:
        rng = f"{rng}:{rng}"
    print(f"## {name} ({rng})")
    cells = [c for row in ws[rng] for c in row if c.value is not None]
    for cell in cells[:MAX_CELLS]:
        v = cell.value
        if isinstance(v, str) and v.startswith("="):
            print(f"{cell.coordinate}: {v}  ->  {wv[cell.coordinate].value!r}")
        else:
            print(f"{cell.coordinate}: {v!r}")
    if len(cells) > MAX_CELLS:
        print(f"... {len(cells) - MAX_CELLS} more cells; pass a range, e.g. A1:H20")
