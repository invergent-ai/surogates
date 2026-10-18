"""The task suite: what a person asks for, which files they put in the workspace first, and the checks on the result.

Inputs are runtime-office's own test fixtures (``runtime-office/test/fixtures``), uploaded under the names a person
would give them. ``tests/test_tasks.py`` pins the fixture contents the edit checks rely on, so a changed fixture fails
the suite instead of silently moving every edit task's baseline.
"""
from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import NamedTuple

from obench import checks as C
from obench import docs
from obench.docs import Deck, PdfDoc, SheetBook, WordDoc, norm

FIXTURES = Path(__file__).resolve().parents[3] / "runtime-office" / "test" / "fixtures"


@dataclass(frozen=True)
class Input:
    name: str  # the workspace path the person put it at
    fixture: str  # file in FIXTURES
    derive: Callable[[str, str], None] | None = None  # (fixture path, output path): a changed copy

    def source(self) -> Path:
        return FIXTURES / self.fixture


@dataclass(frozen=True)
class Task:
    id: str
    kind: str  # docx | xlsx | pptx | pdf
    mode: str  # create | edit
    prompt: str
    target: str  # workspace path of the document the checks read
    checks: tuple[C.Check, ...]
    inputs: tuple[Input, ...] = ()
    pilot: bool = False
    notes: str = ""

    @property
    def original_input(self) -> Input | None:
        """For an edit task, the input that is the target document."""
        for i in self.inputs:
            if i.name == self.target:
                return i
        return None


REPORT_DOCX = Input("quarterly-report.docx", "sample.docx")
SALES_XLSX = Input("sales.xlsx", "sample.xlsx")
REVIEW_PPTX = Input("quarterly-review.pptx", "sample.pptx")
AGREEMENT_PDF = Input("service-agreement.pdf", "sample.pdf")

# What the samples hold, as the checks of the edit tasks rely on it (pinned by tests/test_tasks.py).
SAMPLE_PDF_TEXT = [
    "Service agreement",
    "The supplier delivers the service from 1 March 2026.",
    "Fees are due within thirty days of each invoice.",
    "Page two: the schedule of services.",
]


# =========================================================================== Word

def _east_first(doc: WordDoc, ctx):
    old = [p.text for p in doc.paragraphs if norm(p.text) == "north grew fastest, then the east."]
    if old:
        return False, "the old sentence is still there"
    for p in doc.paragraphs:
        t = norm(p.text)
        if "fastest" in t and "east" in t and "north" in t and t.index("east") < t.index("north") and p.heading is None \
                and p.list_kind is None:
            return True, f"now: {p.text!r}"
    return False, "no sentence with the east first and the north second"


def _para_count_same(doc: WordDoc, ctx):
    if ctx.original is None:
        return None, "no original"
    a = len([p for p in ctx.original.paragraphs if p.text.strip()])
    b = len([p for p in doc.paragraphs if p.text.strip()])
    return a == b, f"{a} paragraphs -> {b}"


def _para_starts(prefix: str, label: str, claim: tuple = ()):
    def fn(doc: WordDoc, ctx):
        hit = [p.text for p in doc.paragraphs if norm(p.text).startswith(norm(prefix))]
        return bool(hit), (f"{hit[0][:60]!r}" if hit else f"no paragraph starting {prefix!r}")
    return C.Check(label, fn, claim)


def _any_para(options: tuple[str, ...], label: str):
    def fn(doc: WordDoc, ctx):
        for p in doc.paragraphs:
            if any(norm(o) in norm(p.text) for o in options):
                return True, f"{p.text[:60]!r}"
        return False, f"none of {options}"
    return C.Check(label, fn)


def _signed_at_end(name: str):
    def fn(doc: WordDoc, ctx):
        paras = [p for p in doc.paragraphs if p.text.strip()]
        idx = [i for i, p in enumerate(paras) if norm(name) in norm(p.text)]
        if not idx:
            return False, f"{name} not found"
        return idx[-1] >= len(paras) * 0.7, f"last mention at paragraph {idx[-1] + 1} of {len(paras)}"
    return C.Check(f"signed by {name} at the end", fn)


def _no_tables_or_charts(doc: WordDoc, ctx):
    return not doc.tables and not doc.charts, f"{len(doc.tables)} tables, {doc.charts} charts"


def _owner_dates(pairs: list[tuple[str, str]]):
    """Each (owner, day) appears together in one paragraph or one table row: an action with its owner and date."""
    def fn(doc: WordDoc, ctx):
        units = [p.text for p in doc.paragraphs] + [" | ".join(r) for t in doc.tables for r in t.rows]
        missing = []
        for owner, day in pairs:
            if not any(norm(owner) in norm(u) and re.search(rf"(?<!\d){day}(?!\d)", u) for u in units):
                missing.append(f"{owner} by the {day}th")
        return not missing, "every action has its owner and date" if not missing else f"missing: {missing}"
    return C.Check("each action item carries its owner and due date", fn, ("action",))


def _bold_or_heading(phrase: str):
    def fn(doc: WordDoc, ctx):
        for p in doc.paragraphs:
            if norm(phrase) not in norm(p.text):
                continue
            if p.heading is not None:
                return True, f"in a heading: {p.text[:60]!r}"
            if any(r.bold and norm(phrase) in norm(r.text) for r in p.runs) or (
                    p.runs and all(r.bold for r in p.runs if r.text.strip())):
                return True, f"bold: {p.text[:60]!r}"
            return False, f"present but plain: {p.text[:80]!r}"
        return False, f"{phrase!r} not found"
    return C.Check(f"{phrase!r} stands out (bold or a heading)", fn, ("bold", "stand"))


def _question_headings(minimum: int):
    def fn(doc: WordDoc, ctx):
        qs = [p.text for p in doc.headings if p.text.strip().endswith("?")]
        return len(qs) >= minimum, f"{len(qs)} question headings: {qs[:7]}"
    return C.Check(f"at least {minimum} questions as headings", fn, ("heading", "question"))


def _section_after(title: str, after_heading: str, min_words: int):
    def fn(doc: WordDoc, ctx):
        blocks = doc.blocks
        idx = [i for i, b in enumerate(blocks) if isinstance(b, docs.Para) and norm(b.text) == norm(title)]
        if not idx:
            return False, f"no paragraph {title!r}"
        i = idx[0]
        if blocks[i].heading is None:
            return False, f"{title!r} is there but not a heading (style {blocks[i].style!r})"
        anchor = [j for j, b in enumerate(blocks) if isinstance(b, docs.Para) and norm(b.text) == norm(after_heading)]
        if anchor and i < anchor[0]:
            return False, f"{title!r} comes before {after_heading!r}"
        body = []
        for b in blocks[i + 1:]:
            if isinstance(b, docs.Para) and b.heading is not None:
                break
            body.append(b.text if isinstance(b, docs.Para) else " ".join(" ".join(r) for r in b.rows))
        n = docs.words(" ".join(body))
        return n >= min_words, f"heading level {blocks[i].heading}, {n} words under it"
    return C.Check(f"a {title!r} section after {after_heading!r}", fn, ("section", "next steps"))


def _heading_level(title: str, level: int):
    def fn(doc: WordDoc, ctx):
        for p in doc.paragraphs:
            if norm(p.text) == norm(title):
                return p.heading == level, f"level {p.heading}"
        return False, "not found"
    return C.Check(f"{title!r} at the other sections' level (heading {level})", fn)


def _table_is(rows: list[list[str]]):
    def fn(doc: WordDoc, ctx):
        got = [[norm(c) for c in row] for row in doc.tables[0].rows] if doc.tables else None
        want = [[norm(c) for c in r] for r in rows]
        return got == want, f"table {doc.tables[0].rows if doc.tables else None}"
    return C.Check(f"the table reads {rows}", fn, ("14%", "table"))


def _absent(phrase: str):
    def fn(doc, ctx):
        return norm(phrase) not in norm(doc.text()), "gone" if norm(phrase) not in norm(doc.text()) else "still there"
    return C.Check(f"{phrase!r} no longer appears", fn)


def _table_style_kept(doc: WordDoc, ctx):
    style = doc.tables[0].style if doc.tables else None
    return style == "Table Grid", f"style {style!r}"


def _list_continued(items: list[str]):
    """The new item is the list's next number, as the reader sees it: one list going 1, 2, 3."""
    def fn(doc: WordDoc, ctx):
        numbered = [p for p in doc.paragraphs if p.list_kind == "number" and p.text.strip()]
        shown = [f"{p.number}. {p.text}" for p in numbered]
        texts = [norm(p.text) for p in numbered]
        want = [norm(i) for i in items]
        for k in range(len(texts) - len(want) + 1):
            if texts[k:k + len(want)] == want:
                got = [p.number for p in numbered[k:k + len(want)]]
                ok = got == list(range(1, len(want) + 1))
                return ok, f"reads {shown}" + ("" if ok else f" (numbering ids {[p.num_id for p in numbered[k:k + len(want)]]})")
        found = [(p.text, p.list_kind) for p in doc.paragraphs if norm(items[-1]) in norm(p.text)]
        if found:
            return False, f"the item is there but not as the list's next number: {found}; numbered now {shown}"
        return False, f"numbered items now {shown}"
    return C.Check(f"the numbered list continues with {items[-1]!r} as item {len(items)}", fn, ("list", "item", "third"))


def _block_text(b) -> str:
    return b.text if isinstance(b, docs.Para) else "\n".join(" | ".join(r) for r in b.rows)


# --------------------------------------------------------------------------- Word: a long document

OLD_BUDGET, NEW_BUDGET = "45,800", "48,500"


class _Area(NamedTuple):
    """An area of the long report, written as a section of the same shape as the others."""
    title: str
    owner: str
    staff: int
    budget_2026: int
    budget_2027: int
    why: str  # why the budget moved
    measure: str  # what the team is judged on
    was: str  # the measure at the start of 2026
    now: str  # the measure at the end of 2026
    target: str  # the measure's target for 2027
    project: str  # the main project of the year
    first: str  # the project's first month
    last: str  # the project's last month
    risk: str  # the biggest risk
    then: str  # what the team does if it happens
    priorities: tuple[str, ...]  # the year's priorities, as a bulleted list; none for some areas
    quarterly: bool  # a table of planned spending by quarter


_AREAS = {
    3: _Area("Staffing", "head of people", 6, 204900, 212400, "the increase covers the pay review agreed in March",
        "staff turnover", "14%", "11%", "below 10%", "a new pay and grading framework", "February", "July",
        "a wave of resignations after the pay review", "bring the mid-year review forward and use the retention fund",
        ("Publish the new pay bands", "Train every manager on the review", "Bring turnover below 10%"), True),
    4: _Area("Recruitment", "recruitment lead", 3, 41300, 38900,
        "the fall reflects fewer agency fees now that most roles are filled directly",
        "the average time to fill a role", "52 days", "41 days", "35 days", "a careers page with live job listings",
        "January", "April", "a shortage of engineers in the local market", "widen the search to remote candidates", (),
        False),
    5: _Area("Training", "learning manager", 2, 24100, 27600,
        "the increase pays for a leadership course for new managers", "training hours per employee", "11", "16", "20",
        "a leadership course for the twelve new managers", "March", "November", "managers too busy to attend",
        "move the sessions online and record them",
        ("Run the leadership course", "Refresh the safety training", "Track training hours per person"), False),
    6: _Area("Customer support", "support manager", 14, 61800, 64300, "the increase covers two extra weekend shifts",
        "the share of tickets answered within a day", "81%", "88%", "92%", "a self-service help centre", "April",
        "September", "a spike in tickets after the spring release",
        "move two product specialists onto the queue for a month", (), True),
    7: _Area("Sales operations", "sales operations lead", 5, 49700, 51200,
        "the increase is the new pricing tool's licence", "the time from a signed order to its first invoice", "9 days",
        "6 days", "4 days", "moving every quote into the pricing tool", "January", "May",
        "errors in quotes during the move", "check every quote by hand for the first two months",
        ("Move all quotes into the pricing tool", "Invoice within four days", "Clean up the customer records"), False),
    8: _Area("Marketing", "marketing manager", 7, 79200, 73500,
        "the fall comes from running the spring campaign online", "qualified leads per month", "140", "185", "220",
        "an online spring campaign", "February", "June", "a campaign that brings visits but few leads",
        "move the spending to the channels that bring leads within the quarter", (), False),
    9: _Area("Product delivery", "head of delivery", 18, 84600, 88100,
        "the increase pays for a second test environment", "releases shipped on the planned date", "70%", "82%", "90%",
        "a second test environment", "March", "August", "a late release in the busy autumn period",
        "freeze new features four weeks before each release",
        ("Ship every release on its planned date", "Set up the second test environment", "Cut open defects by a third"),
        True),
    10: _Area("Quality", "quality lead", 3, 18400, 19700, "the increase covers the external audit",
        "defects found after a release", "46", "31", "below 25", "the external quality audit", "May", "June",
        "audit findings that need large changes", "plan the changes into the autumn releases", (), False),
    11: _Area("IT systems", "IT manager", 6, 91000, 96700, "the increase is the move of the file servers to the cloud",
        "unplanned downtime", "19 hours", "7 hours", "under 5 hours", "moving the file servers to the cloud", "January",
        "October", "data lost during the move",
        "keep the old servers running beside the new ones until every folder is checked",
        ("Move the file servers", "Replace the oldest laptops", "Halve unplanned downtime"), False),
    12: _Area("Security", "security officer", 2, 29800, 33400, "the increase pays for a yearly penetration test",
        "the share of staff who pass the phishing test", "72%", "86%", "95%", "two-factor login for every system",
        "February", "May", "a successful phishing attack",
        "lock the affected accounts and reset every password in the team", (), True),
    13: _Area("Procurement", "procurement lead", 2, 13500, 12800, "the fall reflects fewer suppliers to manage",
        "the share of spending under contract", "64%", "77%", "85%", "a single list of approved suppliers", "April",
        "July", "a key supplier raising its prices", "invite two other suppliers to quote",
        ("Publish the approved supplier list", "Bring the top twenty suppliers under contract",
         "Pay invoices within thirty days"), False),
    14: _Area("Logistics", "logistics manager", 9, 39900, 41600, "the increase covers higher fuel costs",
        "orders delivered on time", "91%", "94%", "97%", "a route plan for the new delivery van", "March", "May",
        "a delay at the main warehouse", "ship from the second warehouse until the delay clears", (), False),
    15: _Area("Travel", "office manager", 1, 26700, 22300, "the fall comes from holding more meetings online",
        "trips taken per month", "38", "29", "25", "a travel booking tool", "January", "March",
        "trips booked outside the tool", "refund only the trips booked through it", (), True),
    16: _Area("Sustainability", "sustainability lead", 1, 11200, 15900, "the increase pays for the first energy audit",
        "electricity used per employee", "2.9 MWh", "2.6 MWh", "2.3 MWh", "the first energy audit of both offices",
        "June", "September", "an audit that finds the old office cannot be improved",
        "fold the findings into the facilities plan in section 17",
        ("Complete the energy audit", "Switch both offices to green power", "Report emissions for the first time"),
        False),
    18: _Area("Legal and compliance", "legal counsel", 2, 31200, 29500,
        "the fall reflects fewer contract reviews by outside lawyers", "contracts reviewed within five days", "68%",
        "80%", "90%", "standard contract templates for sales", "February", "April",
        "a change in the data protection rules", "review the affected contracts within a month", (), False),
    19: _Area("Business continuity", "operations director", 2, 16000, 17200,
        "the increase pays for a second internet line at each office", "systems restored within four hours in a test",
        "3 of 6", "5 of 6", "6 of 6", "a tested recovery plan for each office", "April", "October",
        "a long power cut at the main office", "move the critical work to the second office", (), True),
}

_SUMMARY = [
    "This report sets out the operations plan for 2027: what each area will do, what it will cost and how its progress "
    "will be measured. It has twenty sections, from staffing to the approvals, and replaces the plan issued in January.",
    "The three largest budget lines are staffing (€212,400, section 3), IT systems (€96,700, section 11) and product "
    f"delivery (€88,100, section 9). The facilities budget, €{OLD_BUDGET}, pays for the move of the support team to the "
    "new office and is set out in section 17.",
    "Each section names the team that owns the area, its budget for 2027 against 2026, the measure it is judged on and "
    "its main project for the year. Every figure in this report is in euros and excludes VAT.",
]
_SCOPE = (
    ["The plan covers every team in the company except the two subsidiaries, which publish their own plans. It runs "
     "from 1 January to 31 December 2027 and is reviewed once a quarter by the operations board.",
     "The following are outside its scope:"],
    ["Capital projects above €250,000, which the board approves one by one",
     "The subsidiaries' own budgets",
     "Pay for the leadership team, which the remuneration committee sets"],
    ["Where a figure here differs from the finance system, the finance system is right, and this report is corrected "
     "in its next edition."],
)
_FACILITIES = [
    f"Facilities is run by the facilities manager, with a team of 4 people. Its budget for 2027 is €{OLD_BUDGET}, "
    "against €43,900 in 2026; the increase pays for the move of the support team to the new office on the second "
    "floor.",
    "The team is judged on the share of maintenance requests closed within a week. Over 2026 it moved from 74% to 83%, "
    "and the target for 2027 is 90%. The main project for the year is the move itself, planned for the summer so that "
    "support keeps its weekend shifts.",
    "The biggest risk is a delay in the fit-out of the new floor. If it happens, the support team stays where it is "
    "until the work is finished, and the move is planned again with the support manager.",
    "The facilities budget over the last three years:",
]
_FACILITIES_TABLE = [["Year", "Budget", "Main item"], ["2025", "€41,200", "Roof repairs"],
                     ["2026", "€43,900", "New heating"], ["2027", f"€{OLD_BUDGET}", "Office move"]]
_APPROVALS = [
    f"The finance committee approved the facilities budget of €{OLD_BUDGET} on 2 October, together with the budgets in "
    "sections 3 to 16. The budgets in sections 18 and 19 were approved by the board on 9 October.",
    "A change to any figure in this report after its approval is made in the section the figure belongs to, with a "
    "sentence saying when it was made. The operations board hears of every change at its next meeting.",
    "This report was prepared by the operations team and is next reviewed at the end of the first quarter of 2027.",
]


def _area_blocks(n: int) -> list[tuple[str, object]]:
    a = _AREAS[n]
    team = "one person" if a.staff == 1 else f"{a.staff} people"
    out: list[tuple[str, object]] = [
        ("h", f"{n}. {a.title}"),
        ("p", f"{a.title} is run by the {a.owner}, with a team of {team}. Its budget for 2027 is €{a.budget_2027:,}, "
              f"against €{a.budget_2026:,} in 2026; {a.why}. The team reports to the operations board every month and "
              "to the whole company once a quarter, at the quarterly review."),
        ("p", f"The team is judged on {a.measure}. Over 2026 it moved from {a.was} to {a.now}, and the target for 2027 "
              f"is {a.target}. The figure is read on the first working day of each month and goes into the monthly "
              "operations pack, so a change in one month shows in the next pack and in the next edition of this report."),
        ("p", f"The main project for the year is {a.project}. It starts in {a.first} and is due to end in {a.last}. "
              f"The {a.owner} leads it and reports on it at each board meeting. If it slips by more than a month, the "
              "team agrees a new date with the operations board, and the new date is recorded in this section."),
    ]
    if a.priorities:
        out.append(("p", "The team's priorities for 2027, in order:"))
        out += [("li", p) for p in a.priorities]
    out += [
        ("p", f"The biggest risk for this area is {a.risk}. If it happens, the team will {a.then}. The risk is reviewed "
              f"at every quarterly review, and its owner is the {a.owner}."),
        ("p", f"Spending on {a.title.lower()} is checked against the plan at the end of each quarter. An overspend of "
              f"more than {3 + n % 4} per cent needs the finance team's agreement before the next payment is made, and "
              "an underspend of the same size is offered back to the central budget for other teams to use."),
    ]
    if a.quarterly:
        q = [round(a.budget_2027 * share / 100) * 100 for share in (0.23, 0.25, 0.26)]
        q.append(a.budget_2027 - sum(q))
        out.append(("p", "Planned spending by quarter:"))
        out.append(("table", [["Quarter", "Planned spend"]] + [[f"Q{i}", f"€{v:,}"] for i, v in enumerate(q, 1)]))
    return out


def _long_report_blocks() -> list[tuple[str, object]]:
    """The long report, block by block: ("h", heading), ("p", paragraph), ("li", bullet), ("table", rows)."""
    out: list[tuple[str, object]] = [("h", "1. Summary")] + [("p", t) for t in _SUMMARY]
    intro, bullets, outro = _SCOPE
    out += [("h", "2. Scope of this report")] + [("p", t) for t in intro] + [("li", t) for t in bullets]
    out += [("p", t) for t in outro]
    for n in range(3, 20):
        if n == 17:
            out += [("h", "17. Facilities")] + [("p", t) for t in _FACILITIES] + [("table", _FACILITIES_TABLE)]
        else:
            out += _area_blocks(n)
    out += [("h", "20. Approvals")] + [("p", t) for t in _APPROVALS]
    return out


def _derive_long_report(src: str, dst: str) -> None:
    """sample.docx's body replaced by a twelve-page operations report in its styles: twenty numbered sections of
    figures, with section 17's budget (€45,800) in the summary, in the section's text and table, and in the approvals."""
    import docx

    d = docx.Document(src)
    body = d.element.body
    for el in list(body):
        if el.tag != docs.W + "sectPr":
            body.remove(el)
    d.add_heading("Operations report 2027", level=1)
    for kind, value in _long_report_blocks():
        if kind == "h":
            d.add_heading(value, level=2)
        elif kind == "p":
            d.add_paragraph(value)
        elif kind == "li":
            d.add_paragraph(value, style="List Bullet")
        else:
            table = d.add_table(rows=len(value), cols=len(value[0]))
            table.style = "Table Grid"
            for r, row in enumerate(value):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
    d.save(dst)


LONG_SECTION = 17  # the section whose budget is wrong


def _old_budget_gone(doc: WordDoc, ctx):
    """The old figure is nowhere but in the correction sentence after the section's table, which may say what it was."""
    note, _ = _section_note(doc, LONG_SECTION)
    quoted = note if note is not None and _correction_sentence(note.text) else None
    left = [_block_text(b)[:70] for b in doc.blocks
            if b is not quoted and C.has_text(_block_text(b), int(OLD_BUDGET.replace(",", "")))]
    if left:
        return False, f"still in {len(left)} block(s): {left[:4]}"
    said = quoted is not None and C.has_text(quoted.text, int(OLD_BUDGET.replace(",", "")))
    return True, "gone" + (" (the correction sentence quotes it)" if said else "")


def _with_new_budget(b):
    if isinstance(b, docs.Para):
        return replace(b, text=b.text.replace(OLD_BUDGET, NEW_BUDGET))
    return docs.Table([[c.replace(OLD_BUDGET, NEW_BUDGET) for c in r] for r in b.rows], b.style)


def _new_budget_in_its_places(doc: WordDoc, ctx):
    """Each block that gave the old figure reads as it did with the new one, in the original order."""
    if ctx.original is None:
        return None, "no original"
    places = [b for b in ctx.original.blocks if OLD_BUDGET in _block_text(b)]
    have = [C.w_sig(b) for b in doc.blocks]
    missing, pos = [], 0
    for b in places:
        try:
            pos = have.index(C.w_sig(_with_new_budget(b)), pos) + 1
        except ValueError:
            missing.append(_block_text(b)[:70])
    if missing:
        return False, f"{len(missing)} of {len(places)} places do not read as before with {NEW_BUDGET}: {missing}"
    return True, f"all {len(places)} places now give {NEW_BUDGET}"


# 14 October as people write it: 14 October, 14th of October, 14-Oct, October 14, Oct. 14th, 14/10, 14.10, 14-10 (a year
# may follow), 10/14 (US), 2026-10-14.
_OCT_14 = re.compile(r"\b14(?:th)?(?: of)?[ -]oct(?:ober)?\b|\boct(?:ober)?\.?[ -]?14(?:th)?\b|\b14[/.-]10\b|\b10/14\b"
                     r"|\b\d{4}-10-14\b")
# Words that say a figure was put right; "incorrect" is not one of them.
_CORRECTED = re.compile(r"\b(?:correct(?:ed|ions?|s)?|updated?|revised?|amend(?:ed|ments?|s)?)\b")


def _sentences(text: str) -> list[str]:
    return [s for s in re.split(r"(?<=[.!?])\s+(?=[A-Z])", (text or "").strip()) if s.strip()]


def _correction_sentence(text: str) -> bool:
    """One sentence saying the figure was corrected (or updated, revised, amended) on 14 October."""
    t = norm(text)
    return bool(_CORRECTED.search(t) and _OCT_14.search(t)) and len(_sentences(text)) == 1


def _section_note(doc: WordDoc, section: int) -> tuple[docs.Para | None, str]:
    """Where a note after the section's table is: the first non-empty block after the section's last table and before
    the next heading, when it is a plain paragraph. Else None, and why."""
    blocks = doc.blocks
    heads = [i for i, b in enumerate(blocks) if isinstance(b, docs.Para) and b.heading is not None]
    start = next((i for i in heads if norm(blocks[i].text).startswith(f"{section}.")), None)
    if start is None:
        return None, f"no section {section} heading"
    end = next((i for i in heads if i > start), len(blocks))
    tables = [i for i in range(start + 1, end) if isinstance(blocks[i], docs.Table)]
    if not tables:
        return None, f"section {section} has no table"
    after = [b for b in blocks[tables[-1] + 1:end] if not (isinstance(b, docs.Para) and not b.text.strip())]
    if not after:
        return None, f"nothing between section {section}'s table and the next heading"
    b = after[0]
    if not isinstance(b, docs.Para) or b.list_kind is not None:
        return None, f"after the table: {'a table' if isinstance(b, docs.Table) else 'a list item'}"
    return b, ""


def _note_after_table(section: int):
    def fn(doc: WordDoc, ctx):
        b, why = _section_note(doc, section)
        if b is None:
            return False, why
        t = norm(b.text)
        if not (_CORRECTED.search(t) and _OCT_14.search(t)):
            return False, f"the paragraph after the table does not say it was corrected on 14 October: {b.text[:100]!r}"
        n = len(_sentences(b.text))
        return n == 1, f"{b.text[:100]!r}" + ("" if n == 1 else f" ({n} sentences)")
    return C.Check(f"a sentence right after section {section}'s table says it was corrected on 14 October", fn,
                   ("sentence", "correct"))


def _one_paragraph_added(doc: WordDoc, ctx):
    if ctx.original is None:
        return None, "no original"
    a = len([p for p in ctx.original.paragraphs if p.text.strip()])
    b = len([p for p in doc.paragraphs if p.text.strip()])
    return b == a + 1, f"{a} paragraphs -> {b}"


# --------------------------------------------------------------------------- Word: a table with a Total row

REGION_TABLE = [["Region", "Q1", "Q2"], ["North", "120", "140"], ["East", "90", "100"], ["Total", "210", "240"]]


def _derive_region_table(src: str, dst: str) -> None:
    """sample.docx with its table made Q1 and Q2 by region with a Total row, in the table's own style."""
    import docx

    d = docx.Document(src)
    old = d.tables[0]
    new = d.add_table(rows=len(REGION_TABLE), cols=len(REGION_TABLE[0]))
    new.style = old.style
    for r, row in enumerate(REGION_TABLE):
        for c, text in enumerate(row):
            new.cell(r, c).text = text
    old._tbl.addprevious(new._tbl)
    old._tbl.getparent().remove(old._tbl)
    d.save(dst)


def _row_above_total(name: str, values: list[str]):
    def fn(doc: WordDoc, ctx):
        if not doc.tables:
            return False, "no table"
        rows = doc.tables[0].rows
        firsts = [norm(r[0]) if r else "" for r in rows]
        if norm(name) not in firsts:
            return False, f"no {name} row: {rows}"
        i = firsts.index(norm(name))
        if [norm(c) for c in rows[i][1:]] != [norm(v) for v in values]:
            return False, f"the {name} row reads {rows[i]}"
        if "total" in firsts and i > firsts.index("total"):
            return False, f"the {name} row is below the Total row: {rows}"
        return True, f"row {i + 1}: {rows[i]}"
    return C.Check(f"a {name} row reading {' and '.join(values)} above the Total row", fn, ("row", name.lower()))


def _total_row(values: list[str]):
    def fn(doc: WordDoc, ctx):
        if not doc.tables:
            return False, "no table"
        rows = doc.tables[0].rows
        totals = [r for r in rows if r and norm(r[0]) == "total"]
        if not totals:
            return False, f"no Total row: {rows}"
        if rows[-1] is not totals[0]:
            return False, f"the Total row {totals[0]} is not the last row: {rows}"
        return [norm(c) for c in totals[0][1:]] == [norm(v) for v in values], f"the Total row reads {totals[0]}"
    return C.Check(f"the Total row reads {' and '.join(values)}", fn, ("total",))


def _rows_kept(rows: list[list[str]]):
    """The header stays the first row; each other row is still there as it was, in any order (a new row may go first,
    the regions may be sorted)."""
    def fn(doc: WordDoc, ctx):
        got = doc.tables[0].rows if doc.tables else []
        have = [[norm(c) for c in r] for r in got]
        want = [[norm(c) for c in r] for r in rows]
        if not have or have[0] != want[0]:
            return False, f"the header is {got[0] if got else None}"
        missing = [r for r, w in zip(rows[1:], want[1:]) if w not in have[1:]]
        return not missing, "kept" if not missing else f"changed or gone: {missing}; the table is {got}"
    return C.Check(f"the header and the {' and '.join(r[0] for r in rows[1:])} rows are kept", fn)


# --------------------------------------------------------------------------- Word: comment threads

# Questions whose answers are in sample.docx's own text, each with what its answer must state (a pattern over the
# reply, normalized) and how the check names it; then the notes that ask nothing.
THREAD_QUESTIONS = {
    "How many engineers do we plan to hire?": (r"\b(?:two|2)\b", "two"),
    # 12, 12%, 12 %, 12 pct, 12 per cent, twelve percent; not 112 or 12.5.
    "How much did North grow this quarter?": (r"(?<![\d.,])12(?:[.,]0+)?(?![\d]|[.,]\d)|\btwelve\b", "12%"),
    "Which city is the new office in?": (r"\blisbon\b", "Lisbon"),
}
THREAD_NOTES = ("Keep 'revenue target' in bold, as finance asked.",
                "The regional leads signed off these figures on 3 October.")
# (the paragraph, the run the comment is on or None for the whole paragraph, the comment, its author)
_THREADS = [
    ("Quarterly report", None, "Which city is the new office in?", "Marta"),
    ("This report covers the quarter. The revenue target was met in all regions.", 1, THREAD_NOTES[0], "Jon"),
    ("North grew fastest, then the east.", None, "How much did North grow this quarter?", "Jon"),
    ("East", None, THREAD_NOTES[1], "Marta"),
    ("Hiring could slip a quarter.", None, "How many engineers do we plan to hire?", "Marta"),
]
MAX_REPLY = 400


def _derive_threads(src: str, dst: str) -> None:
    """sample.docx with five open comment threads: three questions its text answers, two notes."""
    import datetime as dt

    import docx

    d = docx.Document(src)
    for text, run, comment, author in _THREADS:
        para = next(p for p in d.paragraphs if p.text == text)
        runs = para.runs if run is None else [para.runs[run]]
        c = d.add_comment(runs, text=comment, author=author, initials=author[0])
        c._comment_elm.date = dt.datetime(2026, 10, 15, 9, 0, tzinfo=dt.timezone.utc)
    d.save(dst)


def _thread(doc: WordDoc, text: str):
    return next((c for c in doc.threads if norm(c.text) == norm(text)), None)


def _answered(question: str):
    pattern, fact = THREAD_QUESTIONS[question]

    def fn(doc: WordDoc, ctx):
        t = _thread(doc, question)
        if t is None:
            return False, "the question's comment is gone or changed"
        if not t.replies:
            loose = [c.text[:60] for c in doc.threads if c is not t and re.search(pattern, norm(c.text))]
            return False, "no reply" + (f"; a new thread says it instead: {loose}" if loose else "")
        long = [len(r.text) for r in t.replies if len(r.text) > MAX_REPLY]
        if long:
            return False, f"a reply of {long[0]} characters ({MAX_REPLY} at most)"
        if not any(re.search(pattern, norm(r.text)) for r in t.replies):
            return False, f"the replies do not say {fact}: {[r.text[:80] for r in t.replies]}"
        return True, f"{t.replies[-1].text[:80]!r}"
    return C.Check(f"the reply to {question!r} says {fact}", fn, ("answer", "repl"))


def _questions_resolved(doc: WordDoc, ctx):
    bad = []
    for q in THREAD_QUESTIONS:
        t = _thread(doc, q)
        if t is None or not t.resolved:
            bad.append(f"{q!r} {'gone' if t is None else 'open'}")
    return not bad, "all three resolved" if not bad else "; ".join(bad)


def _notes_left(doc: WordDoc, ctx):
    if ctx.original is None:
        return None, "no original"
    bad = []
    for note in THREAD_NOTES:
        was, now = _thread(ctx.original, note), _thread(doc, note)
        if now is None:
            bad.append(f"{note[:30]!r} gone or changed")
            continue
        if now.replies:
            bad.append(f"{note[:30]!r} has a reply: {now.replies[0].text[:60]!r}")
        if now.resolved:
            bad.append(f"{note[:30]!r} resolved")
        if was is not None and (norm(now.anchor) != norm(was.anchor) or now.author != was.author):
            bad.append(f"{note[:30]!r} now on {now.anchor!r} by {now.author!r}")
    return not bad, "both open, unanswered, as they were" if not bad else "; ".join(bad)


WORD_TASKS = [
    Task(
        id="docx-report",
        kind="docx",
        mode="create",
        pilot=True,
        target="four-day-week-report.docx",
        prompt=(
            "Can you write up a two-page report on our four-day week pilot in a new Word doc, four-day-week-report.docx? "
            "Here's what we have: 40 staff in 3 teams worked four days a week from March to May 2026. Output per person "
            "went up 6%, sick days dropped from 31 to 18 over the three months, and customer satisfaction held at 4.4 "
            "out of 5. Overtime costs fell by €12,400. Two things went badly: clients who call on Fridays couldn't reach "
            "anyone, and handovers between the teams got messy. We want to recommend rolling it out to all 120 staff from "
            "September, with a Friday on-call rota. Please use proper headings, put the problems in a bulleted list, add "
            "a table comparing before and after, and a chart of sick days before vs after."
        ),
        checks=(
            C.w_headings(3),
            C.w_words(600, 1400),
            C.w_list("bullet", 2),
            C.w_table(2, 2, containing=[31, 18]),
            C.w_chart(),
            C.contains_all([6, 31, 18, 4.4, 12400, "September", 120, ("on-call", "on call")]),
            C.w_no_pending(),
        ),
    ),
    Task(
        id="docx-reword",
        kind="docx",
        mode="edit",
        pilot=True,
        inputs=(REPORT_DOCX,),
        target=REPORT_DOCX.name,
        prompt=(
            "In quarterly-report.docx, the sentence 'North grew fastest, then the east.' is the wrong way round - "
            "the east grew fastest, then the north. Can you fix it? Don't change anything else."
        ),
        checks=(
            C.Check("the sentence now puts the east first", _east_first, ("east",)),
            C.w_keeps_rest(changed=lambda b: isinstance(b, docs.Para) and norm(b.text) == "north grew fastest, then the east."),
            C.Check("no paragraph added or removed", _para_count_same),
            C.w_bold("revenue target"),
        ),
    ),
    Task(
        id="docx-letter",
        kind="docx",
        mode="create",
        target="lease-termination-letter.docx",
        prompt=(
            "Write a letter to our landlord, Mr Daniel Costa of Costa Properties, Rua Augusta 112, Lisbon, saying we're "
            "ending the lease on our office at Rua do Ouro 45 on 31 January 2027, as the contract's 90-day notice requires. "
            "Ask him to arrange the final inspection in the week of 18 January and to return our €4,500 deposit to the "
            "account he has on file. It's from Ana Ribeiro, Operations Manager, Northwind Lda. Save it as "
            "lease-termination-letter.docx."
        ),
        checks=(
            C.contains_all(["Costa", "Rua Augusta 112", "Rua do Ouro 45", ("31 January 2027", "January 31, 2027"),
                            4500, "18 January", "Ana Ribeiro", "Operations Manager", "Northwind", 90]),
            _para_starts("Dear", "opens with a salutation"),
            _any_para(("sincerely", "kind regards", "yours faithfully", "best regards", "regards", "yours truly"),
                      "closes with a sign-off"),
            _signed_at_end("Ana Ribeiro"),
            C.w_words(80, 450),
            C.Check("no table or chart in a letter", _no_tables_or_charts),
            C.w_no_pending(),
        ),
    ),
    Task(
        id="docx-minutes",
        kind="docx",
        mode="create",
        target="minutes-2026-10-14.docx",
        prompt=(
            "Can you turn my notes into proper minutes in minutes-2026-10-14.docx? Product sync, 14 Oct. There: Maria (PM), "
            "Jon (eng), Priya (design), Tom (sales). Decided: ship v2.3 on 28 Oct, and drop the dark-mode toggle from this "
            "release. Discussed: onboarding drop-off is 38% at step 2, Priya will mock up a shorter flow. Actions: Jon fixes "
            "the export bug by 21 Oct, Priya shares the onboarding mockups by 18 Oct, Tom tells the 3 pilot customers about "
            "the date by 16 Oct. Next meeting 21 Oct."
        ),
        checks=(
            C.w_headings(2),
            C.contains_all(["Maria", "Jon", "Priya", "Tom", "2.3", ("28 Oct", "28 October", "28/10"), "38%",
                            ("dark-mode", "dark mode"), "export"]),
            _owner_dates([("Jon", "21"), ("Priya", "18"), ("Tom", "16")]),
            C.w_words(80, 600),
            C.w_no_pending(),
        ),
    ),
    Task(
        id="docx-guide",
        kind="docx",
        mode="create",
        target="time-off-guide.docx",
        prompt=(
            "Write a one-page guide for staff on how to request time off, time-off-guide.docx. The steps: submit the request "
            "in the HR portal at least 10 working days ahead; your manager approves or declines within 3 working days; HR "
            "confirms by email; then add it to the team calendar. Include a table of the leave types: annual 25 days, sick "
            "10, parental 20, unpaid up to 30. And make it stand out that unused annual leave over 5 days expires on 31 March."
        ),
        checks=(
            C.w_headings(1),
            C.w_list("number", 4),
            C.w_table(4, 2, containing=["annual", 25, "sick", 10, "parental", 20, "unpaid", 30]),
            _bold_or_heading("31 March"),
            C.w_words(120, 650),
            C.w_no_pending(),
        ),
    ),
    Task(
        id="docx-proposal",
        kind="docx",
        mode="create",
        target="website-proposal.docx",
        prompt=(
            "Draft a one-page proposal for the board to redo our website, website-proposal.docx. Background: the site is 6 "
            "years old and 64% of visits are on mobile, where it's slow. Scope: a new design, moving to a new CMS, and fixing "
            "accessibility. Timeline: discovery 2 weeks, design 4 weeks, build 6 weeks, launch 1 week. Costs: design "
            "€18,000, build €32,000, CMS licences €4,800 a year, contingency €5,000. Put the timeline and the costs in "
            "tables, and give the total cost."
        ),
        checks=(
            C.w_headings(3),
            C.w_table(2, 2, count=2),
            C.w_table(4, 2, containing=["discovery", "design", "build", "launch"]),
            C.w_table(4, 2, containing=[18000, 32000, 4800, 5000], label="a cost table holding all four costs"),
            C.contains_all([59800], name="gives the total cost (59,800)"),
            C.contains_all(["64%", 6]),
            C.w_words(150, 650),
            C.w_no_pending(),
        ),
    ),
    Task(
        id="docx-faq",
        kind="docx",
        mode="create",
        target="pricing-faq.docx",
        prompt=(
            "We're raising the Basic plan from €9 to €12 a month on 1 January. Write a customer FAQ, pricing-faq.docx, with "
            "six questions: why we're doing it, when it happens, what existing customers pay (they keep €9 until 1 April), "
            "whether features change (no - and we're adding two-factor login and daily backups), how to cancel (Settings > "
            "Billing, no fee), and who to contact (billing@shelfie.example). Each question as a heading with a short answer."
        ),
        checks=(
            _question_headings(6),
            C.contains_all([("€9", "9 €", "€ 9", "9 euros", "EUR 9"), ("€12", "12 €", "€ 12", "12 euros", "EUR 12"),
                            "1 January", "1 April", ("two-factor", "2fa", "two factor"), "backup",
                            "billing@shelfie.example", "Billing"]),
            C.w_words(150, 800),
            C.w_no_pending(),
        ),
    ),
    Task(
        id="docx-add-section",
        kind="docx",
        mode="edit",
        inputs=(REPORT_DOCX,),
        target=REPORT_DOCX.name,
        prompt=(
            "Add a section called 'Next steps' to the end of quarterly-report.docx, after Risks, with two or three "
            "sentences on what we'll do next quarter."
        ),
        checks=(
            _section_after("Next steps", "Hiring could slip a quarter.", 15),
            _heading_level("Next steps", 2),
            C.w_keeps_rest(),
            C.w_bold("revenue target"),
            C.w_no_pending(),
        ),
    ),
    Task(
        id="docx-fix-table",
        kind="docx",
        mode="edit",
        inputs=(REPORT_DOCX,),
        target=REPORT_DOCX.name,
        prompt="The growth figure for North in the table in quarterly-report.docx is wrong - it should be 14%, not 12%.",
        checks=(
            _table_is([["Region", "Growth"], ["North", "14%"]]),
            _absent("12%"),
            C.Check("table style kept", _table_style_kept),
            C.w_keeps_rest(changed=lambda b: isinstance(b, docs.Table)),
        ),
    ),
    Task(
        id="docx-continue-list",
        kind="docx",
        mode="edit",
        inputs=(REPORT_DOCX,),
        target=REPORT_DOCX.name,
        prompt="Add 'Renew the support contract' as a third item on the numbered list in quarterly-report.docx.",
        checks=(
            _list_continued(["Hire two engineers", "Open the Lisbon office", "Renew the support contract"]),
            C.w_keeps_rest(),
            C.w_list("bullet", 2),
        ),
    ),
    Task(
        id="docx-bold",
        kind="docx",
        mode="edit",
        inputs=(REPORT_DOCX,),
        target=REPORT_DOCX.name,
        prompt="In quarterly-report.docx, make 'Hiring could slip a quarter' bold.",
        checks=(
            C.w_bold("Hiring could slip a quarter"),
            C.w_bold("revenue target"),
            C.w_bold("This report covers the quarter", bold=False),
            C.w_keeps_rest(),
            C.Check("no paragraph added or removed", _para_count_same),
        ),
    ),
    Task(
        id="docx-long-edit",
        kind="docx",
        mode="edit",
        inputs=(Input("operations-report.docx", "sample.docx", derive=_derive_long_report),),
        target="operations-report.docx",
        prompt=(
            "In section 17 of operations-report.docx the budget is wrong: it is €48,500, not €45,800. Fix it "
            "everywhere it appears, and add one sentence after the section's table saying the figure was corrected "
            "on 14 October."
        ),
        notes=(
            "Input: sample.docx rebuilt as a twelve-page operations report in its styles, twenty numbered sections; "
            "€45,800 is in the summary, section 17's text and table, and the approvals."
        ),
        checks=(
            C.Check(f"'{OLD_BUDGET}' no longer appears anywhere except in the correction sentence", _old_budget_gone,
                    ("everywhere", "occurrence")),
            C.Check(f"{NEW_BUDGET} stands in every place {OLD_BUDGET} was", _new_budget_in_its_places,
                    ("everywhere", NEW_BUDGET)),
            _note_after_table(LONG_SECTION),
            C.w_keeps_rest(changed=lambda b: OLD_BUDGET in _block_text(b)),
            C.Check("one paragraph added, none removed", _one_paragraph_added),
        ),
    ),
    Task(
        id="docx-table-edit",
        kind="docx",
        mode="edit",
        inputs=(Input("quarterly-report.docx", "sample.docx", derive=_derive_region_table),),
        target="quarterly-report.docx",
        prompt="Add a row for South (Q1 70, Q2 65) to the table in quarterly-report.docx and fix its Total row.",
        notes="Input: sample.docx with its table made Region, Q1, Q2: North 120, 140; East 90, 100; Total 210, 240.",
        checks=(
            _row_above_total("South", ["70", "65"]),
            _total_row(["280", "305"]),
            _rows_kept(REGION_TABLE[:3]),
            C.Check("table style kept", _table_style_kept),
            C.w_keeps_rest(changed=lambda b: isinstance(b, docs.Table)),
        ),
    ),
    Task(
        id="docx-many-threads",
        kind="docx",
        mode="edit",
        inputs=(Input("quarterly-report.docx", "sample.docx", derive=_derive_threads),),
        target="quarterly-report.docx",
        prompt=(
            "Reply to the three questions in the comments on quarterly-report.docx and resolve them; leave the other "
            "two."
        ),
        notes=(
            "Input: sample.docx with five open comment threads: three questions its text answers (engineers to hire: "
            "two; North's growth: 12%; the new office: Lisbon) and two notes that ask nothing."
        ),
        checks=(
            *(_answered(q) for q in THREAD_QUESTIONS),
            C.Check("the three questions are resolved", _questions_resolved, ("resolv",)),
            C.Check("the two notes are left open, unanswered and as they were", _notes_left),
            C.w_keeps_rest(),
        ),
    ),
]


# =========================================================================== Spreadsheet

def _derive_wrong_east_total(src: str, dst: str) -> None:
    """sample.xlsx with East's total summing the wrong cells (=B3+B2: 210, not 200)."""
    import openpyxl

    book = openpyxl.load_workbook(src)
    book["Sales"]["D3"] = "=B3+B2"
    book.save(dst)


def _values_on(sheet: str, values: list[float]):
    def fn(book: SheetBook, ctx):
        name = book.sheet_named(sheet)
        if name is None:
            return False, f"no sheet {sheet}"
        have = [c.raw for c in book.cells[name].values() if isinstance(c.raw, (int, float))]
        missing = [v for v in values if not any(abs(h - v) < 1e-6 for h in have)]
        return not missing, "all there" if not missing else f"missing on {name}: {missing}"
    return C.Check(f"{sheet} holds {values}", fn)


def _formula_on(sheet: str, value: float, label: str, must_ref_other_sheet: bool = False):
    def fn(book: SheetBook, ctx):
        name = book.sheet_named(sheet)
        if name is None:
            return False, f"no sheet {sheet}"
        vals, problems = C.x_eval_all(book)
        for (s, r), v in vals.items():
            if s == name and isinstance(v, (int, float)) and abs(v - value) < 0.01:
                f = book.raw(s, r)
                if must_ref_other_sheet and "!" not in f:
                    return False, f"{s}!{r} {f} does not point at the month sheets"
                return True, f"{s}!{r} {f}"
        return (None if problems else False), f"no formula on {name} gives {value:g}" + (f"; {problems[:2]}" if problems else "")
    return C.Check(f"{label} = {value:g} by formula on {sheet}", fn, ("formula", "total", "summary"))


def _text_cells(items: list):
    def fn(book: SheetBook, ctx):
        text = " ".join(str(c.raw) for s in book.sheets for c in book.cells[s].values())
        missing = [i for i in items if not C.has_text(text, i)]
        return not missing, "all present" if not missing else f"missing: {missing}"
    return C.Check(f"mentions {items}", fn)


SALES_ROWS = {"Lisbon": (412, 48200), "Porto": (530, 61750), "Madrid": (351, 39900), "Barcelona": (618, 72300),
              "Valencia": (160, 18450), "Seville": (233, 25600)}


def _sales_rows(book: SheetBook) -> dict[str, tuple[str, int, list]]:
    """city -> (sheet, row, numbers on that row)."""
    out = {}
    vals, _ = C.x_eval_all(book)
    for s in book.sheets:
        rows: dict[int, list] = {}
        for ref, c in book.cells[s].items():
            m = re.match(r"([A-Z]+)(\d+)$", ref)
            rows.setdefault(int(m.group(2)), []).append((ref, c))
        for r, cells in rows.items():
            for ref, c in cells:
                if isinstance(c.raw, str) and c.raw.strip() in SALES_ROWS:
                    nums = [vals.get((s, rr), cc.raw) for rr, cc in cells]
                    out[c.raw.strip()] = (s, r, [n for n in nums if isinstance(n, (int, float))])
    return out


def _sales_complete(book: SheetBook, ctx):
    rows = _sales_rows(book)
    bad = []
    for city, (units, revenue) in SALES_ROWS.items():
        if city not in rows:
            bad.append(f"{city} missing")
            continue
        nums = rows[city][2]
        if not (any(abs(n - units) < 1e-6 for n in nums) and any(abs(n - revenue) < 1e-6 for n in nums)):
            bad.append(f"{city} row has {nums}")
    return not bad, "all six rows, figures as numbers" if not bad else f"{bad}"


def _sales_sorted(book: SheetBook, ctx):
    rows = _sales_rows(book)
    if len(rows) < 6:
        return False, f"only {len(rows)} cities found"
    order = [c for c, _ in sorted(rows.items(), key=lambda kv: kv[1][1])]
    want = [c for c, _ in sorted(SALES_ROWS.items(), key=lambda kv: -kv[1][1])]
    return order == want, f"rows top to bottom: {order}"


def _header_bold(book: SheetBook, ctx):
    rows = _sales_rows(book)
    if not rows:
        return False, "no data rows"
    sheet = next(iter(rows.values()))[0]
    first = min(r for _, r, _ in rows.values())
    header = [c for ref, c in book.cells[sheet].items() if re.match(rf"[A-Z]+{first - 1}$", ref) and c.raw not in (None, "")]
    if not header:
        return False, f"no header row above row {first}"
    return all(c.bold for c in header), f"header {[c.raw for c in header]} bold={[c.bold for c in header]}"


def _chart_covers(points: int):
    def fn(book: SheetBook, ctx):
        from obench.formula import expand_range, split_ref

        best = 0
        for ch in book.charts:
            for ref in ch.refs:
                try:
                    _, cells = split_ref(ref, "")
                    best = max(best, len(expand_range(cells)))
                except Exception:  # noqa: BLE001
                    continue
        return best >= points, f"largest series range holds {best} cells"
    return C.Check(f"the chart covers all {points} points", fn, ("chart",))


def _dates_real(minimum: int):
    def fn(book: SheetBook, ctx):
        import datetime as dt

        dates = [c.raw for s in book.sheets for c in book.cells[s].values() if isinstance(c.raw, dt.datetime)]
        nov = [d for d in dates if d.year == 2026 and d.month == 11]
        texty = [c.raw for s in book.sheets for c in book.cells[s].values()
                 if isinstance(c.raw, str) and re.search(r"\bnov", c.raw, re.IGNORECASE) and re.search(r"\d", c.raw)]
        return len(nov) >= minimum, f"{len(nov)} real dates in Nov 2026" + (f"; dates typed as text: {texty[:3]}" if texty else "")
    return C.Check(f"start and end are real dates ({minimum}+ in Nov 2026)", fn, ("date",))


DURATIONS = {"end-start": [4, 9, 11, 7, 0], "inclusive": [5, 10, 12, 8, 1], "working days": [4, 7, 9, 6, 1]}


def _durations(book: SheetBook, ctx):
    vals, problems = C.x_eval_all(book)
    got = sorted(v for v in vals.values() if isinstance(v, (int, float)))
    for name, want in DURATIONS.items():
        pool = list(got)
        ok = True
        for w in want:
            hit = next((g for g in pool if abs(g - w) < 1e-6), None)
            if hit is None:
                ok = False
                break
            pool.remove(hit)
        if ok:
            return True, f"durations by formula ({name}): {want}"
    if problems:
        return None, f"cannot evaluate: {problems[:2]}"
    return False, f"formula values {got[:12]}"


def _growth_column(book: SheetBook, ctx):
    sheet = book.sheet_named("Sales")
    if sheet is None:
        return False, "no Sales sheet"
    header = [ref for ref, c in book.cells[sheet].items() if re.match(r"[A-Z]+1$", ref) and "growth" in str(c.raw).lower()]
    if not header:
        return False, "no Growth header in row 1"
    col = re.match(r"([A-Z]+)", header[0]).group(1)
    vals, problems = C.x_eval_all(book)
    want = [(140 - 120) / 120, (110 - 90) / 90, (65 - 70) / 70]
    got, formulas, fmts = [], [], []
    for row in (2, 3, 4):
        ref = f"{col}{row}"
        raw = book.raw(sheet, ref)
        formulas.append(raw)
        got.append(vals.get((sheet, ref)))
        fmts.append(book.cells[sheet].get(ref).number_format if book.cells[sheet].get(ref) else None)
    if not all(isinstance(f, str) and f.startswith("=") for f in formulas):
        return False, f"{col}2:{col}4 are not all formulas: {formulas}"
    if any(not isinstance(g, (int, float)) for g in got):
        return (None if problems else False), f"values {got}; {problems[:2]}"
    as_fraction = all(abs(g - w) < 0.001 for g, w in zip(got, want))
    as_number = all(abs(g - 100 * w) < 0.1 for g, w in zip(got, want))
    if as_fraction:
        pct = all("%" in (f or "") for f in fmts)
        return pct, f"{formulas}, formats {fmts}" + ("" if pct else ": fractions shown without a % format")
    if as_number:
        return True, f"{formulas} as percentage numbers"
    return False, f"{formulas} give {got}, want {[round(w, 4) for w in want]}"


def _east_total_fixed(book: SheetBook, ctx):
    sheet = book.sheet_named("Sales")
    raw = book.raw(sheet, "D3") if sheet else None
    vals, _ = C.x_eval_all(book)
    v = vals.get((sheet, "D3"))
    ok = isinstance(raw, str) and raw.startswith("=") and isinstance(v, (int, float)) and abs(v - 200) < 1e-6 and "C3" in raw.replace("$", "")
    return ok, f"D3 = {raw!r} -> {v}"


def _sheet_order(names: list[str]):
    def fn(book: SheetBook, ctx):
        return book.sheets == names, f"sheets {book.sheets}"
    return C.Check(f"sheets are {names}", fn, ("rename", "renamed"))


def _cell_is(sheet: str, ref: str, value):
    def fn(book: SheetBook, ctx):
        name = book.sheet_named(sheet)
        got = book.raw(name, ref) if name else None
        return got == value, f"{sheet}!{ref} = {got!r}"
    return C.Check(f"{sheet}!{ref} kept", fn)


BUDGET_AMOUNTS = [3200, 650, 1150, 420, 180, 140, 95, 500, 200]

SHEET_TASKS = [
    Task(
        id="xlsx-budget",
        kind="xlsx",
        mode="create",
        pilot=True,
        target="budget.xlsx",
        prompt=(
            "Make me a monthly household budget in budget.xlsx. Income: salary 3,200 and freelance 650. Expenses: rent "
            "1,150, groceries 420, utilities 180, transport 140, insurance 95, savings 500, fun 200. I want totals done "
            "with formulas, a line showing what's left at the end of the month, the money formatted as euros, and a pie "
            "chart of the expenses."
        ),
        checks=(
            C.x_numbers(BUDGET_AMOUNTS),
            C.x_formula_value(3850, "total income"),
            C.x_formula_value(2685, "total expenses"),
            C.x_formula_value(1165, "left over"),
            C.x_format(BUDGET_AMOUNTS + [3850, 2685, 1165], r'€|\[\$EUR|"EUR', "euros", suspect=r"\[Euro\]"),
            C.x_chart(types={"pie", "doughnut"}),
            C.x_cached_consistent(),
        ),
    ),
    Task(
        id="xlsx-sales-table",
        kind="xlsx",
        mode="create",
        target="sales-2026.xlsx",
        prompt=(
            "Put this sales data into sales-2026.xlsx - city, units sold, revenue: Lisbon 412, €48,200; Porto 530, "
            "€61,750; Madrid 351, €39,900; Barcelona 618, €72,300; Valencia 160, €18,450; Seville 233, €25,600. Sort it by "
            "revenue, highest first, make the header row bold, and show revenue with thousands separators."
        ),
        checks=(
            C.Check("six rows with their units and revenue as numbers", _sales_complete),
            C.Check("sorted by revenue, highest first", _sales_sorted, ("sorted",)),
            C.Check("header row bold", _header_bold, ("bold", "header")),
            C.x_format([v for _, v in SALES_ROWS.values()], r"#,##0|#,#", "thousands separators"),
        ),
    ),
    Task(
        id="xlsx-invoice",
        kind="xlsx",
        mode="create",
        target="invoice-1043.xlsx",
        prompt=(
            "Make an invoice in invoice-1043.xlsx for Acme Corp, invoice number 1043: website redesign, 1 at 4,800; "
            "hosting, 12 months at 35; support, 18 hours at 85. Work out each line total, the subtotal, VAT at 23% and the "
            "total due, all with formulas."
        ),
        checks=(
            _text_cells(["Acme", 1043]),
            C.x_formula_value(4800, "redesign line total"),
            C.x_formula_value(420, "hosting line total"),
            C.x_formula_value(1530, "support line total"),
            C.x_formula_value(6750, "subtotal"),
            C.x_formula_value(1552.5, "VAT"),
            C.x_formula_value(8302.5, "total due"),
            C.x_cached_consistent(),
        ),
    ),
    Task(
        id="xlsx-multi-sheet",
        kind="xlsx",
        mode="create",
        target="team-hours.xlsx",
        prompt=(
            "Set up team-hours.xlsx with one sheet per month - Jan, Feb, Mar - listing hours per person (Ana: 120, 132, "
            "128; Ben: 98, 110, 105; Carla: 140, 138, 150), and a Summary sheet that adds up each person's three months "
            "with formulas that point at the month sheets."
        ),
        checks=(
            C.x_sheets(["Jan", "Feb", "Mar", "Summary"]),
            _values_on("Jan", [120, 98, 140]),
            _values_on("Feb", [132, 110, 138]),
            _values_on("Mar", [128, 105, 150]),
            _formula_on("Summary", 380, "Ana's total", must_ref_other_sheet=True),
            _formula_on("Summary", 313, "Ben's total", must_ref_other_sheet=True),
            _formula_on("Summary", 428, "Carla's total", must_ref_other_sheet=True),
            C.x_cached_consistent(),
        ),
    ),
    Task(
        id="xlsx-line-chart",
        kind="xlsx",
        mode="create",
        target="web-traffic.xlsx",
        prompt=(
            "In a new web-traffic.xlsx, put our weekly visitors for the last 8 weeks - 1,200, 1,350, 1,410, 1,380, 1,620, "
            "1,750, 1,900, 2,050 - and add a line chart of it called 'Weekly visitors'."
        ),
        checks=(
            C.x_numbers([1200, 1350, 1410, 1380, 1620, 1750, 1900, 2050]),
            C.x_chart(types={"line"}, title="Weekly visitors"),
            _chart_covers(8),
        ),
    ),
    Task(
        id="xlsx-dates",
        kind="xlsx",
        mode="create",
        target="launch-plan.xlsx",
        prompt=(
            "Make a launch plan in launch-plan.xlsx with task, owner, start date, end date, and a column that works out "
            "how many days each task takes. Copy final - Inês - 3 to 7 Nov; Landing page - Rui - 5 to 14 Nov; Email "
            "campaign - Ana - 10 to 21 Nov; Press kit - Inês - 12 to 19 Nov; Launch event - Marta - 24 Nov. All in 2026."
        ),
        checks=(
            _text_cells(["Copy final", "Landing page", "Email campaign", "Press kit", "Launch event", "Inês", "Rui", "Ana", "Marta"]),
            _dates_real(9),
            C.Check("durations computed by formula", _durations, ("formula", "days")),
            C.x_cached_consistent(),
        ),
    ),
    Task(
        id="xlsx-add-column",
        kind="xlsx",
        mode="edit",
        inputs=(SALES_XLSX,),
        target=SALES_XLSX.name,
        prompt=(
            "In sales.xlsx, add a Growth column to the Sales sheet showing how much Q2 grew over Q1 for each region, as a "
            "percentage, with formulas."
        ),
        checks=(
            C.Check("a Growth column of formulas with the right percentages", _growth_column, ("growth", "formula")),
            C.x_keeps("Sales", [f"{c}{r}" for c in "ABCD" for r in range(1, 5)], "the existing table is kept"),
            C.x_keeps("Notes", name="the Notes sheet is kept"),
            C.x_cached_consistent(),
        ),
    ),
    Task(
        id="xlsx-fix-formula",
        kind="xlsx",
        mode="edit",
        inputs=(Input("sales.xlsx", "sample.xlsx", derive=_derive_wrong_east_total),),
        target="sales.xlsx",
        prompt="The East total in sales.xlsx looks wrong to me - can you fix it?",
        notes="Input: sample.xlsx with D3 changed to =B3+B2 (210 instead of 200).",
        checks=(
            C.Check("East's total is =B3+C3 (200)", _east_total_fixed, ("fixed", "200")),
            C.x_keeps("Sales", [f"{c}{r}" for c in "ABCD" for r in range(1, 5) if f"{c}{r}" != "D3"], "the other cells are kept"),
            C.x_keeps("Notes", name="the Notes sheet is kept"),
        ),
    ),
    Task(
        id="xlsx-rename-sheet",
        kind="xlsx",
        mode="edit",
        inputs=(SALES_XLSX,),
        target=SALES_XLSX.name,
        prompt="Rename the Notes sheet in sales.xlsx to 'Read me'.",
        checks=(
            _sheet_order(["Sales", "Read me"]),
            _cell_is("Read me", "A1", "Figures in thousands of euros."),
            C.x_keeps("Sales", name="the Sales sheet is kept"),
        ),
    ),
]


# =========================================================================== Deck

def _bullet_added(deck: Deck, ctx):
    if not deck.slides:
        return False, "no slides"
    first = deck.slides[0]
    for s in first.shapes:
        lines = [norm(t) for t, _ in s.paragraphs]
        if any("4% under budget" in ln for ln in lines):
            if "revenue met target" in lines and "hiring slipped" in lines:
                order = [ln for ln in lines if ln]
                return order[:2] == ["revenue met target", "hiring slipped"], f"bullets: {[t for t, _ in s.paragraphs]}"
            return False, f"added in a separate box {s.name!r}, not with the other bullets"
    return False, f"not on slide 1: {first.text()[:200]!r}"


def _title_kept(index: int, title: str):
    def fn(deck: Deck, ctx):
        if index > len(deck.slides):
            return False, f"no slide {index}"
        got = deck.slides[index - 1].title
        return norm(got) == norm(title), f"title {got!r}"
    return fn


def _deck_table(containing: list, slide: int | None = None):
    def fn(deck: Deck, ctx):
        slides = deck.slides if slide is None else deck.slides[slide - 1:slide]
        tables = [s for sl in slides for s in sl.shapes if s.kind == "table"]
        for t in tables:
            text = "\n".join(" | ".join(r) for r in t.table)
            missing = [c for c in containing if not C.has_text(text, c)]
            if not missing:
                return True, f"native table {len(t.table)}x{max(len(r) for r in t.table)}"
        if tables:
            return False, f"tables lack some of {containing}: {[t.table for t in tables][:2]}"
        kinds = sorted({s.kind for sl in slides for s in sl.shapes})
        return False, f"no native table (shapes: {kinds})"
    where = f" on slide {slide}" if slide else ""
    return C.Check(f"a native table{where} holding {containing}", fn, ("table",))


def _slide_has(index: int, items: list, label: str = ""):
    def fn(deck: Deck, ctx):
        if index > len(deck.slides):
            return False, f"no slide {index}"
        text = deck.slides[index - 1].text()
        missing = [i for i in items if not C.has_text(text, i)]
        return not missing, "all there" if not missing else f"missing: {missing}"
    return C.Check(label or f"slide {index} holds {items}", fn)


def _input_unchanged(name: str):
    def fn(doc, ctx):
        return name not in ctx.inputs_changed, "unchanged" if name not in ctx.inputs_changed else f"{name} was modified"
    return C.Check(f"{name} left as it was", fn)


def _chart_on_slide(index: int):
    def fn(deck: Deck, ctx):
        if index > len(deck.slides):
            return False, f"no slide {index}"
        shapes = deck.slides[index - 1].shapes
        kinds = [s.kind for s in shapes]
        if "chart" in kinds:
            return True, "native chart"
        if "picture" in kinds:
            return True, "a picture (inline SVG drawn as a picture)"
        drawn = kinds.count("other") + sum(1 for s in shapes if s.kind == "text" and not s.text)
        if drawn >= 6:
            return True, f"drawn from {drawn} shapes"
        if "table" in kinds:
            return False, "a table, not a chart"
        return False, f"no chart (shapes: {kinds})"
    return C.Check(f"a chart on slide {index}", fn, ("chart",))


def _chart_values(index: int, values: list):
    def fn(deck: Deck, ctx):
        if index > len(deck.slides):
            return False, f"no slide {index}"
        sl = deck.slides[index - 1]
        kinds = [s.kind for s in sl.shapes]
        text = sl.text()
        hits = [v for v in values if C.has_text(text, v)]
        if len(hits) >= len(values) - 1:
            return True, f"{len(hits)}/{len(values)} values labelled"
        if "picture" in kinds or "chart" in kinds:
            return None, f"chart is a {'picture' if 'picture' in kinds else 'native chart'}; its values are not readable here"
        return False, f"only {hits} of {values} on the slide"
    return C.Check("the chart carries the six monthly values", fn)


def _title_text_changed(index: int, title: str):
    def fn(deck: Deck, ctx):
        if index > len(deck.slides):
            return False, f"no slide {index}"
        got = deck.slides[index - 1].title
        return norm(got) == norm(title), f"title {got!r}"
    return C.Check(f"slide {index} title is {title!r}", fn, ("title",))


def _slide_kept_but(index: int, skip_title: bool = False, skip_notes: bool = False, name: str = ""):
    """Slide *index* keeps every shape (text, place, size, font sizes) except its title and/or notes."""
    def fn(deck: Deck, ctx):
        if ctx.original is None:
            return None, "no original"
        if index > len(deck.slides):
            return False, f"no slide {index}"
        a, b = ctx.original.slides[index - 1], deck.slides[index - 1]
        sa = [x for s, x in zip(a.shapes, C.p_sig(a)) if not (skip_title and s.placeholder in ("title", "ctrTitle"))]
        sb = [x for s, x in zip(b.shapes, C.p_sig(b)) if not (skip_title and s.placeholder in ("title", "ctrTitle"))]
        if skip_title and len(sb) == len(sa) + 1:
            # The title may have been rebuilt as a plain text box: drop the shape holding the new title.
            sb = [x for x in sb if norm(x[1]) != norm(b.title)]
        if len(sa) != len(sb):
            return False, f"{len(sa)} other shapes -> {len(sb)}"
        for x, y in zip(sa, sb):
            if x[:3] != y[:3]:
                return False, f"{x[1][:40]!r} -> {y[1][:40]!r}"
            if any(abs(p - q) > 2 for p, q in zip(x[3:7], y[3:7])):
                return False, f"{x[1][:30]!r} moved/resized {x[3:7]} -> {y[3:7]}"
            if x[7] != y[7]:
                return False, f"{x[1][:30]!r} font sizes {x[7]} -> {y[7]}"
        if not skip_notes and norm(a.notes) != norm(b.notes):
            return False, f"notes {a.notes!r} -> {b.notes!r}"
        return True, "kept"
    return C.Check(name or f"the rest of slide {index} is kept", fn)


def _titles(order: list[str]):
    def fn(deck: Deck, ctx):
        got = [s.title for s in deck.slides]
        return [norm(t) for t in got] == [norm(t) for t in order], f"order {got}"
    return C.Check(f"slides in the order {order}", fn, ("moved", "first", "order"))


def _notes_contain(index: int, phrases: list[str]):
    def fn(deck: Deck, ctx):
        if index > len(deck.slides):
            return False, f"no slide {index}"
        notes = deck.slides[index - 1].notes
        missing = [p for p in phrases if norm(p) not in norm(notes)]
        return not missing, f"notes {notes[:120]!r}"
    return C.Check(f"slide {index} speaker notes say it", fn, ("notes",))


DECK_TASKS = [
    Task(
        id="pptx-pitch",
        kind="pptx",
        mode="create",
        pilot=True,
        target="shelfie-pitch.pptx",
        prompt=(
            "I need a 6-slide pitch deck, shelfie-pitch.pptx, for Shelfie - an app that gets a small bookshop selling its "
            "stock online in 10 minutes. Problem: there are 2,300 independent bookshops in Portugal and Spain and 70% of "
            "them don't sell online. Then the solution, how it works, the business model (€29 a month plus 2% per sale), "
            "traction (85 shops in the beta, €41k of sales in 3 months), and the ask: we're raising €600k for 18 months."
        ),
        checks=(
            C.p_slides(5, 7),
            C.contains_all([2300, "70%", 29, "2%", 85, 41, 600, 18]),
            C.p_every_slide_has_text(),
            C.p_no_overflow(),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-add-bullet",
        kind="pptx",
        mode="edit",
        pilot=True,
        inputs=(REVIEW_PPTX,),
        target=REVIEW_PPTX.name,
        prompt="Add 'Costs came in 4% under budget' as another bullet on the first slide of quarterly-review.pptx.",
        checks=(
            C.Check("the bullet joins the slide's bullet list", _bullet_added, ("bullet",)),
            C.Check("slide 1 title kept", _title_kept(1, "Quarterly review")),
            C.p_slide_kept(2),
            C.p_slides(2, 2),
            C.p_no_overflow(),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-report-summary",
        kind="pptx",
        mode="create",
        inputs=(REPORT_DOCX,),
        target="report-summary.pptx",
        prompt="Turn quarterly-report.docx into a 3-slide summary deck, report-summary.pptx.",
        checks=(
            C.p_slides(3, 3),
            C.contains_all(["12%", "north", "lisbon", ("two engineers", "hire", "hiring"), ("slip",)]),
            C.p_every_slide_has_text(),
            C.p_no_overflow(),
            C.p_in_bounds(),
            _input_unchanged("quarterly-report.docx"),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-onboarding",
        kind="pptx",
        mode="create",
        target="support-onboarding.pptx",
        prompt=(
            "Make a 5-slide onboarding deck for new support staff, support-onboarding.pptx: a welcome slide; our three "
            "support tiers with their response targets (Tier 1: 1 hour, Tier 2: 4 hours, Tier 3: 24 hours) as a table; "
            "the tools we use (Zendesk, Slack, Notion); what the first week looks like; and who to ask (Marta, team lead; "
            "Rui, tooling)."
        ),
        checks=(
            C.p_slides(5, 5),
            _deck_table(["tier", 1, 4, 24]),
            C.contains_all(["Zendesk", "Slack", "Notion", "Marta", "Rui"]),
            C.p_every_slide_has_text(),
            C.p_no_overflow(),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-kpi-chart",
        kind="pptx",
        mode="create",
        target="kpi.pptx",
        prompt=(
            "Create kpi.pptx: one slide titled 'Revenue is up 50% since January' with a chart of monthly revenue from Jan "
            "to Jun: 42, 45, 51, 49, 58, 63 (thousand euros)."
        ),
        checks=(
            C.p_slides(1, 1),
            _slide_has(1, ["revenue is up 50%"], "the title says it"),
            _chart_on_slide(1),
            _chart_values(1, [42, 45, 51, 49, 58, 63]),
            C.p_no_overflow(),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-workshop",
        kind="pptx",
        mode="create",
        target="workshop.pptx",
        prompt=(
            "Make workshop.pptx with a title slide, 'Pricing workshop - 12 November', and an agenda slide: 09:00 Welcome; "
            "09:15 Where our prices are today; 10:00 Competitor review; 10:45 Break; 11:00 Three pricing options; 12:00 "
            "Vote and next steps; 12:30 Close."
        ),
        checks=(
            C.p_slides(2, 2),
            _slide_has(1, ["pricing workshop", "12 november"], "title slide"),
            _slide_has(2, ["09:00", "09:15", "10:00", "10:45", "11:00", "12:00", "12:30", "welcome",
                           "where our prices are today", "competitor review", "break", "three pricing options",
                           "vote and next steps", "close"], "agenda slide has every item and time"),
            C.p_no_overflow(),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-change-title",
        kind="pptx",
        mode="edit",
        inputs=(REVIEW_PPTX,),
        target=REVIEW_PPTX.name,
        prompt="Change the title of the second slide in quarterly-review.pptx to 'Q3 plan'.",
        checks=(
            _title_text_changed(2, "Q3 plan"),
            _slide_kept_but(2, skip_title=True),
            C.p_slide_kept(1),
            C.p_slides(2, 2),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-reorder",
        kind="pptx",
        mode="edit",
        inputs=(REVIEW_PPTX,),
        target=REVIEW_PPTX.name,
        prompt="In quarterly-review.pptx, move the 'Next quarter' slide to the front.",
        checks=(
            _titles(["Next quarter", "Quarterly review"]),
            C.p_slide_kept(2, new_index=1, name="the moved slide is intact (notes too)"),
            C.p_slide_kept(1, new_index=2, name="the other slide is intact"),
            C.p_slides(2, 2),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-table-slide",
        kind="pptx",
        mode="edit",
        inputs=(REVIEW_PPTX,),
        target=REVIEW_PPTX.name,
        prompt=(
            "Add a slide at the end of quarterly-review.pptx with a table of Q1 and Q2 revenue by region: North 120 and "
            "140, East 90 and 110, South 70 and 65."
        ),
        checks=(
            C.p_slides(3, 3),
            _deck_table(["North", "East", "South", 120, 140, 90, 110, 70, 65], slide=3),
            C.p_slide_kept(1),
            C.p_slide_kept(2),
            C.p_no_overflow(),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
    Task(
        id="pptx-notes",
        kind="pptx",
        mode="edit",
        inputs=(REVIEW_PPTX,),
        target=REVIEW_PPTX.name,
        prompt=(
            "Add speaker notes to the first slide of quarterly-review.pptx: 'Thank the regional leads. Hiring slipped "
            "because two offers were declined.'"
        ),
        checks=(
            _notes_contain(1, ["Thank the regional leads.", "Hiring slipped because two offers were declined."]),
            _slide_kept_but(1, skip_notes=True, name="slide 1's visible content is kept"),
            C.p_slide_kept(2),
            C.p_slides(2, 2),
            C.p_in_bounds(),
            C.p_contrast(),
        ),
    ),
]


# =========================================================================== PDF

def _field(name: str, want: str):
    def fn(pdf: PdfDoc, ctx):
        got = pdf.fields.get(name)
        return (got or "").strip() == want, f"{name} = {got!r}"
    return fn


def _checked(name: str):
    def fn(pdf: PdfDoc, ctx):
        got = pdf.fields.get(name)
        return got not in (None, "", "/Off", "Off"), f"{name} = {got!r}"
    return fn


def _note_near(phrase: str, words_in_note: str, within_pt: float = 60.0):
    def fn(pdf: PdfDoc, ctx):
        where = pdf.y_of(phrase)
        notes = [a for a in pdf.annotations if norm(words_in_note) in norm(a.contents)]
        if not notes:
            other = [(a.page, a.subtype, a.contents[:40]) for a in pdf.annotations]
            return False, f"no annotation saying {words_in_note!r} (annotations: {other})"
        if where is None:
            return None, "the phrase is no longer in the text"
        page, y = where
        for a in notes:
            if a.page == page and a.rect:
                y1, y2 = sorted((a.rect[1], a.rect[3]))
                gap = 0.0 if y1 <= y <= y2 else min(abs(y - y1), abs(y - y2))
                if gap <= within_pt:
                    return True, f"{a.subtype} on page {a.page}, {gap:.0f} pt from the line"
        a = notes[0]
        return False, f"{a.subtype} on page {a.page} at {a.rect}, the line is at y={y:.0f} on page {page}"
    return C.Check(f"a note saying {words_in_note!r} by {phrase!r}", fn, ("note", "annotation"))


def _text_lacks(phrase: str):
    def fn(pdf: PdfDoc, ctx):
        gone = norm(phrase) not in norm(pdf.text())
        return gone, "no longer in the text" if gone else "still extractable from the text"
    return C.Check(f"{phrase!r} is gone from the text", fn, ("redact",))


PDF_TASKS = [
    Task(
        id="pdf-fill",
        kind="pdf",
        mode="edit",
        pilot=True,
        inputs=(AGREEMENT_PDF,),
        target=AGREEMENT_PDF.name,
        prompt="Please fill in service-agreement.pdf: the client is Northwind Lda, and tick the Signed box.",
        checks=(
            C.Check("client field says Northwind Lda", _field("client", "Northwind Lda"), ("northwind", "client")),
            C.Check("signed box ticked", _checked("signed"), ("tick", "signed", "checked")),
            C.f_pages(2),
            C.f_text_keeps(SAMPLE_PDF_TEXT),
            C.f_fields_kept(["client", "signed"]),
        ),
    ),
    Task(
        id="pdf-annotate",
        kind="pdf",
        mode="edit",
        inputs=(AGREEMENT_PDF,),
        target=AGREEMENT_PDF.name,
        prompt=(
            "Put a note next to the 'thirty days' line in service-agreement.pdf saying: Finance wants 45 days, check "
            "with legal."
        ),
        checks=(
            _note_near("thirty days", "45 days"),
            C.f_pages(2),
            C.f_text_keeps(SAMPLE_PDF_TEXT),
            C.f_fields_kept(["client", "signed"]),
        ),
    ),
    Task(
        id="pdf-redact",
        kind="pdf",
        mode="edit",
        inputs=(AGREEMENT_PDF,),
        target=AGREEMENT_PDF.name,
        prompt="Before I send service-agreement.pdf out, redact the start date - 'from 1 March 2026'.",
        checks=(
            _text_lacks("1 March 2026"),
            C.f_text_keeps(["Service agreement", "The supplier delivers the service",
                            "Fees are due within thirty days of each invoice.", "Page two: the schedule of services."]),
            C.f_pages(2),
            C.f_fields_kept(["client", "signed"]),
        ),
    ),
]


TASKS: list[Task] = WORD_TASKS + SHEET_TASKS + DECK_TASKS + PDF_TASKS


def by_id() -> dict[str, Task]:
    return {t.id: t for t in TASKS}
