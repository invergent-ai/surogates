"""Read a saved document into a plain model the checks can ask questions of.

One reader per type, each tolerant of what an editor writes rather than of
what python-docx/openpyxl/python-pptx would write: styles found by name or
by outline level, lists by their numbering definition, charts by the parts
in the package. A reader never judges; the checks in ``tasks.py`` do.
"""
from __future__ import annotations

import math
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}
W = "{%s}" % NS["w"]
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]

KINDS = {".docx": "docx", ".xlsx": "xlsx", ".pptx": "pptx", ".pdf": "pdf"}


def kind_of(path: str) -> str | None:
    return KINDS.get(posixpath.splitext(path.lower())[1])


def norm(text: str) -> str:
    """Whitespace collapsed, typographic quotes and dashes made plain, lowercased."""
    text = (text or "").replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    text = text.replace("–", "-").replace("—", "-").replace(" ", " ").replace(" ", " ")
    return re.sub(r"\s+", " ", text).strip().lower()


def words(text: str) -> int:
    return len(re.findall(r"\w+(?:[-'.,]\w+)*", text or ""))


# --------------------------------------------------------------------------- Word


@dataclass
class Run:
    text: str
    bold: bool


@dataclass
class Para:
    text: str
    style: str
    heading: int | None  # 0 for Title, 1..9 for headings, None otherwise
    list_kind: str | None  # "bullet" | "number" | None
    num_id: str | None
    level: int
    runs: list[Run] = field(default_factory=list)
    number: int | None = None  # the number a reader sees on a numbered item


@dataclass
class Table:
    rows: list[list[str]]
    style: str


@dataclass
class WordDoc:
    blocks: list[Para | Table]
    charts: int
    pictures: int
    pending: list[str]
    chart_xml: list[str] = field(default_factory=list)
    num_abstract: dict[str, str] = field(default_factory=dict)  # numId -> abstractNumId
    num_restarts: dict[str, int] = field(default_factory=dict)  # numId -> the value its startOverride restarts at

    @property
    def paragraphs(self) -> list[Para]:
        return [b for b in self.blocks if isinstance(b, Para)]

    @property
    def tables(self) -> list[Table]:
        return [b for b in self.blocks if isinstance(b, Table)]

    @property
    def headings(self) -> list[Para]:
        return [p for p in self.paragraphs if p.heading is not None]

    def text(self) -> str:
        parts: list[str] = []
        for b in self.blocks:
            if isinstance(b, Para):
                parts.append(b.text)
            else:
                parts.extend(" | ".join(r) for r in b.rows)
        return "\n".join(parts)

    def body_words(self) -> int:
        return words(self.text())


def read_docx(path: str) -> WordDoc:
    import docx  # python-docx

    document = docx.Document(path)
    numbering, num_abstract, num_restarts, starts = _numbering_formats(document)
    style_info = _word_styles(document)

    blocks: list[Para | Table] = []
    pending: list[str] = []

    def para_of(p_el) -> Para:
        from docx.text.paragraph import Paragraph

        para = Paragraph(p_el, document)
        style_id = None
        ppr = p_el.find(W + "pPr")
        if ppr is not None and ppr.find(W + "pStyle") is not None:
            style_id = ppr.find(W + "pStyle").get(W + "val")
        style_id = style_id or style_info["default"]
        name, outline, style_num, style_bold = style_info["styles"].get(style_id, (style_id or "", None, None, False))
        heading = _heading_level(name, outline)
        if ppr is not None and ppr.find(W + "outlineLvl") is not None and heading is None:
            lvl = int(ppr.find(W + "outlineLvl").get(W + "val") or 9)
            heading = lvl + 1 if lvl < 9 else None
        num_id, level = None, 0
        numpr = ppr.find(W + "numPr") if ppr is not None else None
        if numpr is not None:
            nid = numpr.find(W + "numId")
            ilvl = numpr.find(W + "ilvl")
            num_id = nid.get(W + "val") if nid is not None else None
            level = int(ilvl.get(W + "val")) if ilvl is not None else 0
        elif style_num:
            num_id = style_num
        list_kind = None
        if num_id and num_id != "0":
            fmt = numbering.get((num_id, level))
            list_kind = "bullet" if fmt == "bullet" else ("number" if fmt else None)
            if list_kind is None:
                lname = name.lower()
                list_kind = "bullet" if "bullet" in lname else ("number" if "number" in lname else "number")
        runs = []
        for r in para.runs:
            bold = r.bold
            if bold is None and r.style is not None and r.style.font is not None:
                bold = r.style.font.bold
            if bold is None:
                bold = style_bold
            runs.append(Run(r.text, bool(bold)))
        return Para(para.text, name, heading, list_kind, num_id if list_kind else None, level, runs)

    def walk(container) -> None:
        for child in container:
            tag = child.tag
            if tag == W + "p":
                blocks.append(para_of(child))
            elif tag == W + "tbl":
                rows = []
                for tr in child.findall(W + "tr"):
                    cells = []
                    for tc in tr.findall(W + "tc"):
                        cells.append(" ".join(
                            "".join(t.text or "" for t in p.iter(W + "t")) for p in tc.findall(W + "p")
                        ).strip())
                    rows.append(cells)
                tstyle = child.find(f"{W}tblPr/{W}tblStyle")
                sid = tstyle.get(W + "val") if tstyle is not None else ""
                blocks.append(Table(rows, style_info["styles"].get(sid, (sid,))[0] if sid else ""))
            elif tag == W + "sdt":
                tag_el = child.find(f"{W}sdtPr/{W}tag")
                val = tag_el.get(W + "val") if tag_el is not None else ""
                if val and val.startswith("pending:"):
                    pending.append(val[len("pending:"):])
                    continue
                content = child.find(W + "sdtContent")
                if content is not None:
                    walk(content)

    walk(document.element.body)
    _number_items(blocks, num_abstract, num_restarts, starts)

    charts, pictures, chart_xml = 0, 0, []
    for part in document.part.package.iter_parts():
        name = str(part.partname)
        if name.startswith("/word/charts/") and name.endswith(".xml") and "/_rels/" not in name and "colors" not in name and "style" not in name:
            charts += 1
            chart_xml.append(part.blob.decode("utf-8", "replace"))
        elif name.startswith("/word/media/"):
            pictures += 1
    return WordDoc(blocks=blocks, charts=charts, pictures=pictures, pending=pending, chart_xml=chart_xml,
                   num_abstract=num_abstract, num_restarts=num_restarts)


def _heading_level(name: str, outline: int | None) -> int | None:
    lname = (name or "").lower()
    if lname == "title":
        return 0
    m = re.match(r"heading\s*(\d)", lname)
    if m:
        return int(m.group(1))
    if outline is not None and outline < 9:
        return outline + 1
    return None


def _word_styles(document) -> dict[str, Any]:
    """style id -> (name, outline level, numId, bold), following basedOn."""
    root = document.styles.element
    raw: dict[str, dict[str, Any]] = {}
    default = None
    for st in root.findall(W + "style"):
        if st.get(W + "type") != "paragraph" and st.get(W + "type") != "table":
            continue
        sid = st.get(W + "styleId")
        name_el = st.find(W + "name")
        based = st.find(W + "basedOn")
        outline = st.find(f"{W}pPr/{W}outlineLvl")
        num = st.find(f"{W}pPr/{W}numPr/{W}numId")
        bold = st.find(f"{W}rPr/{W}b")
        raw[sid] = {
            "name": name_el.get(W + "val") if name_el is not None else sid,
            "based": based.get(W + "val") if based is not None else None,
            "outline": int(outline.get(W + "val")) if outline is not None and outline.get(W + "val", "").isdigit() else None,
            "num": num.get(W + "val") if num is not None else None,
            "bold": (bold.get(W + "val") not in ("0", "false")) if bold is not None else None,
        }
        if st.get(W + "default") in ("1", "true") and st.get(W + "type") == "paragraph":
            default = sid

    def resolve(sid: str, key: str, depth: int = 0):
        st = raw.get(sid)
        if st is None or depth > 10:
            return None
        if st[key] is not None:
            return st[key]
        return resolve(st["based"], key, depth + 1) if st["based"] else None

    styles = {}
    for sid, st in raw.items():
        name = st["name"]
        # python-docx's display names for the built-in styles.
        if re.fullmatch(r"heading \d", name.lower()):
            name = "Heading " + name.split()[-1]
        styles[sid] = (name, resolve(sid, "outline"), resolve(sid, "num"), bool(resolve(sid, "bold")))
    return {"styles": styles, "default": default}


def _number_items(blocks: list, num_abstract: dict[str, str], restarts: dict[str, int],
                  starts: dict[tuple[str, int], int]) -> None:
    """Give each numbered item the number Word shows: one counter per abstract numbering definition and level (lists
    sharing a definition continue each other), restarted where a numbering instance has a startOverride, and deeper
    levels reset when a shallower one advances."""
    counters: dict[tuple[str, int], int] = {}
    started: set[str] = set()
    for b in blocks:
        if not isinstance(b, Para) or b.list_kind != "number" or not b.num_id:
            continue
        aid = num_abstract.get(b.num_id, b.num_id)
        key = (aid, b.level)
        if b.num_id in restarts and b.num_id not in started:
            counters[key] = restarts[b.num_id] - 1
        started.add(b.num_id)
        counters[key] = counters.get(key, starts.get(key, 1) - 1) + 1
        for (a, lvl) in list(counters):
            if a == aid and lvl > b.level:
                del counters[(a, lvl)]
        b.number = counters[key]


def _numbering_formats(document):
    """(numId, level) -> numFmt (bullet, decimal, ...); numId -> abstractNumId; numId -> its startOverride value;
    (abstractNumId, level) -> start value."""
    try:
        part = document.part.numbering_part
    except Exception:  # noqa: BLE001 - no numbering part at all
        return {}, {}, {}, {}
    root = part.element
    abstract: dict[str, dict[int, str]] = {}
    starts: dict[tuple[str, int], int] = {}
    for an in root.findall(W + "abstractNum"):
        levels = {}
        for lvl in an.findall(W + "lvl"):
            fmt = lvl.find(W + "numFmt")
            levels[int(lvl.get(W + "ilvl"))] = fmt.get(W + "val") if fmt is not None else "decimal"
            start = lvl.find(W + "start")
            if start is not None and (start.get(W + "val") or "").isdigit():
                starts[(an.get(W + "abstractNumId"), int(lvl.get(W + "ilvl")))] = int(start.get(W + "val"))
        abstract[an.get(W + "abstractNumId")] = levels
    out: dict[tuple[str, int], str] = {}
    to_abstract: dict[str, str] = {}
    restarts: dict[str, int] = {}
    for num in root.findall(W + "num"):
        nid = num.get(W + "numId")
        aid = num.find(W + "abstractNumId")
        to_abstract[nid] = aid.get(W + "val") if aid is not None else ""
        override = num.find(f"{W}lvlOverride/{W}startOverride")
        if override is not None and (override.get(W + "val") or "").isdigit():
            restarts[nid] = int(override.get(W + "val"))
        levels = abstract.get(aid.get(W + "val") if aid is not None else "", {})
        for lvl, fmt in levels.items():
            out[(nid, lvl)] = fmt
        for ov in num.findall(W + "lvlOverride"):
            lvl_el = ov.find(W + "lvl")
            if lvl_el is not None and lvl_el.find(W + "numFmt") is not None:
                out[(nid, int(ov.get(W + "ilvl")))] = lvl_el.find(W + "numFmt").get(W + "val")
    return out, to_abstract, restarts, starts


# --------------------------------------------------------------------------- Spreadsheet


@dataclass
class Chart:
    sheet: str
    types: list[str]  # pie, doughnut, line, column, bar, area, scatter, ...
    title: str
    refs: list[str]  # the c:f formulas the series name
    anchor: tuple[int, int] | None  # (col, row), 1-based, of the top-left corner


@dataclass
class Cell:
    raw: Any  # the formula ("=...") or the value as stored
    cached: Any  # the stored computed value of a formula (None when the file stores none)
    number_format: str
    bold: bool


@dataclass
class SheetBook:
    sheets: list[str]
    cells: dict[str, dict[str, Cell]]  # sheet -> "A1" -> Cell
    charts: list[Chart]

    def raw(self, sheet: str, ref: str) -> Any:
        cell = self.cells.get(sheet, {}).get(ref.upper())
        return cell.raw if cell else None

    def sheet_named(self, name: str) -> str | None:
        for s in self.sheets:
            if s.lower() == name.lower():
                return s
        return None

    def evaluator(self):
        from obench.formula import Evaluator

        return Evaluator(lambda sheet, ref: self.raw(sheet, ref))

    def formula_cells(self) -> list[tuple[str, str, str]]:
        return [
            (s, ref, c.raw) for s in self.sheets for ref, c in self.cells.get(s, {}).items()
            if isinstance(c.raw, str) and c.raw.startswith("=")
        ]


def read_xlsx(path: str) -> SheetBook:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    wb_vals = openpyxl.load_workbook(path, data_only=True)
    cells: dict[str, dict[str, Cell]] = {}
    for ws in wb.worksheets:
        vs = wb_vals[ws.title]
        sheet_cells: dict[str, Cell] = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                raw = c.value
                if hasattr(raw, "text"):  # ArrayFormula and friends
                    raw = getattr(raw, "text", str(raw))
                cached = vs[c.coordinate].value if isinstance(raw, str) and raw.startswith("=") else None
                sheet_cells[c.coordinate] = Cell(
                    raw=raw, cached=cached, number_format=c.number_format or "General",
                    bold=bool(c.font and c.font.b),
                )
        cells[ws.title] = sheet_cells
    return SheetBook(sheets=[ws.title for ws in wb.worksheets], cells=cells, charts=_xlsx_charts(path))


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, str]:
    """Relationship id -> target part path, for *part*."""
    base, name = posixpath.split(part)
    rels_path = posixpath.join(base, "_rels", name + ".rels")
    if rels_path not in z.namelist():
        return {}
    out = {}
    for rel in ET.fromstring(z.read(rels_path)).findall("rel:Relationship", NS):
        target = rel.get("Target")
        if rel.get("TargetMode") == "External":
            continue
        out[rel.get("Id")] = posixpath.normpath(
            target.lstrip("/") if target.startswith("/") else posixpath.join(base, target)
        )
    return out


_CHART_TAGS = {
    "pieChart": "pie", "pie3DChart": "pie", "doughnutChart": "doughnut", "lineChart": "line",
    "line3DChart": "line", "areaChart": "area", "area3DChart": "area", "scatterChart": "scatter",
    "radarChart": "radar", "bubbleChart": "bubble", "stockChart": "stock",
}


def chart_info(xml: bytes | str) -> tuple[list[str], str, list[str]]:
    """(types, title, series refs) of one chart part."""
    root = ET.fromstring(xml if isinstance(xml, bytes) else xml.encode())
    types: list[str] = []
    plot = root.find(".//c:plotArea", NS)
    if plot is not None:
        for child in plot:
            local = child.tag.split("}")[-1]
            if local in ("barChart", "bar3DChart"):
                d = child.find("c:barDir", NS)
                types.append("bar" if d is not None and d.get("val") == "bar" else "column")
            elif local in _CHART_TAGS:
                types.append(_CHART_TAGS[local])
    title_el = root.find(".//c:title", NS)
    title = ""
    if title_el is not None:
        title = "".join(t.text or "" for t in title_el.iter("{%s}t" % NS["a"]))
        if not title:
            f = title_el.find(".//c:f", NS)
            title = f.text if f is not None and f.text else ""
    refs = [f.text for f in root.iter("{%s}f" % NS["c"]) if f.text]
    return types, title, refs


def _xlsx_charts(path: str) -> list[Chart]:
    charts: list[Chart] = []
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        if "xl/workbook.xml" not in names:
            return charts
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        wb_rels = _rels(z, "xl/workbook.xml")
        for sheet in wb.findall("s:sheets/s:sheet", NS):
            sheet_part = wb_rels.get(sheet.get("{%s}id" % NS["r"]))
            if not sheet_part or sheet_part not in names:
                continue
            for drawing in _rels(z, sheet_part).values():
                if "drawings/" not in drawing or drawing not in names:
                    continue
                drels = _rels(z, drawing)
                droot = ET.fromstring(z.read(drawing))
                for anchor in list(droot):
                    chart_ref = anchor.find(".//c:chart", NS)
                    if chart_ref is None:
                        continue
                    target = drels.get(chart_ref.get("{%s}id" % NS["r"]))
                    if not target or target not in names:
                        continue
                    frm = anchor.find("xdr:from", NS)
                    pos = None
                    if frm is not None:
                        pos = (int(frm.find("xdr:col", NS).text) + 1, int(frm.find("xdr:row", NS).text) + 1)
                    types, title, refs = chart_info(z.read(target))
                    charts.append(Chart(sheet.get("name"), types, title, refs, pos))
    return charts


# --------------------------------------------------------------------------- Deck

EMU_PER_PT = 12700


@dataclass
class Shape:
    kind: str  # text | table | picture | chart | group | other
    name: str
    x: float  # points
    y: float
    w: float
    h: float
    text: str
    paragraphs: list[tuple[str, float | None]]  # (text, font size pt or None)
    placeholder: str | None  # title, body, ... for placeholders
    table: list[list[str]] | None = None
    in_group: bool = False  # a group's child: its offsets are in the group's own space, not the slide's
    group_box: tuple[float, float, float, float] | None = None  # a group child's outermost group's slide box
    overflow: str | None = None  # measured on the render: "" fits, else how far the text runs; None: not measured
    color: str | None = None  # the first run's explicit sRGB colour ("RRGGBB"); None for a theme colour or none set
    fill: str | None = None  # a solid sRGB fill; None for no fill or one whose colour is not read here
    opaque: bool = False  # paints over its box: any fill, a picture, a table or a chart


@dataclass
class Slide:
    shapes: list[Shape]
    notes: str
    title: str
    background: str | None = None  # the solid sRGB background the slide shows; None for a theme or picture one

    def text(self) -> str:
        parts = []
        for s in self.shapes:
            parts.append(s.text)
            if s.table:
                parts.extend(" | ".join(r) for r in s.table)
        return "\n".join(p for p in parts if p)


@dataclass
class Deck:
    width: float  # points
    height: float
    slides: list[Slide]
    render_note: str = "the deck was not rendered"  # why no text box's overflow was measured, when none was

    def text(self) -> str:
        return "\n\n".join(s.text() for s in self.slides)


def read_pptx(path: str) -> Deck:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(path)
    width = (deck.slide_width or 0) / EMU_PER_PT
    height = (deck.slide_height or 0) / EMU_PER_PT
    slides: list[Slide] = []
    for slide in deck.slides:
        shapes: list[Shape] = []

        def add(shape, dx: float = 0.0, dy: float = 0.0, in_group: bool = False,
                group_box: tuple[float, float, float, float] | None = None) -> None:
            x = ((shape.left or 0) / EMU_PER_PT) + dx
            y = ((shape.top or 0) / EMU_PER_PT) + dy
            w = (shape.width or 0) / EMU_PER_PT
            h = (shape.height or 0) / EMU_PER_PT
            ph = None
            if shape.is_placeholder:
                try:
                    ph = str(shape.placeholder_format.type).split(".")[-1].split(" ")[0].lower()
                except Exception:  # noqa: BLE001
                    ph = "other"
                ph = {"center_title": "ctrTitle", "subtitle": "subTitle"}.get(ph, ph)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for child in shape.shapes:
                    add(child, 0.0, 0.0, True, group_box or (x, y, w, h))
                return
            fill, opaque = _paint(shape._element)
            paint = {"in_group": in_group, "group_box": group_box, "fill": fill, "opaque": opaque}
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
                shapes.append(Shape("table", shape.name, x, y, w, h, "", [], ph, table=rows, **paint))
                return
            if getattr(shape, "has_chart", False) and shape.has_chart:
                shapes.append(Shape("chart", shape.name, x, y, w, h, "", [], ph, **paint))
                return
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape.__class__.__name__ == "Picture":
                shapes.append(Shape("picture", shape.name, x, y, w, h, "", [], ph, **paint))
                return
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                paras = []
                for p in shape.text_frame.paragraphs:
                    sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
                    size = max(sizes) if sizes else (p.font.size.pt if p.font.size is not None else None)
                    paras.append((p.text if hasattr(p, "text") else "".join(r.text for r in p.runs), size))
                text = "\n".join(t for t, _ in paras).strip()
                shapes.append(Shape("text", shape.name, x, y, w, h, text, paras, ph,
                                    color=_text_color(shape.text_frame), **paint))
                return
            shapes.append(Shape("other", shape.name, x, y, w, h, "", [], ph, **paint))

        for shape in slide.shapes:
            add(shape)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
        slides.append(Slide(shapes=shapes, notes=notes or "", title=_slide_title(shapes),
                            background=_background(slide)))
    return Deck(width=width, height=height, slides=slides)


def _srgb(solid_fill) -> str | None:
    """The colour of an ``<a:solidFill>`` holding a plain sRGB value. A theme, scheme, preset or system colour, or an
    sRGB value with a transform (alpha, lumMod, ...), is not read: None."""
    clr = solid_fill.find(A + "srgbClr") if solid_fill is not None else None
    if clr is None or len(clr):
        return None
    val = (clr.get("val") or "").upper()
    return val if re.fullmatch(r"[0-9A-F]{6}", val) else None


def _paint(el) -> tuple[str | None, bool]:
    """(solid sRGB fill, whether it paints over its box) of a shape element. A picture, table or chart paints a box
    whose colour is not read; a shape with no fill of its own takes its style's fill, a theme one."""
    if el.tag in (P + "pic", P + "graphicFrame"):
        return None, True
    sppr = el.find(P + "spPr")
    for child in sppr if sppr is not None else ():
        name = child.tag.removeprefix(A) if isinstance(child.tag, str) else ""
        if name == "noFill":
            return None, False
        if name == "solidFill":
            return _srgb(child), True
        if name in ("gradFill", "blipFill", "pattFill", "grpFill"):
            return None, True
    ref = el.find(f"{P}style/{A}fillRef")
    return None, ref is not None and ref.get("idx", "0") != "0"


def _text_color(frame) -> str | None:
    """The explicit sRGB colour of the first run holding text; None when it inherits one (theme or placeholder)."""
    for p in frame.paragraphs:
        for r in p.runs:
            if r.text.strip():
                rpr = r._r.find(A + "rPr")
                return _srgb(rpr.find(A + "solidFill")) if rpr is not None else None
    return None


def _background(slide) -> str | None:
    """The solid sRGB background the slide shows: the first ``<p:bg>`` of the slide, its layout, its master."""
    parts = [slide]
    try:
        parts += [slide.slide_layout, slide.slide_layout.slide_master]
    except Exception:  # noqa: BLE001 - a slide with no layout shows only its own background
        pass
    for part in parts:
        bg = part._element.find(f"{P}cSld/{P}bg")
        if bg is not None:
            return _srgb(bg.find(f"{P}bgPr/{A}solidFill"))
    return None


def _slide_title(shapes: list[Shape]) -> str:
    for s in shapes:
        if s.placeholder in ("title", "ctrTitle") and s.text:
            return s.text
    for s in shapes:
        # runtime-office names a slide's heading shape "Title" (what office_read's outline reports as its title).
        if s.kind == "text" and s.name.strip().lower() == "title" and s.text:
            return s.text
    texts = [s for s in shapes if s.kind == "text" and s.text]
    if not texts:
        return ""
    # The largest text, topmost first: a slide built from HTML has no title placeholder.
    best = max(texts, key=lambda s: (max((sz or 0) for _, sz in s.paragraphs) if s.paragraphs else 0, -s.y))
    return best.text


# --------------------------------------------------------------------------- PDF


@dataclass
class Annotation:
    page: int  # 1-based
    subtype: str
    contents: str
    rect: tuple[float, float, float, float] | None  # PDF user space, origin bottom-left


@dataclass
class PdfDoc:
    pages: list[str]  # extracted text per page
    fields: dict[str, Any]  # name -> value as stored (str, or "/Yes" style names for checkboxes)
    field_pages: dict[str, int]
    annotations: list[Annotation]
    line_y: dict[int, list[tuple[float, str]]]  # page -> [(baseline y, text run)]

    def text(self) -> str:
        return "\n".join(self.pages)

    def y_of(self, phrase: str) -> tuple[int, float] | None:
        """(page, baseline y) of the first text run holding *phrase*."""
        want = norm(phrase)
        for page, runs in self.line_y.items():
            for y, text in runs:
                if want in norm(text):
                    return page, y
        return None


def read_pdf(path: str) -> PdfDoc:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages: list[str] = []
    line_y: dict[int, list[tuple[float, str]]] = {}
    annotations: list[Annotation] = []
    field_pages: dict[str, int] = {}
    for i, page in enumerate(reader.pages, start=1):
        runs: list[tuple[float, str]] = []

        def visitor(text, cm, tm, font_dict, font_size, _runs=runs):
            if text and text.strip():
                _runs.append((float(tm[5]) * float(cm[3]) + float(cm[5]), text))

        try:
            pages.append(page.extract_text(visitor_text=visitor) or "")
        except Exception:  # noqa: BLE001 - a page pypdf cannot read is empty text, not a crash
            pages.append("")
        line_y[i] = runs
        for ref in page.get("/Annots") or []:
            try:
                annot = ref.get_object()
            except Exception:  # noqa: BLE001
                continue
            subtype = str(annot.get("/Subtype", ""))
            rect = annot.get("/Rect")
            rect_t = tuple(float(v) for v in rect) if rect else None
            if subtype == "/Widget":
                name = annot.get("/T") or (annot.get("/Parent") or {}).get("/T")
                if name:
                    field_pages[str(name)] = i
                continue
            contents = annot.get("/Contents") or ""
            if not contents and annot.get("/RC"):
                contents = re.sub(r"<[^>]+>", " ", str(annot.get("/RC")))
            annotations.append(Annotation(i, subtype.lstrip("/"), str(contents), rect_t))
    fields: dict[str, Any] = {}
    for name, f in (reader.get_fields() or {}).items():
        value = f.get("/V")
        fields[name] = str(value) if value is not None else None
    return PdfDoc(pages=pages, fields=fields, field_pages=field_pages, annotations=annotations, line_y=line_y)


def read(path: str):
    kind = kind_of(path)
    return {"docx": read_docx, "xlsx": read_xlsx, "pptx": read_pptx, "pdf": read_pdf}[kind](path)


# --------------------------------------------------------------------------- Producer

# How the document editor names itself in what it saves, before a slash and its version.
EDITOR_PRODUCER = "euro-office"
_APPLICATION = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Application"


def producer(path: Path) -> str | None:
    """The program that saved the document, as it names itself: a package's ``docProps/app.xml`` ``<Application>``, a
    PDF's ``/Producer``. None when the document does not say, or that part cannot be read."""
    if kind_of(str(path)) == "pdf":
        from pypdf import PdfReader

        try:
            name = (PdfReader(path).metadata or {}).get("/Producer")
        except Exception:  # noqa: BLE001 - a PDF whose metadata pypdf cannot read names no producer
            return None
    else:
        try:
            with zipfile.ZipFile(path) as z:
                app = ET.fromstring(z.read("docProps/app.xml")).find(_APPLICATION)
        except (KeyError, zipfile.BadZipFile, ET.ParseError):
            return None
        name = app.text if app is not None else None
    return str(name or "").strip() or None
