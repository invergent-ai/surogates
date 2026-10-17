"""Every shape stays on its slide. A group's children are placed in the group's own space, so they are judged by
where the group itself lies on the slide."""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from obench.checks import Ctx, p_in_bounds
from obench.docs import read_pptx

IN_BOUNDS = "every shape inside the slide"


def _blank():
    deck = Presentation()  # 720 x 540 pt
    return deck, deck.slides.add_slide(deck.slide_layouts[6])


def _score(deck, tmp_path):
    path = str(tmp_path / "deck.pptx")
    deck.save(path)
    return p_in_bounds().run(read_pptx(path), Ctx())


def _group(slide, name="Badge"):
    """A group holding one 200x100 pt rectangle at (100, 100)."""
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(100), Pt(100), Pt(200), Pt(100))
    child.name = name
    return group


def test_a_shape_inside_the_slide_passes(tmp_path):
    deck, slide = _blank()
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(100), Pt(100), Pt(200), Pt(100))
    result = _score(deck, tmp_path)
    assert result.name == IN_BOUNDS
    assert result.passed is True, result.detail


def test_a_shape_past_the_slide_fails(tmp_path):
    deck, slide = _blank()
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(600), Pt(100), Pt(200), Pt(100)).name = "Wide"
    result = _score(deck, tmp_path)
    assert result.passed is False
    assert "Wide" in result.detail


def test_a_group_inside_the_slide_passes(tmp_path):
    deck, slide = _blank()
    _group(slide)
    assert _score(deck, tmp_path).passed is True


def test_a_group_moved_off_the_slide_fails(tmp_path):
    """The group's child still reads (100, 100) in the group's space; the group itself lies at x = 900 on a 720 pt
    slide."""
    deck, slide = _blank()
    group = _group(slide)
    group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(120), Pt(120), Pt(40), Pt(40))
    group.left = Pt(900)
    result = _score(deck, tmp_path)
    assert result.passed is False
    assert "Badge" in result.detail and "(900,100)" in result.detail
    assert result.detail.count("the group holding") == 1  # one group, named once whatever it holds
