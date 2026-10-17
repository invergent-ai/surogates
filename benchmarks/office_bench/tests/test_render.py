"""Text overflow measured on the document server's own render of a deck.

The fixtures were made with python-pptx and rendered to PDF by the document server's converter (x2t):
``fits.pptx`` is one slide with a 600x100 pt text box holding one short line at 24 pt; ``overflows.pptx`` is one slide
with a 300x40 pt box named ``Crowded`` holding four lines at 28 pt.
"""
import os
import pathlib
import re
import shlex
import sys

import pytest

from obench.checks import Ctx, p_no_overflow
from obench.docs import Deck, Shape, Slide, read_pptx
from obench.render import (
    RENDER_SHELL_ENV,
    X2T_ENV,
    Line,
    measure_lines,
    measure_overflow,
    render_pdf,
    rendered_lines,
)

FIXTURES = pathlib.Path(__file__).parent / "fixtures"


def _python(code: str) -> str:
    """A renderer command prefix that runs *code*; the arguments the renderer appends land in its sys.argv."""
    return f"{shlex.quote(sys.executable)} -c {shlex.quote(code)}"


def _measured(name: str):
    deck = read_pptx(str(FIXTURES / f"{name}.pptx"))
    measure_overflow((FIXTURES / f"{name}.pdf").read_bytes(), deck)
    return deck


# --------------------------------------------------------------------------- the renderer


def test_no_renderer_means_unknown(monkeypatch):
    monkeypatch.delenv(RENDER_SHELL_ENV, raising=False)
    result = render_pdf(b"x", "pptx")
    assert result.pdf is None and "OFFICEBENCH_RENDER_SHELL" in result.reason


def test_renderer_output_is_the_pdf(monkeypatch):
    pdf = FIXTURES / "fits.pdf"
    monkeypatch.setenv(RENDER_SHELL_ENV, _python(
        f"import shutil, sys; shutil.copyfileobj(open({str(pdf)!r}, 'rb'), sys.stdout.buffer)"))
    result = render_pdf((FIXTURES / "fits.pptx").read_bytes(), "pptx")
    assert result.pdf == pdf.read_bytes()


def test_failing_renderer_is_unknown(monkeypatch):
    monkeypatch.setenv(RENDER_SHELL_ENV, _python("import sys; sys.exit(1)"))
    result = render_pdf(b"x", "pptx")
    assert result.pdf is None and "status 1" in result.reason

    monkeypatch.setenv(RENDER_SHELL_ENV, _python("pass"))
    result = render_pdf(b"x", "pptx")
    assert result.pdf is None and "empty" in result.reason

    monkeypatch.setenv(RENDER_SHELL_ENV, _python("import time; time.sleep(10)"))
    result = render_pdf(b"x", "pptx", timeout=0.5)
    assert result.pdf is None and "timed out" in result.reason


def test_a_renderer_that_cannot_be_parsed_is_unknown(monkeypatch):
    """A prefix with an unbalanced quote is a reason, not an exception: scoring goes on after the session was spent."""
    monkeypatch.setenv(RENDER_SHELL_ENV, "sh -c 'unclosed")
    result = render_pdf(b"x", "pptx")
    assert result.pdf is None and "OFFICEBENCH_RENDER_SHELL cannot be parsed" in result.reason


def test_the_converter_runs_from_its_directory_on_a_parameters_file(monkeypatch, tmp_path):
    """The script the renderer runs: the document in, x2t on a parameters file from x2t's directory, the PDF out."""
    seen = tmp_path / "seen"
    x2t = tmp_path / "bin" / "x2t"
    x2t.parent.mkdir()
    x2t.write_text(
        f"#!{sys.executable}\n"
        "import os, re, sys\n"
        "params = open(sys.argv[1]).read()\n"
        "src = re.search('<m_sFileFrom>(.*?)</m_sFileFrom>', params)[1]\n"
        "dst = re.search('<m_sFileTo>(.*?)</m_sFileTo>', params)[1]\n"
        "fmt = re.search('<m_nFormatTo>(.*?)</m_nFormatTo>', params)[1]\n"
        f"open({str(seen)!r}, 'w').write('\\n'.join([os.getcwd(), src, fmt]))\n"
        "open(dst, 'wb').write(b'%PDF ' + open(src, 'rb').read())\n"
    )
    x2t.chmod(0o755)
    monkeypatch.setenv(RENDER_SHELL_ENV, "env")
    monkeypatch.setenv(X2T_ENV, str(x2t))
    result = render_pdf(b"the deck", "pptx")
    assert result.pdf == b"%PDF the deck", result.reason
    cwd, src, fmt = seen.read_text().splitlines()
    assert cwd == str(x2t.parent) and src.endswith("/in.pptx") and fmt == "513"
    assert not os.path.exists(os.path.dirname(src)), "the working directory is removed"


# --------------------------------------------------------------------------- measuring


def test_text_that_fits():
    deck = _measured("fits")
    [box] = deck.slides[0].shapes
    assert box.overflow == ""


def test_text_that_overflows():
    deck = _measured("overflows")
    crowded = next(s for s in deck.slides[0].shapes if s.name == "Crowded")
    match = re.fullmatch(r"text ends (\d+) pt past the box's bottom", crowded.overflow or "")
    assert match, crowded.overflow
    # The fourth 28 pt line's baseline sits about 91 pt below the 40 pt box.
    assert int(match[1]) == 91


def test_a_line_ends_at_its_baseline():
    """Descenders below a box are not an overflow a reader sees: a line's bottom is its baseline."""
    deck = read_pptx(str(FIXTURES / "fits.pptx"))
    [[line]] = rendered_lines((FIXTURES / "fits.pdf").read_bytes(), deck.width, deck.height)
    # The PDF puts the 24 pt line's baseline at y = 453.936 on a 540 pt page.
    assert line.bottom == pytest.approx(540 - 453.936, abs=0.01)
    assert line.top == pytest.approx(line.bottom - 1.2 * 24, abs=0.01)


def test_text_of_no_shape_is_ignored():
    deck = read_pptx(str(FIXTURES / "fits.pptx"))
    pages = rendered_lines((FIXTURES / "fits.pdf").read_bytes(), deck.width, deck.height)
    [box] = deck.slides[0].shapes
    bottom = box.y + box.h
    # Under the box and inside its width: it would overflow the box, were it the box's text.
    pages[0].append(Line("Nothing like this is on the slide", box.x + 10, bottom + 20, box.x + 300, bottom + 50))
    measure_lines(pages, deck)
    assert [s.overflow for s in deck.slides[0].shapes] == [""]


def test_group_children_are_not_measured(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt

    # fits.pptx's text box, inside a group: its offsets are the group's, so it is not measured.
    deck = Presentation(str(FIXTURES / "fits.pptx"))
    slide = deck.slides[0]
    box = slide.shapes[0]
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(box.left, box.top, box.width, box.height)
    child.text_frame.text = box.text_frame.text
    child.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    box._element.getparent().remove(box._element)
    path = str(tmp_path / "grouped.pptx")
    deck.save(path)
    model = read_pptx(path)
    measure_overflow((FIXTURES / "fits.pdf").read_bytes(), model)
    [grouped] = model.slides[0].shapes
    assert grouped.text == "A short line that fits" and grouped.overflow is None


def _slide(*shapes: Shape) -> Deck:
    return Deck(720.0, 540.0, [Slide(list(shapes), "", "")])


def _text(name: str, x: float, y: float, w: float, h: float, text: str) -> Shape:
    return Shape("text", name, x, y, w, h, text, [(t, 24.0) for t in text.split("\n")], None)


def test_a_line_past_the_right_side_overflows():
    deck = _slide(_text("Wide", 60, 60, 100, 40, "This single line is far wider than its box"))
    measure_lines([[Line("This single line is far wider than its box", 67, 62, 349, 84)]], deck)
    assert deck.slides[0].shapes[0].overflow == "text ends 189 pt past the box's right side"


def test_a_charts_own_labels_are_not_the_title_running_over():
    title = _text("Title", 40, 30, 640, 60, "Revenue is up 50% since January")
    chart = Shape("chart", "Chart", 40, 120, 640, 380, "", [], None)
    deck = _slide(title, chart)
    lines = [Line("Revenue is up 50% since January", 47, 37, 520, 80), Line("Jan", 172, 460, 196, 481),
             Line("50", 80, 265, 98, 286)]
    measure_lines([lines], deck)
    assert title.overflow == ""


def test_a_list_marker_the_text_does_not_hold_is_still_its_line():
    body = _text("Body", 60, 100, 400, 80, "Hire two engineers\nOpen the Lisbon office")
    deck = _slide(body)
    lines = [Line("• Hire two engineers", 70, 102, 260, 130), Line("2. Open the Lisbon office", 70, 190, 300, 218)]
    measure_lines([lines], deck)
    assert body.overflow == "text ends 38 pt past the box's bottom"


def test_text_running_into_the_box_below_overflows():
    title = _text("Title", 60, 40, 600, 70, "A long title that wraps onto a second line")
    subtitle = _text("Subtitle", 60, 115, 600, 40, "The subtitle")
    deck = _slide(title, subtitle)
    lines = [Line("A long title that wraps onto", 67, 64, 560, 100), Line("a second line", 67, 114, 300, 150),
             Line("The subtitle", 67, 125, 200, 150)]
    measure_lines([lines], deck)
    assert title.overflow == "text ends 40 pt past the box's bottom"
    assert subtitle.overflow == ""


def test_a_list_number_or_a_slide_number_is_not_a_distant_box_running_over():
    # The render draws a list's "1." and a slide's number "3"; neither box stores that text, but the title holds both.
    title = _text("Title", 60, 40, 600, 50, "Q1. Results and the Q3 plan")
    body = _text("Body", 60, 120, 600, 200, "Hire two engineers")
    number = _text("Slide Number", 600, 500, 60, 24, "‹#›")
    deck = _slide(title, body, number)
    lines = [Line("Q1. Results and the Q3 plan", 67, 44, 500, 80), Line("1.", 70, 124, 84, 148),
             Line("Hire two engineers", 90, 124, 300, 148), Line("3", 640, 502, 646, 520)]
    measure_lines([lines], deck)
    assert [title.overflow, body.overflow, number.overflow] == ["", "", None]


def test_a_render_with_other_pages_than_slides_is_not_measured():
    deck = read_pptx(str(FIXTURES / "fits.pptx"))
    pages = rendered_lines((FIXTURES / "fits.pdf").read_bytes(), deck.width, deck.height)
    measure_lines(pages + [[]], deck)
    assert deck.slides[0].shapes[0].overflow is None
    assert "2 pages for 1 slide" in deck.render_note


# --------------------------------------------------------------------------- the check


def test_overflow_check_reads_the_measurement():
    check = p_no_overflow()
    assert check.name == "no text overflows its box"

    unrendered = read_pptx(str(FIXTURES / "fits.pptx"))
    unrendered.render_note = "OFFICEBENCH_RENDER_SHELL is not set"
    passed, detail = check.fn(unrendered, Ctx())
    assert passed is None and "OFFICEBENCH_RENDER_SHELL is not set" in detail

    passed, detail = check.fn(_measured("fits"), Ctx())
    assert passed is True

    passed, detail = check.fn(_measured("overflows"), Ctx())
    assert passed is False and "Crowded" in detail and "pt past" in detail
