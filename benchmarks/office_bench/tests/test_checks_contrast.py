"""Text against what lies under it, by the colours the deck stores.

A text box is judged by its first run's explicit sRGB colour against its own solid fill, else the fill of the nearest
shape drawn before it whose box holds its centre, else the slide's solid background. Theme colours are not resolved:
text that uses them, or that sits on a picture or a theme fill, is not judged, and a deck where nothing could be
judged is unknown.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from obench.checks import Ctx, contrast_ratio, p_contrast
from obench.docs import read_pptx
from obench.tasks import DECK_TASKS, FIXTURES

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x14, 0x28, 0x5A)
CONTRAST = "text readable against what is under it"


def _blank():
    deck = Presentation()
    return deck, deck.slides.add_slide(deck.slide_layouts[6])


def _solid(fill, color) -> None:
    fill.solid()
    fill.fore_color.rgb = color


def _card(shapes, color=None):
    """A 400x200 pt rectangle at (100, 100), filled with *color*; with no colour it keeps its theme fill."""
    card = shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(100), Pt(100), Pt(400), Pt(200))
    if color is not None:
        _solid(card.fill, color)
    return card


def _text(shapes, text, color, fill=None):
    """A text box over the middle of the card."""
    box = shapes.add_textbox(Pt(120), Pt(150), Pt(360), Pt(40))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.color.rgb = color
    if fill is not None:
        _solid(box.fill, fill)
    return box


def _score(deck, tmp_path):
    path = str(tmp_path / "deck.pptx")
    deck.save(path)
    return p_contrast().run(read_pptx(path), Ctx())


def test_contrast_ratio():
    assert contrast_ratio("FFFFFF", "000000") == 21.0
    assert round(contrast_ratio("FFFFFF", "FFFFFF"), 1) == 1.0
    assert contrast_ratio("000000", "FFFFFF") == contrast_ratio("FFFFFF", "000000")


def test_white_on_white_card_fails(tmp_path):
    deck, slide = _blank()
    _card(slide.shapes, WHITE)
    _text(slide.shapes, "Ticketing and help centre", WHITE)
    result = _score(deck, tmp_path)
    assert result.name == CONTRAST
    assert result.passed is False
    assert "Ticketing and help centre" in result.detail and "1.0:1" in result.detail


def test_white_on_dark_band_passes(tmp_path):
    deck, slide = _blank()
    _card(slide.shapes, NAVY)
    _text(slide.shapes, "Tools we use", WHITE)
    assert _score(deck, tmp_path).passed is True


def test_own_fill_counts_first(tmp_path):
    deck, slide = _blank()
    _solid(slide.background.fill, WHITE)
    _text(slide.shapes, "Tools we use", WHITE, fill=NAVY)
    assert _score(deck, tmp_path).passed is True


def test_theme_colours_are_unknown():
    deck = read_pptx(str(FIXTURES / "sample.pptx"))
    result = p_contrast().run(deck, Ctx())
    assert result.passed is None, result.detail


def test_slide_background_is_the_layouts_or_masters_when_the_slide_has_none(tmp_path):
    deck, slide = _blank()
    _solid(deck.slide_master.background.fill, NAVY)
    _text(slide.shapes, "Tools we use", WHITE)
    assert _score(deck, tmp_path).passed is True

    _solid(slide.background.fill, WHITE)
    assert _score(deck, tmp_path).passed is False


def test_text_on_a_theme_fill_or_a_picture_is_not_judged(tmp_path):
    """What lies under the text paints its box with a colour not read here: unknown, not the slide's background."""
    deck, slide = _blank()
    _solid(slide.background.fill, WHITE)
    _card(slide.shapes)
    _text(slide.shapes, "Tools we use", WHITE)
    assert _score(deck, tmp_path).passed is None

    from PIL import Image

    image = str(tmp_path / "photo.png")
    Image.new("RGB", (40, 20), (20, 40, 90)).save(image)
    deck, slide = _blank()
    _solid(slide.background.fill, WHITE)
    slide.shapes.add_picture(image, Pt(100), Pt(100), Pt(400), Pt(200))
    _text(slide.shapes, "Tools we use", WHITE)
    assert _score(deck, tmp_path).passed is None


def test_a_box_with_no_fill_lets_what_is_under_it_show(tmp_path):
    deck, slide = _blank()
    _solid(slide.background.fill, WHITE)
    _card(slide.shapes, NAVY)
    clear = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(110), Pt(110), Pt(380), Pt(180))
    clear.fill.background()  # no fill
    _text(slide.shapes, "Tools we use", WHITE)
    assert _score(deck, tmp_path).passed is True


def test_shapes_in_a_group_are_not_judged(tmp_path):
    """A group's children are placed in the group's own space, not the slide's."""
    deck, slide = _blank()
    _solid(slide.background.fill, WHITE)
    group = slide.shapes.add_group_shape()
    _card(group.shapes, WHITE)
    _text(group.shapes, "Ticketing and help centre", WHITE)
    assert _score(deck, tmp_path).passed is None


def test_every_deck_task_checks_contrast():
    assert all(CONTRAST in [c.name for c in t.checks] for t in DECK_TASKS)
