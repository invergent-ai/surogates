"""How ``docs.read`` reads a Word document's comments: threads, replies and resolved state as OOXML stores them.

Each comment, a reply included, is a ``w:comment`` in ``word/comments.xml``; ``word/commentsExtended.xml`` keys a
``w15:commentEx`` by the ``w14:paraId`` of the comment's last paragraph, links a reply with ``w15:paraIdParent`` and
marks a resolved thread with ``w15:done``. The documents here are built with python-docx and those parts written by
hand, so each case is exactly the XML it names.
"""
import pytest

from obench import docs
from obench.tasks import FIXTURES

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"


def _saved(tmp_path, comments: list[tuple[str, str | None, list[str]]],
           extended: list[tuple[str, str | None, str | None]] | None) -> docs.WordDoc:
    """A document of two paragraphs, "Alpha beta." and "Gamma delta.", with *comments*: (text, where, paraIds), where
    is "both" (on both paragraphs), "first" or None (no range marks, as a reply may be), and paraIds one per paragraph
    of the comment's text (its lines). *extended*: the commentsExtended entries (paraId, paraIdParent, done), any of
    the last two None to leave the attribute out; None writes no commentsExtended part at all."""
    import docx
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml.ns import qn

    d = docx.Document()
    first, second = d.add_paragraph("Alpha beta."), d.add_paragraph("Gamma delta.")
    for text, where, para_ids in comments:
        if where == "both":
            c = d.add_comment([first.runs[0], second.runs[0]], text=text, author="Ana")
        elif where == "first":
            c = d.add_comment(first.runs[0], text=text, author="Ana")
        else:
            c = d.comments.add_comment(text=text, author="Rui")
        paras = c._comment_elm.findall(qn("w:p"))
        assert len(paras) == len(para_ids)
        for p, pid in zip(paras, para_ids):
            p.set(f"{{{W14}}}paraId", pid)
    if extended is not None:
        entries = []
        for pid, parent, done in extended:
            attrs = f'w15:paraId="{pid}"'
            attrs += f' w15:paraIdParent="{parent}"' if parent is not None else ""
            attrs += f' w15:done="{done}"' if done is not None else ""
            entries.append(f"<w15:commentEx {attrs}/>")
        blob = f'<w15:commentsEx xmlns:w15="{W15}">{"".join(entries)}</w15:commentsEx>'.encode()
        part = Part(PackURI("/word/commentsExtended.xml"),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml", blob,
                    d.part.package)
        d.part.relate_to(part, "http://schemas.microsoft.com/office/2011/relationships/commentsExtended")
    path = str(tmp_path / "comments.docx")
    d.save(path)
    return docs.read(path)


def test_a_document_without_comments_has_none():
    assert docs.read(str(FIXTURES / "sample.docx")).comments == []


def test_without_a_commentsextended_part_every_comment_is_an_open_thread(tmp_path):
    d = _saved(tmp_path, [("One", "first", ["0000000A"]), ("Two", None, ["0000000B"]), ("Three", "both", ["0000000C"])],
               None)
    assert [c.text for c in d.threads] == ["One", "Two", "Three"]
    assert not any(c.resolved or c.parent or c.replies for c in d.comments)


def test_anchor_is_the_text_between_the_range_marks(tmp_path):
    d = _saved(tmp_path, [("On both", "both", ["0000000A"]), ("On one", "first", ["0000000B"]),
                          ("Marks none", None, ["0000000C"])], None)
    assert [(c.text, c.anchor, c.author) for c in d.comments] == [
        ("On both", "Alpha beta. Gamma delta.", "Ana"), ("On one", "Alpha beta.", "Ana"), ("Marks none", "", "Rui")]


def test_a_reply_and_a_reply_to_it_belong_to_the_thread(tmp_path):
    d = _saved(tmp_path, [("Question", "first", ["000000A1"]), ("Answer", None, ["000000B1"]),
                          ("Answer to the answer", None, ["000000C1"])],
               [("000000A1", None, "1"), ("000000B1", "000000A1", "0"), ("000000C1", "000000B1", "0")])
    assert [c.text for c in d.threads] == ["Question"]
    thread = d.threads[0]
    assert thread.resolved
    assert [(r.text, r.parent) for r in thread.replies] == [("Answer", thread.id), ("Answer to the answer", thread.id)]


@pytest.mark.parametrize("done,resolved", [("1", True), ("true", True), ("on", True), ("0", False), ("false", False),
                                           (None, False)])
def test_done_marks_a_thread_resolved(tmp_path, done, resolved):
    d = _saved(tmp_path, [("Question", "first", ["0000000A"])], [("0000000A", None, done)])
    assert d.threads[0].resolved is resolved


def test_the_entry_is_matched_by_the_last_paragraph_or_another_of_the_comments(tmp_path):
    # Two entries for one comment: the last paragraph's wins.
    d = _saved(tmp_path, [("First line\nSecond line", "first", ["0000000A", "0000000B"])],
               [("0000000A", None, "0"), ("0000000B", None, "1")])
    assert d.threads[0].text == "First line\nSecond line" and d.threads[0].resolved
    # An entry naming the first paragraph, and a reply naming it as its parent, in lower case.
    d = _saved(tmp_path, [("First line\nSecond line", "first", ["0000000A", "0000000B"]),
                          ("Reply", None, ["0000000C"])],
               [("0000000a", None, "1"), ("0000000C", "0000000a", "0")])
    assert d.threads[0].resolved and [r.text for r in d.threads[0].replies] == ["Reply"]


def test_a_reply_whose_parent_is_not_there_is_a_thread_of_its_own(tmp_path):
    d = _saved(tmp_path, [("Question", "first", ["0000000A"]), ("Stray", None, ["0000000B"])],
               [("0000000A", None, "0"), ("0000000B", "0000FFFF", "0")])
    assert [c.text for c in d.threads] == ["Question", "Stray"]
