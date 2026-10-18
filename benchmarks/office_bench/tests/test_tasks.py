"""The suite's checks against documents built here: a correct result passes every check, the untouched input fails.

This is what makes a check trustworthy: it passes the document a careful person would produce, and it catches the
document the agent left unchanged. The fixtures' contents the edit checks rely on are pinned too.
"""
import copy
import os
from collections.abc import Callable

import pytest

from obench import docs
from obench.checks import Ctx
from obench.tasks import FIXTURES, TASKS, _block_text, by_id

T = by_id()


def run_checks(task_id: str, path: str, inputs_changed: list[str] | None = None):
    task = T[task_id]
    inp = task.original_input
    original = None
    if inp is not None:
        src = str(inp.source())
        if inp.derive:
            derived = path + ".input" + os.path.splitext(path)[1]
            inp.derive(src, derived)
            src = derived
        original = docs.read(src)
    ctx = Ctx(original=original, final_message="", inputs_changed=inputs_changed or [])
    doc = docs.read(path)
    return [c.run(doc, ctx) for c in task.checks]


def failures(results):
    return [(r.name, r.detail) for r in results if r.passed is False]


# --------------------------------------------------------------------------- the suite itself


def test_suite_shape():
    ids = [t.id for t in TASKS]
    assert len(ids) == len(set(ids))
    assert 30 <= len(TASKS) <= 36
    # Creating from scratch stays close to half the suite as harder edits join it.
    assert sum(t.mode == "create" for t in TASKS) >= 0.45 * len(TASKS)
    assert {t.kind for t in TASKS} == {"docx", "xlsx", "pptx", "pdf"}
    pilot = [t for t in TASKS if t.pilot]
    assert len(pilot) == 6 and {t.kind for t in pilot} == {"docx", "xlsx", "pptx", "pdf"}
    assert sum(t.mode == "create" for t in pilot) >= 3


def test_every_deck_task_keeps_its_shapes_on_the_slide():
    decks = [t for t in TASKS if t.kind == "pptx"]
    assert decks
    assert [t.id for t in decks if "every shape inside the slide" not in [c.name for c in t.checks]] == []


def test_minutes_floor():
    """The notes the minutes are written from hold about 70 words; minutes that cover them run from about 80."""
    lengths = [c.name for c in T["docx-minutes"].checks if c.name.startswith("length ")]
    assert lengths == ["length 80-600 words"]


@pytest.mark.parametrize("task", TASKS, ids=lambda t: t.id)
def test_task_is_well_formed(task):
    assert docs.kind_of(task.target) == task.kind
    assert task.target in task.prompt, "the person names the file"
    assert len(task.checks) >= 3
    for inp in task.inputs:
        assert inp.source().is_file(), inp.fixture
        assert inp.name in task.prompt
    if task.mode == "edit":
        assert task.original_input is not None


def test_fixture_contents_are_what_the_checks_assume():
    d = docs.read(str(FIXTURES / "sample.docx"))
    texts = [p.text for p in d.paragraphs]
    assert "North grew fastest, then the east." in texts
    assert "Hiring could slip a quarter." in texts
    assert d.tables[0].rows == [["Region", "Growth"], ["North", "12%"]]
    assert [p.text for p in d.paragraphs if p.list_kind == "number"] == ["Hire two engineers", "Open the Lisbon office"]
    x = docs.read(str(FIXTURES / "sample.xlsx"))
    assert x.sheets == ["Sales", "Notes"] and x.raw("Sales", "D3") == "=B3+C3"
    p = docs.read(str(FIXTURES / "sample.pptx"))
    assert [s.title for s in p.slides] == ["Quarterly review", "Next quarter"]
    assert p.slides[1].notes == "Mention the lease."
    f = docs.read(str(FIXTURES / "sample.pdf"))
    assert set(f.fields) == {"client", "signed"} and len(f.pages) == 2
    assert f.y_of("thirty days") == (1, 712.0)


@pytest.mark.parametrize("task", [t for t in TASKS if t.mode == "edit"], ids=lambda t: t.id)
def test_untouched_input_fails_the_task(task, tmp_path):
    """Handing back the input unchanged must fail; what it keeps must pass."""
    src = str(task.original_input.source())
    path = str(tmp_path / task.target)
    if task.original_input.derive:
        task.original_input.derive(src, path)
    else:
        with open(src, "rb") as a, open(path, "wb") as b:
            b.write(a.read())
    results = run_checks(task.id, path)
    assert failures(results), "an unchanged document passed every check"
    assert not [r for r in results if r.passed is None and "crashed" in r.detail]


# --------------------------------------------------------------------------- Word, done right


def _docx(tmp_path, name="quarterly-report.docx"):
    import docx

    return docx.Document(str(FIXTURES / "sample.docx")), str(tmp_path / name)


def _para(d, text):
    return next(p for p in d.paragraphs if p.text == text)


def test_docx_reword_done_right(tmp_path):
    d, path = _docx(tmp_path)
    _para(d, "North grew fastest, then the east.").runs[0].text = "The east grew fastest, then the north."
    d.save(path)
    assert failures(run_checks("docx-reword", path)) == []


def test_docx_reword_that_loses_the_bold_fails(tmp_path):
    d, path = _docx(tmp_path)
    _para(d, "North grew fastest, then the east.").runs[0].text = "The east grew fastest, then the north."
    for r in d.paragraphs[1].runs:
        r.bold = False
    d.save(path)
    assert [n for n, _ in failures(run_checks("docx-reword", path))] == ["'revenue target' is bold"]


def test_docx_fix_table_done_right(tmp_path):
    d, path = _docx(tmp_path)
    d.tables[0].cell(1, 1).text = "14%"
    d.save(path)
    assert failures(run_checks("docx-fix-table", path)) == []


def test_docx_bold_done_right(tmp_path):
    d, path = _docx(tmp_path)
    _para(d, "Hiring could slip a quarter.").runs[0].bold = True
    d.save(path)
    assert failures(run_checks("docx-bold", path)) == []


def test_docx_continue_list_done_right_and_restarted(tmp_path):
    d, path = _docx(tmp_path)
    last = _para(d, "Open the Lisbon office")
    new = copy.deepcopy(last._p)
    last._p.addnext(new)
    from docx.text.paragraph import Paragraph

    Paragraph(new, last._parent).runs[0].text = "Renew the support contract"
    d.save(path)
    assert failures(run_checks("docx-continue-list", path)) == []

    # The same item in a list of its own (another numId that restarts) is caught.
    from docx.oxml import parse_xml

    numbering = d.part.numbering_part.element
    abstract = numbering.xpath("./w:num[@w:numId='5']/w:abstractNumId/@w:val")[0]
    numbering.append(parse_xml(
        '<w:num xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:numId="99">'
        f'<w:abstractNumId w:val="{abstract}"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
    ))
    ppr = new.get_or_add_pPr()
    numpr = parse_xml('<w:numPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      '<w:ilvl w:val="0"/><w:numId w:val="99"/></w:numPr>')
    ppr.append(numpr)
    d.save(path)
    fails = failures(run_checks("docx-continue-list", path))
    assert fails and "1. Renew the support contract" in fails[0][1]


def test_docx_add_section_done_right(tmp_path):
    d, path = _docx(tmp_path)
    d.add_heading("Next steps", level=2)
    d.add_paragraph("Next quarter we hire the two engineers, sign the Lisbon lease and review regional targets "
                    "with each lead by the end of the first month.")
    d.save(path)
    assert failures(run_checks("docx-add-section", path)) == []


def test_docx_add_section_as_bold_text_is_not_a_heading(tmp_path):
    d, path = _docx(tmp_path)
    d.add_paragraph().add_run("Next steps").bold = True
    d.add_paragraph("Next quarter we hire the two engineers, sign the Lisbon lease and review regional targets.")
    d.save(path)
    assert "not a heading" in failures(run_checks("docx-add-section", path))[0][1]


def test_docx_letter_done_right(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Northwind Lda")
    d.add_paragraph("Mr Daniel Costa, Costa Properties, Rua Augusta 112, Lisbon")
    d.add_paragraph("Dear Mr Costa,")
    d.add_paragraph("We are writing to give notice that we are ending the lease on our office at Rua do Ouro 45 on "
                    "31 January 2027, as the 90-day notice period in the contract requires. Please arrange the final "
                    "inspection in the week of 18 January. We also ask you to return our deposit of €4,500 to the "
                    "account you have on file. Thank you for your help over the years.")
    d.add_paragraph("Kind regards,")
    d.add_paragraph("Ana Ribeiro")
    d.add_paragraph("Operations Manager, Northwind Lda")
    path = str(tmp_path / "lease-termination-letter.docx")
    d.save(path)
    assert failures(run_checks("docx-letter", path)) == []


# --------------------------------------------------------------------------- Word, derived inputs


def _derived(task_id: str, tmp_path, name: str = "input.docx") -> str:
    inp = T[task_id].original_input
    path = str(tmp_path / name)
    inp.derive(str(inp.source()), path)
    return path


def test_derived_inputs_leave_the_fixture_as_it_is(tmp_path):
    import hashlib

    source = FIXTURES / "sample.docx"
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    for task_id in ("docx-long-edit", "docx-table-edit", "docx-many-threads"):
        _derived(task_id, tmp_path, f"{task_id}.docx")
        assert T[task_id].original_input.fixture == "sample.docx"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


def test_docx_long_edit_input(tmp_path):
    from obench.tasks import NEW_BUDGET, OLD_BUDGET

    d = docs.read(_derived("docx-long-edit", tmp_path))
    sections = [p.text for p in d.headings if p.heading == 2]
    assert [s.split(".")[0] for s in sections] == [str(n) for n in range(1, 21)]
    # About twelve pages at the sample's 11 pt body text, its margins and its spacing (some 450-500 words a page).
    assert 5000 <= d.body_words() <= 6500, d.body_words()
    places = [b for b in d.blocks if OLD_BUDGET in _block_text(b)]
    assert len(places) == 4
    assert sum(isinstance(b, docs.Table) for b in places) == 1
    assert NEW_BUDGET not in d.text()
    # Section 17 holds the figure in its text and in its table, which closes the section.
    blocks = d.blocks
    start = next(i for i, b in enumerate(blocks) if isinstance(b, docs.Para) and b.text.startswith("17. "))
    end = next(i for i, b in enumerate(blocks) if isinstance(b, docs.Para) and b.text.startswith("18. "))
    inside = blocks[start + 1:end]
    assert isinstance(inside[-1], docs.Table) and OLD_BUDGET in _block_text(inside[-1])
    assert any(isinstance(b, docs.Para) and OLD_BUDGET in b.text for b in inside)
    assert len(d.tables) >= 5
    # The same document every time.
    again = docs.read(_derived("docx-long-edit", tmp_path, "again.docx"))
    assert [_block_text(b) for b in again.blocks] == [_block_text(b) for b in d.blocks]


def _long_edit(tmp_path, places: str = "all", note: str | None = "The budget figure was corrected on 14 October.",
               note_after_table: bool = True) -> str:
    import docx

    from obench.tasks import NEW_BUDGET, OLD_BUDGET

    d = docx.Document(_derived("docx-long-edit", tmp_path))
    heading = next(p for p in d.paragraphs if p.text.startswith("17. "))
    table = next(t for t in d.tables if OLD_BUDGET in t._tbl.xml)
    body = list(d.element.body)
    in_17 = range(body.index(heading._p), body.index(table._tbl) + 1)
    cells = [p for row in table.rows for c in row.cells for p in c.paragraphs]
    for p in list(d.paragraphs) + cells:
        # "section 17": only the section's own text and table; the summary and the approvals keep the old figure.
        if places == "section 17" and p not in cells and body.index(p._p) not in in_17:
            continue
        for r in p.runs:
            r.text = r.text.replace(OLD_BUDGET, NEW_BUDGET)
    if note:
        p = d.add_paragraph(note)
        (table._tbl if note_after_table else heading._p).addnext(p._p)
    path = str(tmp_path / "operations-report.docx")
    d.save(path)
    return path


def test_docx_long_edit_done_right(tmp_path):
    assert failures(run_checks("docx-long-edit", _long_edit(tmp_path))) == []


OLD_GONE = "'45,800' no longer appears anywhere except in the correction sentence"
NOTE = "a sentence right after section 17's table says it was corrected on 14 October"


def test_docx_long_edit_that_misses_a_place_fails(tmp_path):
    names = [n for n, _ in failures(run_checks("docx-long-edit", _long_edit(tmp_path, places="section 17")))]
    assert names == [OLD_GONE, "48,500 stands in every place 45,800 was"]


@pytest.mark.parametrize("note", [
    "The budget was corrected from €45,800 to €48,500 on 14 October.",
    "This figure was updated on 2026-10-14.",
    "The budget was revised on 14/10.",
    "The 2027 figure was amended on October 14.",
    "Correction made on the 14th of October: the budget is €48,500.",
    "The figure was fixed on 14 October.",
    "Fixed on 14 October: the budget is €48,500.",
])
def test_docx_long_edit_correction_sentence_in_other_words_passes(tmp_path, note):
    assert failures(run_checks("docx-long-edit", _long_edit(tmp_path, note=note))) == []


def test_docx_long_edit_old_figure_outside_the_correction_sentence_fails(tmp_path):
    # In a second sentence of the note.
    fails = dict(failures(run_checks("docx-long-edit", _long_edit(
        tmp_path, note="The budget was corrected on 14 October. It was €45,800."))))
    assert list(fails) == [OLD_GONE, NOTE] and "2 sentences" in fails[NOTE]
    # In a correction sentence put under the heading, not after the table.
    fails = dict(failures(run_checks("docx-long-edit", _long_edit(
        tmp_path, note="The budget was corrected from €45,800 to €48,500 on 14 October.", note_after_table=False))))
    assert list(fails) == [OLD_GONE, NOTE]


@pytest.mark.parametrize("note", ["The figure here was incorrect until 14 October.",
                                  "The budget figure was corrected.",
                                  "The budget figure was corrected on 4 October."])
def test_docx_long_edit_note_that_does_not_say_corrected_on_14_october_fails(tmp_path, note):
    assert [n for n, _ in failures(run_checks("docx-long-edit", _long_edit(tmp_path, note=note)))] == [NOTE]


def test_docx_long_edit_note_in_the_wrong_place_or_missing_fails(tmp_path):
    fails = failures(run_checks("docx-long-edit", _long_edit(tmp_path, note_after_table=False)))
    assert [n for n, _ in fails] == [NOTE]
    fails = failures(run_checks("docx-long-edit", _long_edit(tmp_path, note=None)))
    assert [n for n, _ in fails] == [NOTE, "one paragraph added, none removed"]
    fails = failures(run_checks("docx-long-edit", _long_edit(
        tmp_path, note="The figure was corrected on 14 October. The old one was a typing error.")))
    assert "2 sentences" in dict(fails)[NOTE]


def test_docx_table_edit_input(tmp_path):
    from obench.tasks import REGION_TABLE

    d = docs.read(_derived("docx-table-edit", tmp_path))
    assert len(d.tables) == 1
    assert d.tables[0].rows == REGION_TABLE
    assert d.tables[0].style == "Table Grid"
    assert "North grew fastest, then the east." in [p.text for p in d.paragraphs]


def _table_edit(tmp_path, south=("South", "70", "65"), total=("Total", "280", "305")) -> str:
    import docx

    d = docx.Document(_derived("docx-table-edit", tmp_path))
    table = d.tables[0]
    total_tr = table.rows[-1]._tr
    row = table.add_row()
    total_tr.addprevious(row._tr)
    for cell, value in zip(row.cells, south):
        cell.text = value
    for cell, value in zip(table.rows[-1].cells, total):
        cell.text = value
    path = str(tmp_path / "quarterly-report.docx")
    d.save(path)
    return path


def test_docx_table_edit_done_right(tmp_path):
    assert failures(run_checks("docx-table-edit", _table_edit(tmp_path))) == []


def test_docx_table_edit_with_the_old_totals_fails(tmp_path):
    fails = failures(run_checks("docx-table-edit", _table_edit(tmp_path, total=("Total", "210", "240"))))
    assert [n for n, _ in fails] == ["the Total row reads 280 and 305"]
    assert "210" in fails[0][1]


def _table_rows(tmp_path, rows: list[list[str]]) -> str:
    """The derived input with its table's rows under the header replaced by *rows*."""
    import docx

    d = docx.Document(_derived("docx-table-edit", tmp_path))
    table = d.tables[0]
    for tr in table._tbl.tr_lst[1:]:
        table._tbl.remove(tr)
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    path = str(tmp_path / "quarterly-report.docx")
    d.save(path)
    return path


@pytest.mark.parametrize("rows", [
    [["South", "70", "65"], ["North", "120", "140"], ["East", "90", "100"], ["Total", "280", "305"]],
    [["East", "90", "100"], ["North", "120", "140"], ["South", "70", "65"], ["Total", "280", "305"]],
], ids=["south-first", "sorted-by-name"])
def test_docx_table_edit_in_another_row_order_passes(tmp_path, rows):
    assert failures(run_checks("docx-table-edit", _table_rows(tmp_path, rows))) == []


def test_docx_table_edit_that_changes_a_region_fails(tmp_path):
    rows = [["North", "125", "140"], ["East", "90", "100"], ["South", "70", "65"], ["Total", "280", "305"]]
    fails = failures(run_checks("docx-table-edit", _table_rows(tmp_path, rows)))
    assert [n for n, _ in fails] == ["the header and the North and East rows are kept"]
    assert "['North', '120', '140']" in fails[0][1]


def test_docx_many_threads_input(tmp_path):
    from obench.tasks import THREAD_NOTES, THREAD_QUESTIONS

    d = docs.read(_derived("docx-many-threads", tmp_path))
    threads = {c.text: c for c in d.threads}
    assert set(threads) == set(THREAD_QUESTIONS) | set(THREAD_NOTES)
    assert len(d.comments) == 5
    assert not any(c.resolved or c.replies or c.parent for c in d.comments)
    assert threads["How many engineers do we plan to hire?"].anchor == "Hiring could slip a quarter."
    assert threads["Keep 'revenue target' in bold, as finance asked."].anchor == "revenue target"
    # Each answer is in the document's text, not in the words its question is anchored to.
    for question, (pattern, _) in THREAD_QUESTIONS.items():
        import re

        assert re.search(pattern, docs.norm(d.text()))
        assert not re.search(pattern, docs.norm(threads[question].anchor + " " + question))
    # The comments add no text to the document: every other Word edit check still reads the sample.
    assert [p.text for p in d.paragraphs] == [p.text for p in docs.read(str(FIXTURES / "sample.docx")).paragraphs]
    # The person asks for replies in the threads, not answers anywhere.
    assert T["docx-many-threads"].prompt.startswith("Reply to the three questions in the comments on ")


W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"


def _threads_saved(tmp_path, replies: dict[str, str], resolved: set[str]) -> str:
    """The derived input with *replies* (comment text -> reply) and *resolved* threads, written as Word and the editor
    write them: each reply its own ``w:comment`` with range marks beside its parent's, every comment paragraph carrying
    a ``w14:paraId``, and ``commentsExtended.xml`` linking a reply to its parent's last paragraph (``w15:paraIdParent``)
    and marking a resolved thread ``w15:done="1"``."""
    import docx
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml.ns import qn
    from docx.oxml.parser import parse_xml

    d = docx.Document(_derived("docx-many-threads", tmp_path))
    comments = d.part._comments_part.element
    by_text = {"\n".join(p.text for p in c.paragraphs): c for c in d.comments}
    body = d.element.body
    for text, reply_text in replies.items():
        parent = by_text[text]
        reply = d.comments.add_comment(text=reply_text, author="Agent", initials="A")
        pid, rid = parent.comment_id, reply.comment_id
        start = body.xpath(f".//w:commentRangeStart[@w:id='{pid}']")[0]
        end = body.xpath(f".//w:commentRangeEnd[@w:id='{pid}']")[0]
        start.addnext(parse_xml(f'<w:commentRangeStart xmlns:w="{docs.NS["w"]}" w:id="{rid}"/>'))
        end.addnext(parse_xml(f'<w:commentRangeEnd xmlns:w="{docs.NS["w"]}" w:id="{rid}"/>'))
        by_text[f"reply to {text}"] = reply
    para_of = {}
    n = 0
    for c in comments.findall(qn("w:comment")):
        for p in c.findall(qn("w:p")):
            n += 1
            p.set(f"{{{W14}}}paraId", f"{n:08X}")
        para_of[c.get(qn("w:id"))] = f"{n:08X}"  # the comment's last paragraph
    entries = []
    for text, c in by_text.items():
        if text.startswith("reply to "):
            parent = by_text[text[len("reply to "):]]
            entries.append(f'<w15:commentEx w15:paraId="{para_of[str(c.comment_id)]}" '
                           f'w15:paraIdParent="{para_of[str(parent.comment_id)]}" w15:done="0"/>')
        else:
            done = "1" if text in resolved else "0"
            entries.append(f'<w15:commentEx w15:paraId="{para_of[str(c.comment_id)]}" w15:done="{done}"/>')
    blob = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w15:commentsEx xmlns:w15="{W15}">'
            + "".join(entries) + "</w15:commentsEx>").encode()
    part = Part(PackURI("/word/commentsExtended.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml", blob,
                d.part.package)
    d.part.relate_to(part, "http://schemas.microsoft.com/office/2011/relationships/commentsExtended")
    path = str(tmp_path / "quarterly-report.docx")
    d.save(path)
    return path


ANSWERS = {
    "How many engineers do we plan to hire?": "Two: the plan lists 'Hire two engineers'.\nBoth are in the hiring plan.",
    "How much did North grow this quarter?": "North grew 12% this quarter, as the table shows.",
    "Which city is the new office in?": "Lisbon.",
}


def test_docx_many_threads_done_right(tmp_path):
    path = _threads_saved(tmp_path, ANSWERS, set(ANSWERS))
    d = docs.read(path)
    replied = {c.text: [r.text for r in c.replies] for c in d.threads}
    assert replied["How many engineers do we plan to hire?"] == [ANSWERS["How many engineers do we plan to hire?"]]
    assert all(r.parent is not None for c in d.threads for r in c.replies)
    assert {c.text for c in d.threads if c.resolved} == set(ANSWERS)
    assert failures(run_checks("docx-many-threads", path)) == []


def test_docx_many_threads_wrong_answer_open_thread_or_answered_note_fails(tmp_path):
    from obench.tasks import THREAD_NOTES

    wrong = dict(ANSWERS, **{"How much did North grow this quarter?": "North grew 14% this quarter."})
    fails = failures(run_checks("docx-many-threads", _threads_saved(tmp_path, wrong, set(ANSWERS))))
    assert [n for n, _ in fails] == ["the reply to 'How much did North grow this quarter?' says 12%"]

    fails = failures(run_checks("docx-many-threads", _threads_saved(tmp_path, ANSWERS, set())))
    assert [n for n, _ in fails] == ["the three questions are resolved"]

    noted = dict(ANSWERS, **{THREAD_NOTES[0]: "Done, it stays bold."})
    fails = failures(run_checks("docx-many-threads", _threads_saved(tmp_path, noted, set(ANSWERS) | {THREAD_NOTES[1]})))
    assert [n for n, _ in fails] == ["the two notes are left open, unanswered and as they were"]
    assert "reply" in fails[0][1] and "resolved" in fails[0][1]

    long = dict(ANSWERS, **{"Which city is the new office in?": "Lisbon. " + "The office opens next year. " * 20})
    fails = failures(run_checks("docx-many-threads", _threads_saved(tmp_path, long, set(ANSWERS))))
    assert [n for n, _ in fails] == ["the reply to 'Which city is the new office in?' says Lisbon"]
    assert "400" in fails[0][1]


@pytest.mark.parametrize("answer,right", [
    ("North grew 12%.", True), ("North grew 12.", True), ("12 % this quarter.", True), ("12 pct.", True),
    ("North grew 12 per cent.", True), ("North grew twelve percent.", True),
    ("North grew 112%.", False), ("North grew 12.5%.", False),
])
def test_docx_many_threads_north_growth_written_any_way(tmp_path, answer, right):
    answers = dict(ANSWERS, **{"How much did North grow this quarter?": answer})
    names = [n for n, _ in failures(run_checks("docx-many-threads", _threads_saved(tmp_path, answers, set(ANSWERS))))]
    assert names == ([] if right else ["the reply to 'How much did North grow this quarter?' says 12%"])


def test_docx_many_threads_answer_left_as_a_new_thread_is_named(tmp_path):
    import docx

    d = docx.Document(_derived("docx-many-threads", tmp_path))
    para = next(p for p in d.paragraphs if p.text == "North grew fastest, then the east.")
    d.add_comment(para.runs[0], text="North grew 12% this quarter.", author="Agent", initials="A")
    path = str(tmp_path / "quarterly-report.docx")
    d.save(path)
    details = dict(failures(run_checks("docx-many-threads", path)))
    assert details["the reply to 'How much did North grow this quarter?' says 12%"] == (
        "no reply; a new thread says it instead: ['North grew 12% this quarter.']")


# --------------------------------------------------------------------------- Spreadsheets, done right


def _sample_xlsx():
    import openpyxl

    return openpyxl.load_workbook(str(FIXTURES / "sample.xlsx"))


def test_xlsx_add_column_done_right(tmp_path):
    book = _sample_xlsx()
    ws = book["Sales"]
    ws["E1"] = "Growth"
    for r in (2, 3, 4):
        ws[f"E{r}"] = f"=(C{r}-B{r})/B{r}"
        ws[f"E{r}"].number_format = "0.0%"
    path = str(tmp_path / "sales.xlsx")
    book.save(path)
    results = run_checks("xlsx-add-column", path)
    assert failures(results) == []


def test_xlsx_add_column_typed_numbers_fail(tmp_path):
    book = _sample_xlsx()
    ws = book["Sales"]
    ws["E1"] = "Growth"
    for r, v in ((2, 0.1667), (3, 0.2222), (4, -0.0714)):
        ws[f"E{r}"] = v
    path = str(tmp_path / "sales.xlsx")
    book.save(path)
    assert "not all formulas" in failures(run_checks("xlsx-add-column", path))[0][1]


def test_xlsx_fix_formula_done_right(tmp_path):
    book = _sample_xlsx()
    path = str(tmp_path / "sales.xlsx")
    book.save(path)
    assert failures(run_checks("xlsx-fix-formula", path)) == []


def test_xlsx_rename_sheet_done_right(tmp_path):
    book = _sample_xlsx()
    book["Notes"].title = "Read me"
    path = str(tmp_path / "sales.xlsx")
    book.save(path)
    assert failures(run_checks("xlsx-rename-sheet", path)) == []


def test_xlsx_budget_done_right(tmp_path):
    import openpyxl
    from openpyxl.chart import PieChart, Reference

    book = openpyxl.Workbook()
    ws = book.active
    rows = [("Income", None), ("Salary", 3200), ("Freelance", 650), ("Total income", "=SUM(B2:B3)"), ("", None),
            ("Expenses", None), ("Rent", 1150), ("Groceries", 420), ("Utilities", 180), ("Transport", 140),
            ("Insurance", 95), ("Savings", 500), ("Fun", 200), ("Total expenses", "=SUM(B7:B13)"), ("", None),
            ("Left over", "=B4-B14")]
    for i, (a, b) in enumerate(rows, start=1):
        ws[f"A{i}"], ws[f"B{i}"] = a, b
        ws[f"B{i}"].number_format = '"€"#,##0.00'
    chart = PieChart()
    chart.add_data(Reference(ws, min_col=2, min_row=7, max_row=13))
    chart.set_categories(Reference(ws, min_col=1, min_row=7, max_row=13))
    ws.add_chart(chart, "D2")
    path = str(tmp_path / "budget.xlsx")
    book.save(path)
    results = run_checks("xlsx-budget", path)
    assert failures(results) == []
    # openpyxl stores no computed values: the consistency check says it cannot tell.
    assert [r.passed for r in results if r.name.startswith("stored values")] == [None]


def test_xlsx_budget_hardcoded_totals_fail(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    ws = book.active
    for i, v in enumerate([3200, 650, 3850, 1150, 420, 180, 140, 95, 500, 200, 2685, 1165], start=1):
        ws[f"B{i}"] = v
    path = str(tmp_path / "budget.xlsx")
    book.save(path)
    details = dict(failures(run_checks("xlsx-budget", path)))
    assert "typed in as a number" in details["a formula computes total expenses = 2685"]


def test_xlsx_sales_table_done_right(tmp_path):
    import openpyxl
    from openpyxl.styles import Font

    from obench.tasks import SALES_ROWS

    book = openpyxl.Workbook()
    ws = book.active
    ws.append(["City", "Units", "Revenue"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for city, (units, revenue) in sorted(SALES_ROWS.items(), key=lambda kv: -kv[1][1]):
        ws.append([city, units, revenue])
        ws.cell(ws.max_row, 3).number_format = "#,##0"
    path = str(tmp_path / "sales-2026.xlsx")
    book.save(path)
    assert failures(run_checks("xlsx-sales-table", path)) == []


def test_xlsx_multi_sheet_and_dates_done_right(tmp_path):
    import datetime as dt

    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Jan"
    months = {"Jan": (120, 98, 140), "Feb": (132, 110, 138), "Mar": (128, 105, 150)}
    for i, (m, vals) in enumerate(months.items()):
        ws = book[m] if i == 0 else book.create_sheet(m)
        for row, (name, v) in enumerate(zip(("Ana", "Ben", "Carla"), vals), start=2):
            ws[f"A{row}"], ws[f"B{row}"] = name, v
    s = book.create_sheet("Summary")
    for row, name in enumerate(("Ana", "Ben", "Carla"), start=2):
        s[f"A{row}"], s[f"B{row}"] = name, f"=Jan!B{row}+Feb!B{row}+'Mar'!B{row}"
    path = str(tmp_path / "team-hours.xlsx")
    book.save(path)
    assert failures(run_checks("xlsx-multi-sheet", path)) == []

    book = openpyxl.Workbook()
    ws = book.active
    ws.append(["Task", "Owner", "Start", "End", "Days"])
    plan = [("Copy final", "Inês", 3, 7), ("Landing page", "Rui", 5, 14), ("Email campaign", "Ana", 10, 21),
            ("Press kit", "Inês", 12, 19), ("Launch event", "Marta", 24, 24)]
    for i, (t, o, a, b) in enumerate(plan, start=2):
        ws.append([t, o, dt.datetime(2026, 11, a), dt.datetime(2026, 11, b), f"=D{i}-C{i}+1"])
    path = str(tmp_path / "launch-plan.xlsx")
    book.save(path)
    assert failures(run_checks("xlsx-dates", path)) == []


HOURS = {"Jan": (120, 98, 140), "Feb": (132, 110, 138), "Mar": (128, 105, 150)}
TOTALS = {"Ana": 380, "Ben": 313, "Carla": 428}


def _team_hours(tmp_path, summary: Callable[[int, int], list]) -> str:
    """team-hours.xlsx with the prompt's month sheets and a Summary whose row for each person (2 to 4) is
    *summary*(row, index): the cells from B on."""
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Jan"
    for i, (month, hours) in enumerate(HOURS.items()):
        ws = book[month] if i == 0 else book.create_sheet(month)
        ws.append(["Person", "Hours"])
        for name, h in zip(TOTALS, hours):
            ws.append([name, h])
    s = book.create_sheet("Summary")
    s.append(["Person", "Jan", "Feb", "Mar", "Total"])
    for i, name in enumerate(TOTALS):
        s.append([name, *summary(i + 2, i)])
    path = str(tmp_path / "team-hours.xlsx")
    book.save(path)
    return path


def _total_checks(results) -> dict[str, tuple[bool | None, str]]:
    return {r.name: (r.passed, r.detail) for r in results if "by formula on Summary" in r.name}


@pytest.mark.parametrize("summary", [
    # Each month's column links to its sheet, and the total adds the row: what the agent most often builds.
    lambda row, i: [f"=Jan!B{row}", f"=Feb!B{row}", f"=Mar!B{row}", f"=SUM(B{row}:D{row})"],
    # The total points at the month sheets itself.
    lambda row, i: [f"=Jan!B{row}+Feb!B{row}+'Mar'!B{row}"],
    lambda row, i: [f"=SUM(Jan!B{row},Feb!B{row},Mar!B{row})"],
    # A link through a link is still a link.
    lambda row, i: [f"=Jan!B{row}", f"=Feb!B{row}", f"=Mar!B{row}", f"=B{row}+C{row}+D{row}", f"=E{row}"],
], ids=["row-links", "direct", "direct-sum", "link-of-links"])
def test_xlsx_multi_sheet_total_that_reaches_the_month_sheets_passes(tmp_path, summary):
    totals = _total_checks(run_checks("xlsx-multi-sheet", _team_hours(tmp_path, summary)))
    assert len(totals) == 3
    assert all(passed for passed, _ in totals.values()), totals


@pytest.mark.parametrize("summary, why", [
    # The month values typed again on Summary and added there: a copy, not a link.
    (lambda row, i: [*(HOURS[m][i] for m in HOURS), f"=SUM(B{row}:D{row})"], "typed on Summary"),
    # One month typed, two linked.
    (lambda row, i: [f"=Jan!B{row}", f"=Feb!B{row}", HOURS["Mar"][i], f"=SUM(B{row}:D{row})"], "typed on Summary"),
    # The total typed as a number, or worked out from numbers in the formula.
    (lambda row, i: [None, None, None, list(TOTALS.values())[i]], "no formula on Summary gives"),
    (lambda row, i: [None, None, None, "=" + "+".join(str(HOURS[m][i]) for m in HOURS)], "does not point at"),
    # Linked, but the wrong months: the value is wrong.
    (lambda row, i: [f"=Jan!B{row}", f"=Feb!B{row}", None, f"=SUM(B{row}:D{row})"], "no formula on Summary gives"),
], ids=["typed-copies", "one-typed-copy", "typed-total", "numbers-in-formula", "wrong-value"])
def test_xlsx_multi_sheet_total_that_does_not_reach_the_month_sheets_fails(tmp_path, summary, why):
    totals = _total_checks(run_checks("xlsx-multi-sheet", _team_hours(tmp_path, summary)))
    assert len(totals) == 3
    for passed, detail in totals.values():
        assert passed is False and why in detail, totals


def test_xlsx_multi_sheet_links_are_followed_a_bounded_number_of_cells_deep(tmp_path):
    from obench.tasks import LINK_STEPS, _reaches_other_sheets

    # B links to Jan; each next column reads the one before it, so column B + n is n cells away from the link.
    columns = "BCDEFGH"
    path = _team_hours(tmp_path, lambda row, i: [f"=Jan!B{row}", *(f"={c}{row}" for c in columns[:LINK_STEPS + 1])])
    book = docs.read(path)
    assert _reaches_other_sheets(book, "Summary", f"{columns[LINK_STEPS]}2")[0] is True
    linked, detail = _reaches_other_sheets(book, "Summary", f"{columns[LINK_STEPS + 1]}2")
    assert linked is False and f"more than {LINK_STEPS} cells" in detail


LAUNCH_PLAN = [("Copy final", "Inês", 3, 7), ("Landing page", "Rui", 5, 14), ("Email campaign", "Ana", 10, 21),
               ("Press kit", "Inês", 12, 19), ("Launch event", "Marta", 24, 24)]


def _store_computed(path: str, values: dict[str, float]) -> None:
    """Store each formula's computed value in the first sheet, as the editor's save does: openpyxl stores none."""
    import re
    import zipfile

    with zipfile.ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    for ref, value in values.items():
        sheet, found = re.subn(rf'(<c r="{ref}"[^>]*><f>[^<]*</f>)<v></v>', rf"\g<1><v>{value}</v>", sheet)
        assert found == 1, ref
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def _launch_plan(path: str, *, formulas: bool, year: int = 2026, number_format: str = "d mmm yyyy") -> str:
    """launch-plan.xlsx with the prompt's rows, each start and end in November of *year*, in *number_format*: dates
    typed in, or =DATE formulas whose values the file stores; the days by formula."""
    import datetime as dt

    import openpyxl

    book = openpyxl.Workbook()
    ws = book.active
    ws.append(["Task", "Owner", "Start", "End", "Days"])
    stored: dict[str, float] = {}
    for row, (task, owner, start, end) in enumerate(LAUNCH_PLAN, start=2):
        ws.append([task, owner, None, None, f"=D{row}-C{row}+1"])
        stored[f"E{row}"] = end - start + 1
        for column, day in (("C", start), ("D", end)):
            cell = ws[f"{column}{row}"]
            date = dt.datetime(year, 11, day)
            if formulas:
                cell.value = f"=DATE({year},11,{day})"
                stored[cell.coordinate] = (date - dt.datetime(1899, 12, 30)).days
            else:
                cell.value = date
            cell.number_format = number_format
    book.save(path)
    _store_computed(path, stored)
    return path


@pytest.mark.parametrize("formulas", [False, True], ids=["typed", "=DATE"])
def test_xlsx_dates_counts_a_date_typed_or_made_by_a_formula(tmp_path, formulas):
    path = _launch_plan(str(tmp_path / "launch-plan.xlsx"), formulas=formulas, number_format="dd-mmm-yyyy")
    results = run_checks("xlsx-dates", path)
    assert failures(results) == []
    # The days are worked out from the dates, whichever way they were written.
    assert [r.passed for r in results if r.name == "durations computed by formula"] == [True]


@pytest.mark.parametrize("formulas", [False, True], ids=["typed", "=DATE"])
@pytest.mark.parametrize("year, number_format", [(2025, "d mmm yyyy"), (2026, "General")], ids=["wrong year", "a number"])
def test_xlsx_dates_fails_a_wrong_date_or_a_number_not_shown_as_a_date(tmp_path, formulas, year, number_format):
    path = _launch_plan(str(tmp_path / "launch-plan.xlsx"), formulas=formulas, year=year, number_format=number_format)
    assert ("start and end are real dates (9+ in Nov 2026)", "0 real dates in Nov 2026") in failures(
        run_checks("xlsx-dates", path)
    )


def test_xlsx_invoice_done_right(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    ws = book.active
    ws.append(["Invoice 1043", "Acme Corp"])
    ws.append(["Item", "Qty", "Price", "Total"])
    for i, (item, q, p) in enumerate([("Website redesign", 1, 4800), ("Hosting", 12, 35), ("Support", 18, 85)], start=3):
        ws.append([item, q, p, f"=B{i}*C{i}"])
    ws.append(["Subtotal", None, None, "=SUM(D3:D5)"])
    ws.append(["VAT 23%", None, None, "=ROUND(D6*23%,2)"])
    ws.append(["Total due", None, None, "=D6+D7"])
    path = str(tmp_path / "invoice-1043.xlsx")
    book.save(path)
    assert failures(run_checks("xlsx-invoice", path)) == []


# --------------------------------------------------------------------------- Decks, done right


def _pptx():
    from pptx import Presentation

    return Presentation(str(FIXTURES / "sample.pptx"))


def test_pptx_add_bullet_done_right_and_in_a_separate_box(tmp_path):
    deck = _pptx()
    deck.slides[0].placeholders[1].text_frame.add_paragraph().text = "Costs came in 4% under budget"
    path = str(tmp_path / "quarterly-review.pptx")
    deck.save(path)
    assert failures(run_checks("pptx-add-bullet", path)) == []

    from pptx.util import Inches

    deck = _pptx()
    deck.slides[0].shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(0.6)).text_frame.text = (
        "Costs came in 4% under budget")
    deck.save(path)
    assert "separate box" in failures(run_checks("pptx-add-bullet", path))[0][1]


def test_pptx_change_title_and_notes_done_right(tmp_path):
    deck = _pptx()
    deck.slides[1].shapes.title.text = "Q3 plan"
    path = str(tmp_path / "quarterly-review.pptx")
    deck.save(path)
    assert failures(run_checks("pptx-change-title", path)) == []

    deck = _pptx()
    deck.slides[0].notes_slide.notes_text_frame.text = (
        "Thank the regional leads. Hiring slipped because two offers were declined.")
    deck.save(path)
    assert failures(run_checks("pptx-notes", path)) == []


def test_pptx_reorder_done_right(tmp_path):
    deck = _pptx()
    ids = deck.slides._sldIdLst
    first = ids[0]
    ids.remove(first)
    ids.append(first)
    path = str(tmp_path / "quarterly-review.pptx")
    deck.save(path)
    assert failures(run_checks("pptx-reorder", path)) == []


def test_pptx_table_slide_done_right(tmp_path):
    from pptx.util import Inches

    deck = _pptx()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Revenue by region"
    table = slide.shapes.add_table(4, 3, Inches(1), Inches(2), Inches(8), Inches(3)).table
    for r, row in enumerate([["Region", "Q1", "Q2"], ["North", "120", "140"], ["East", "90", "110"], ["South", "70", "65"]]):
        for c, v in enumerate(row):
            table.cell(r, c).text = v
    path = str(tmp_path / "quarterly-review.pptx")
    deck.save(path)
    assert failures(run_checks("pptx-table-slide", path)) == []


# --------------------------------------------------------------------------- PDFs, done right


def test_pdf_fill_done_right(tmp_path):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(str(FIXTURES / "sample.pdf")))
    writer.update_page_form_field_values(writer.pages[0], {"client": "Northwind Lda", "signed": "/Yes"},
                                         auto_regenerate=False)
    path = str(tmp_path / "service-agreement.pdf")
    writer.write(path)
    assert failures(run_checks("pdf-fill", path)) == []


def test_pdf_annotate_done_right_and_far_away(tmp_path):
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import FreeText

    def write(rect):
        writer = PdfWriter(clone_from=PdfReader(str(FIXTURES / "sample.pdf")))
        writer.add_annotation(0, FreeText(text="Finance wants 45 days, check with legal.", rect=rect))
        path = str(tmp_path / "service-agreement.pdf")
        writer.write(path)
        return path

    assert failures(run_checks("pdf-annotate", write((380, 700, 560, 730)))) == []
    fails = failures(run_checks("pdf-annotate", write((60, 100, 260, 140))))
    assert fails and "the line is at y=712" in fails[0][1]
