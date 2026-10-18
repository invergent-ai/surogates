"""Whether a saved document holds a marker, read back from the file the document server saved."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable, Iterable, Iterator

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape
from pypdf import PdfReader
from pypdf.errors import PyPdfError
from pypdf.generic import StreamObject

# A file these libraries cannot read holds no marker: the save is broken.
UNREADABLE = (zipfile.BadZipFile, KeyError, ValueError, etree.XMLSyntaxError, PyPdfError)


def _docx(data: bytes) -> Iterator[str]:
    # Every paragraph, in tables and content controls too, read whole: the editor splits a paragraph's text into runs.
    for paragraph in Document(io.BytesIO(data)).element.body.iter(qn("w:p")):
        yield "".join(text.text or "" for text in paragraph.iter(qn("w:t")))


def _xlsx(data: bytes) -> Iterator[str]:
    for sheet in load_workbook(io.BytesIO(data)).worksheets:
        for row in sheet.iter_rows(values_only=True):
            yield from (value for value in row if isinstance(value, str))


def _shapes(shapes: Iterable[BaseShape]) -> Iterator[str]:
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _shapes(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame.text
        elif shape.has_table:
            yield from (cell.text for row in shape.table.rows for cell in row.cells)


def _pptx(data: bytes) -> Iterator[str]:
    for slide in Presentation(io.BytesIO(data)).slides:
        yield from _shapes(slide.shapes)


def _pdf(data: bytes) -> Iterator[str]:
    # A note's text is its /Contents; one saved with only rich text has it in /RC, as XHTML.
    for page in PdfReader(io.BytesIO(data)).pages:
        for ref in page.get("/Annots") or []:
            note = ref.get_object()
            yield str(note.get("/Contents") or "")
            rich = note.get("/RC")
            if rich is not None:
                rich = rich.get_object()
                text = rich.get_data().decode("utf-8", "replace") if isinstance(rich, StreamObject) else str(rich)
                yield re.sub(r"<[^>]+>", " ", text)


READERS: dict[str, Callable[[bytes], Iterator[str]]] = {"docx": _docx, "xlsx": _xlsx, "pptx": _pptx, "pdf": _pdf}


def has_marker(kind: str, data: bytes, marker: str) -> bool:
    """Whether *data*, a saved *kind* document, holds *marker* as a word of its own.

    A marker another one starts with (``…-1`` beside ``…-12``) does not count, nor does an unreadable file.
    """
    try:
        texts = READERS[kind]
    except KeyError:
        raise ValueError(f"the tool reads {', '.join(READERS)}, not {kind!r}") from None
    word = re.compile(rf"(?<![A-Za-z0-9-]){re.escape(marker)}(?![A-Za-z0-9-])")
    try:
        return any(word.search(text) for text in texts(data))
    except UNREADABLE:
        return False
