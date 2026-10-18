"""The marker check on a saved file: each edit's op, applied with the format's own library as the editor's save."""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from typing import Any

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import FreeText
from pypdf.generic import NameObject, TextStringObject

from oload.markers import has_marker
from oload.workloads import KIND_SHARES, edit_ops, fixture

MARKER = "load-r1-3-7-0-ab12"


def _saved(build: Callable[[Any], None], document: Any) -> bytes:
    build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def saved_docx(op: dict[str, Any]) -> bytes:
    return _saved(lambda doc: doc.add_paragraph(op["content"]["markdown"]), Document(io.BytesIO(fixture("docx"))))


def saved_xlsx(op: dict[str, Any]) -> bytes:
    sheet, cell = op["range"].split("!")

    def build(book: Any) -> None:
        book[sheet][cell] = op["values"][0][0]

    return _saved(build, load_workbook(io.BytesIO(fixture("xlsx"))))


def saved_pptx(op: dict[str, Any]) -> bytes:
    [text] = re.findall(r"<p>(.*?)</p>", op["slides"][0]["html"])

    def build(deck: Any) -> None:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text_frame.text = text

    return _saved(build, Presentation(io.BytesIO(fixture("pptx"))))


def _pdf_writer() -> PdfWriter:
    return PdfWriter(clone_from=PdfReader(io.BytesIO(fixture("pdf"))))


def _bytes(writer: PdfWriter) -> bytes:
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _free_text(writer: PdfWriter, op: dict[str, Any]) -> None:
    height = float(writer.pages[op["page"] - 1].mediabox.height)
    x1, y1, x2, y2 = op["rect"]  # from the top-left corner; PDF's own space is from the bottom-left
    note = FreeText(text=op["text"], rect=(x1, height - y2, x2, height - y1))
    writer.add_annotation(page_number=op["page"] - 1, annotation=note)


def saved_pdf(op: dict[str, Any]) -> bytes:
    writer = _pdf_writer()
    _free_text(writer, op)
    return _bytes(writer)


SAVED: dict[str, Callable[[dict[str, Any]], bytes]] = {
    "docx": saved_docx, "xlsx": saved_xlsx, "pptx": saved_pptx, "pdf": saved_pdf,
}


@pytest.mark.parametrize("kind", list(KIND_SHARES))
def test_the_saved_file_holds_the_edits_marker(kind: str) -> None:
    [op] = edit_ops(kind, MARKER, 4)
    assert has_marker(kind, SAVED[kind](op), MARKER)


@pytest.mark.parametrize("kind", list(KIND_SHARES))
def test_the_untouched_fixture_holds_no_marker(kind: str) -> None:
    assert not has_marker(kind, fixture(kind), MARKER)


@pytest.mark.parametrize("kind", list(KIND_SHARES))
def test_a_later_marker_that_starts_with_this_one_is_not_this_one(kind: str) -> None:
    [op] = edit_ops(kind, MARKER + "0", 4)
    saved = SAVED[kind](op)
    assert has_marker(kind, saved, MARKER + "0")
    assert not has_marker(kind, saved, MARKER)
    assert not has_marker(kind, saved, MARKER[1:])


@pytest.mark.parametrize("kind", list(KIND_SHARES))
def test_an_unreadable_save_holds_no_marker(kind: str) -> None:
    assert not has_marker(kind, b"", MARKER)
    assert not has_marker(kind, b"not a document " + MARKER.encode(), MARKER)


def test_a_word_paragraph_the_editor_split_into_runs_is_read_whole() -> None:
    doc = Document(io.BytesIO(fixture("docx")))
    paragraph = doc.add_paragraph(MARKER[:7])
    paragraph.add_run(MARKER[7:]).bold = True
    assert has_marker("docx", _saved(lambda _: None, doc), MARKER)


def test_a_word_marker_in_a_table_cell_is_found() -> None:
    doc = Document(io.BytesIO(fixture("docx")))
    doc.tables[0].cell(0, 0).paragraphs[0].add_run(" " + MARKER)
    assert has_marker("docx", _saved(lambda _: None, doc), MARKER)


def test_a_cell_on_any_sheet_is_read() -> None:
    def build(book: Any) -> None:
        book["Notes"]["C9"] = MARKER

    assert has_marker("xlsx", _saved(build, load_workbook(io.BytesIO(fixture("xlsx")))), MARKER)


def test_slide_text_in_a_group_or_a_table_on_any_slide_is_read() -> None:
    def grouped(deck: Any) -> None:
        group = deck.slides[0].shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = MARKER

    def tabled(deck: Any) -> None:
        table = deck.slides[1].shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
        table.cell(1, 1).text = MARKER

    for build in (grouped, tabled):
        assert has_marker("pptx", _saved(build, Presentation(io.BytesIO(fixture("pptx")))), MARKER)


def test_a_pdf_note_whose_text_is_only_rich_text_is_read() -> None:
    writer = _pdf_writer()
    [op] = edit_ops("pdf", MARKER, 0)
    _free_text(writer, op)
    note = writer.pages[0]["/Annots"][-1].get_object()
    del note[NameObject("/Contents")]
    note[NameObject("/RC")] = TextStringObject(f'<?xml version="1.0"?><body><p><span>{MARKER}</span></p></body>')
    assert has_marker("pdf", _bytes(writer), MARKER)


def test_unknown_kinds_are_refused() -> None:
    with pytest.raises(ValueError, match="odt"):
        has_marker("odt", b"", MARKER)
