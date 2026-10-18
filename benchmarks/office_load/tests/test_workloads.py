"""What a virtual session does: its document, its edits and its pace."""

from __future__ import annotations

import io
import random
import statistics
from collections import Counter

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from oload.pace import RECORDED_EDITS, RECORDED_GAPS, Pace
from oload.workloads import FIXTURES, KIND_SHARES, edit_ops, fixture, pick_kind

MARKER = "load-r1-3-7-0-ab12"


def test_docx_edit_inserts_a_paragraph_at_the_end() -> None:
    assert edit_ops("docx", MARKER, 0) == [
        {"op": "insert", "target": {"end": True}, "content": {"markdown": MARKER}},
    ]


def test_xlsx_edit_sets_one_cell_of_the_first_sheet_in_row_n_plus_2_of_a_free_column() -> None:
    assert edit_ops("xlsx", MARKER, 0) == [{"op": "set", "range": "Sales!E2", "values": [[MARKER]]}]
    assert edit_ops("xlsx", MARKER, 5) == [{"op": "set", "range": "Sales!E7", "values": [[MARKER]]}]
    book = load_workbook(io.BytesIO(fixture("xlsx")))
    sheet = book.worksheets[0]
    assert sheet.title == "Sales"
    assert all(sheet.cell(row=row, column=5).value is None for row in range(1, 200))


def test_pptx_edit_adds_one_slide_with_the_marker_as_its_text() -> None:
    [op] = edit_ops("pptx", MARKER, 0)
    assert op["op"] == "slides"
    [slide] = op["slides"]
    assert set(slide) == {"html"}  # no "after": the slide goes last
    html = slide["html"]
    assert html.startswith('<section style="background:') and html.endswith("</section>")
    assert f"<p>{MARKER}</p>" in html


def test_pdf_edit_places_a_note_by_rect_on_page_1_inside_the_page_and_below_its_text() -> None:
    page = PdfReader(io.BytesIO(fixture("pdf"))).pages[0]
    width, height = float(page.mediabox.width), float(page.mediabox.height)
    for n in range(40):
        [op] = edit_ops("pdf", MARKER, n)
        assert op == {"op": "annotate", "kind": "free_text", "page": 1, "rect": op["rect"], "text": MARKER}
        x1, y1, x2, y2 = op["rect"]
        # From the page's top-left corner: the text and fields end about 230 points down.
        assert 0 <= x1 < x2 <= width and 240 <= y1 < y2 <= height


@pytest.mark.parametrize("marker", ["", "1-starts-with-a-digit", "has space", "under_score", "*bold*", "<p>", "x" * 65])
def test_a_marker_is_a_word_every_format_keeps_as_it_is(marker: str) -> None:
    with pytest.raises(ValueError, match="marker"):
        edit_ops("docx", marker, 0)


def test_unknown_kinds_and_negative_edit_numbers_are_refused() -> None:
    with pytest.raises(ValueError, match="odt"):
        edit_ops("odt", MARKER, 0)
    with pytest.raises(ValueError, match="from 0"):
        edit_ops("xlsx", MARKER, -1)


def test_fixtures_are_the_runtime_office_samples() -> None:
    assert FIXTURES.parts[-3:] == ("runtime-office", "test", "fixtures")
    for kind in KIND_SHARES:
        assert fixture(kind) == (FIXTURES / f"sample.{kind}").read_bytes()
    with pytest.raises(ValueError):
        fixture("odt")


def test_pick_kind_follows_the_mix_within_2_percent_over_1000_sessions() -> None:
    counts = Counter(pick_kind(i) for i in range(1000))
    assert KIND_SHARES == {"docx": 0.35, "xlsx": 0.25, "pptx": 0.30, "pdf": 0.10}
    for kind, share in KIND_SHARES.items():
        assert abs(counts[kind] / 1000 - share) <= 0.02, (kind, counts[kind])


def test_pick_kind_is_deterministic_and_mixes_early() -> None:
    assert [pick_kind(i) for i in range(50)] == [pick_kind(i) for i in range(50)]
    assert {pick_kind(i) for i in range(6)} == set(KIND_SHARES)


def test_rollout_never_idles() -> None:
    assert Pace("rollout").idle_s() == 0
    pace = Pace("rollout", random.Random(1))
    assert all(pace.idle_s() == 0 for _ in range(100))


def test_chat_idles_30_to_90_s() -> None:
    pace = Pace("chat", random.Random(2))
    idles = [pace.idle_s() for _ in range(2000)]
    assert all(30 <= idle <= 90 for idle in idles)
    assert min(idles) < 35 and max(idles) > 85


def test_edits_in_turn_draws_only_recorded_values() -> None:
    pace = Pace("rollout", random.Random(3))
    drawn = [pace.edits_in_turn() for _ in range(5000)]
    assert set(drawn) <= set(RECORDED_EDITS)
    assert len(set(drawn)) > 5


def test_gap_draws_recorded_gaps_about_10_s() -> None:
    pace = Pace("chat", random.Random(4))
    drawn = [pace.gap_s() for _ in range(5000)]
    assert set(drawn) <= set(RECORDED_GAPS)
    assert 5 <= statistics.mean(RECORDED_GAPS) <= 15


def test_recorded_edits_have_median_1_p90_6_and_max_28() -> None:
    assert all(isinstance(edits, int) and edits >= 1 for edits in RECORDED_EDITS)
    assert statistics.median(RECORDED_EDITS) == 1
    assert statistics.quantiles(RECORDED_EDITS, n=10)[8] == 6
    assert max(RECORDED_EDITS) == 28


def test_a_pace_is_for_rollout_or_chat() -> None:
    with pytest.raises(ValueError, match="rollout"):
        Pace("batch")  # type: ignore[arg-type]
