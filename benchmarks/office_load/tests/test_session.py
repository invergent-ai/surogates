"""A virtual session: open, edit turn by turn, save each turn and check the save holds the turn's last marker.

The load client is faked in process; the document server is a fake over HTTP, so the forcesave, its signed callback
and the saved file's download go through the tool's own routes.
"""

from __future__ import annotations

import asyncio
import io
import json
import random
import time
from collections.abc import AsyncIterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import aiohttp
import jwt
import pytest
from aiohttp import web
from aiohttp.test_utils import TestServer, unused_port
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from oload import docserver, session
from oload.clients import (
    CallTimedOut,
    ClientBusy,
    ClientDown,
    Fleet,
    FleetFull,
    PageGone,
    RequestRefused,
    RuntimeOfficeError,
    Timed,
)
from oload.docserver import CallbackState, make_app
from oload.events import Event, EventLog
from oload.session import VirtualSession
from oload.settings import Settings
from oload.workloads import OUTLINE_READ, pick_kind

SECRET = "a-test-secret-of-thirty-two-bytes-or-more"
DS_HOST = "http://office-docs.example.svc"  # how the document server names its saved files: rewritten by the tool
READS: dict[str, Any] = json.loads((Path(__file__).parent / "fixtures" / "sample_pptx_reads.json").read_text())


def index_of(kind: str) -> int:
    return next(i for i in range(20) if pick_kind(i) == kind)


def saved_file(kind: str, markers: list[str]) -> bytes:
    """The document as the document server saves it: every marker its editors sent, a deck's text box the last."""
    out = io.BytesIO()
    if kind == "docx":
        document = Document()
        for marker in markers:
            document.add_paragraph(marker)
        document.save(out)
    else:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = (markers or [""])[-1]
        deck.save(out)
    return out.getvalue()


class FakeDocumentServer:
    """Keeps what each document's editors sent; answers forcesave and calls back with the saved file, as the document
    server does. Each knob names saves by their order, from 0."""

    def __init__(self) -> None:
        self.sent: dict[str, list[str]] = {}  # doc_key -> the markers its editors sent, in order
        self.saved: dict[str, int] = {}  # doc_key -> how many of them its last save held
        self.callback_urls: dict[str, str] = {}
        self.kinds: dict[str, str] = {}
        self.files: dict[str, bytes] = {}
        self.commands: list[dict[str, Any]] = []
        self.errors: dict[int, int] = {}  # save -> the error answered; 4 also drops what was unsaved
        self.http_errors: dict[int, int] = {}  # save -> the HTTP status /command answers
        self.slow: set[int] = set()  # saves whose /command never answers in time
        self.stale: set[int] = set()  # saves whose file is the save before it
        self.partial: set[int] = set()  # saves whose file lacks the last edit sent
        self.held: dict[int, int] = {}  # save -> the save whose /command releases its callback
        self.delays: dict[int, float] = {}  # save -> seconds before its callback
        self.download_delays: dict[int, float] = {}  # save -> seconds before its file is served
        self.gates: dict[int, asyncio.Event] = {}
        self.tasks: set[asyncio.Task[None]] = set()

    def app(self) -> web.Application:
        app = web.Application()
        app.router.add_post("/command", self._command)
        app.router.add_get("/cache/{name}", self._cache)
        return app

    async def opened(self, config: dict[str, Any]) -> None:
        """An editor opened a document: the input fetched from the tool, as the document server fetches it."""
        key = config["document"]["key"]
        async with aiohttp.ClientSession() as http, http.get(config["document"]["url"]) as response:
            assert response.status == 200, "the tool does not serve the document"
        self.callback_urls[key] = config["editorConfig"]["callbackUrl"]
        self.kinds[key] = config["document"]["fileType"]
        self.sent.setdefault(key, [])

    async def _command(self, request: web.Request) -> web.Response:
        body = await request.json()
        assert jwt.decode(body["token"], SECRET, algorithms=["HS256"])["key"] == body["key"]
        save = len(self.commands)
        self.commands.append(body)
        for held, by in self.held.items():
            if by == save:
                self.gates.setdefault(held, asyncio.Event()).set()
        if save in self.slow:
            await asyncio.sleep(2)
        if save in self.http_errors:
            return web.json_response({"detail": "no"}, status=self.http_errors[save])
        key, sent = body["key"], self.sent[body["key"]]
        before = self.saved.get(key, 0)
        error = self.errors.get(save, 0 if len(sent) > before else 4)
        self.saved[key] = len(sent)
        if error == 0:
            markers = sent[:before] if save in self.stale else sent[:-1] if save in self.partial else list(sent)
            self.files[f"{save}"] = saved_file(self.kinds[key], markers)
            task = asyncio.create_task(self._call_back(save, key, body["userdata"]))
            self.tasks.add(task)
            task.add_done_callback(self.tasks.discard)
        return web.json_response({"error": error})

    async def _call_back(self, save: int, key: str, userdata: str) -> None:
        if save in self.held:
            await self.gates.setdefault(save, asyncio.Event()).wait()
        await asyncio.sleep(self.delays.get(save, 0.0))
        payload = {"key": key, "status": 6, "userdata": userdata, "url": f"{DS_HOST}/cache/{save}"}
        async with aiohttp.ClientSession() as http:
            await http.post(self.callback_urls[key], json={**payload, "token": jwt.encode(payload, SECRET)})

    async def _cache(self, request: web.Request) -> web.Response:
        name = request.match_info["name"]
        await asyncio.sleep(self.download_delays.get(int(name), 0.0))
        return web.Response(body=self.files[name])

    async def stop(self) -> None:
        for task in self.tasks:
            task.cancel()
        await asyncio.gather(*self.tasks, return_exceptions=True)


@dataclass(frozen=True)
class AfterApplying:
    """An edit that applies, then fails as *error*: an answer lost after the editor ran it."""

    error: Exception


class FakeOffice:
    """A load client's pages, in process. Each edit sends its marker to the document server at once.

    *script* gives a call's next answers in turn: an error raised, an answer given, AfterApplying, or None for the
    usual one. *on_edit* is called at each edit.
    """

    def __init__(self, ds: FakeDocumentServer) -> None:
        self.ds = ds
        self.calls: list[str] = []
        self.opened: list[tuple[str, str, dict[str, Any]]] = []  # doc_key, session_id, config
        self.closed: list[str] = []
        self.placed_at_close: list[tuple[int, ...]] = []
        self.fleet: Fleet | None = None
        self.script: dict[str, list[Any]] = {}
        self.pages: dict[str, str] = {}  # page -> doc_key
        self.hash = 0  # the deck slide's, changed by each edit
        self.on_edit: Any = None

    def _scripted(self, call: str) -> Any:
        queue = self.script.get(call)
        item = queue.pop(0) if queue else None
        if isinstance(item, BaseException):
            raise item
        return item

    async def open(self, doc_key: str, session_id: str, config: dict[str, Any],
                   document_server_url: str) -> Timed[str]:
        self.calls.append(f"open {session_id}")
        assert document_server_url == DS_HOST
        self.opened.append((doc_key, session_id, config))
        self._scripted("open")
        await self.ds.opened(config)
        self.pages[f"{doc_key}.{session_id}"] = doc_key
        return Timed(f"{doc_key}.{session_id}", 0.01)

    async def read(self, page: str, request: dict[str, Any]) -> Timed[dict[str, Any]]:
        self.calls.append("read outline" if request == OUTLINE_READ else "read slide")
        self._scripted("read")
        if request == OUTLINE_READ:
            return Timed(READS["outline"], 0.01)
        assert request == {"view": "slide", "target": "1473"}
        return Timed({**READS["slide"], "hash": f"{self.hash:08x}"}, 0.01)

    async def edit(self, page: str, ops: list[Any]) -> Timed[dict[str, Any]]:
        self.calls.append("edit")
        if self.on_edit is not None:
            self.on_edit()
        answer = self._scripted("edit")
        if answer is not None and not isinstance(answer, AfterApplying):
            return Timed(answer, 0.01)
        [op] = ops
        if op["op"] == "text":
            if op["hash"] != f"{self.hash:08x}":
                return Timed({"applied": False, "refused": "the slide changed since it was read"}, 0.01)
            self.hash += 1
        self.ds.sent[self.pages[page]].append(op["text"] if op["op"] == "text" else op["content"]["markdown"])
        if isinstance(answer, AfterApplying):
            raise answer.error
        return Timed({"applied": True, "results": [{}], "flushed": True}, 0.01)

    async def flush(self, page: str) -> Timed[bool]:
        self.calls.append("flush")
        answer = self._scripted("flush")
        return Timed(True if answer is None else answer, 0.01)

    async def close(self, page: str) -> None:
        self.calls.append("close")
        assert self.fleet is not None
        self.placed_at_close.append(self.fleet.placed)
        self.closed.append(page)
        self.pages.pop(page, None)
        self._scripted("close")


class ScriptedPace:
    """Each turn's edit count from *edits*; gaps and idle as given."""

    def __init__(self, edits: list[int], *, gap: float = 0.0, idle: float = 0.0) -> None:
        self.edits = list(edits)
        self.gap = gap
        self.idle = idle

    def edits_in_turn(self) -> int:
        return self.edits.pop(0) if self.edits else 1

    def gap_s(self) -> float:
        return self.gap

    def idle_s(self) -> float:
        return self.idle


@dataclass
class World:
    settings: Settings
    state: CallbackState
    ds: FakeDocumentServer
    office: FakeOffice
    fleet: Fleet
    log: EventLog

    def session(self, kind: str = "docx", workload: str = "rollout", edits: list[int] | None = None,
                **pace: float) -> VirtualSession:
        return VirtualSession(index_of(kind), workload, self.fleet, self.settings, self.state, self.log,
                              random.Random(0), run_id="r1", pace=ScriptedPace(edits or [1], **pace))

    def seen(self) -> list[tuple[str, bool, str | None]]:
        return [(event.what, event.ok, event.cause) for event in self.log.events]


@pytest.fixture
async def world() -> AsyncIterator[World]:
    tool_port, ds_port = unused_port(), unused_port()
    settings = Settings(
        ds_url=DS_HOST,
        command_url=f"http://127.0.0.1:{ds_port}",
        jwt_secret=SECRET,
        client_token="a-test-client-token",
        listen=f"127.0.0.1:{tool_port}",
        public_base=f"http://127.0.0.1:{tool_port}",
        namespace="office-load",
    )
    state = CallbackState(settings)
    ds = FakeDocumentServer()
    office = FakeOffice(ds)
    fleet = Fleet([office], pages_per_client=2)  # type: ignore[list-item]
    office.fleet = fleet
    async with TestServer(make_app(state), port=tool_port), TestServer(ds.app(), port=ds_port):
        yield World(settings, state, ds, office, fleet, EventLog())
        await ds.stop()


async def run(session: VirtualSession, turns: int) -> None:
    await session.run(time.monotonic() + 60, turns=turns)


OK_TURN = [("edit", True, None), ("flush", True, None), ("forcesave", True, None), ("callback", True, None),
           ("marker", True, None)]


# A turn.

async def test_a_turn_logs_open_edits_flush_forcesave_callback_marker(world: World) -> None:
    world.ds.delays[0] = 0.2
    world.ds.download_delays[0] = 0.3
    one = world.session(edits=[2])
    await run(one, turns=1)

    assert world.seen() == [
        ("open", True, None), ("edit", True, None), ("edit", True, None), ("flush", True, None),
        ("forcesave", True, None), ("callback", True, None), ("marker", True, None),
    ]
    assert {event.session for event in world.log.events} == {one.index}
    assert {event.kind for event in world.log.events} == {"docx"}
    callback = next(event for event in world.log.events if event.what == "callback")
    assert 0.5 <= callback.seconds < 5  # from the forcesave's sending to the save's file in hand, downloaded
    assert all(event.at > 0 and event.seconds >= 0 for event in world.log.events)
    # The forcesave names the document, and the turn its callback is matched by.
    [command] = world.ds.commands
    assert command["c"] == "forcesave" and command["key"] == one.doc_key and command["userdata"]
    assert one.doc_key.startswith(f"load-r1-{one.index}-")
    # Two distinct markers, each one word.
    assert len(set(world.ds.sent[one.doc_key])) == 2
    # Its page closed, its placement given back and the saved file's bytes dropped once checked.
    assert world.office.closed == [f"{one.doc_key}.agent-{one.index}"]
    assert (world.fleet.placed, world.fleet.leaked) == ((0,), (0,))
    assert all(callback.data is None for callback in world.state.callbacks)


async def test_turns_go_on_until_their_count(world: World) -> None:
    holding: list[list[bool]] = []  # at each edit: which saves' files are still held
    world.office.on_edit = lambda: holding.append([c.data is not None for c in world.state.callbacks])
    await run(world.session(edits=[1, 3]), turns=2)

    assert world.seen() == [("open", True, None), *OK_TURN, ("edit", True, None), ("edit", True, None), *OK_TURN]
    assert holding[1:] == [[False]] * 3  # the first turn's save dropped once checked


async def test_the_session_ends_at_until_even_in_its_idle(world: World) -> None:
    one = world.session(workload="chat", edits=[1, 1, 1], idle=100.0)
    started = time.monotonic()
    await one.run(started + 1.0)
    assert time.monotonic() - started < 3
    assert world.seen()[-5:] == OK_TURN
    assert [event.what for event in world.log.events].count("forcesave") == 1
    assert world.fleet.placed == (0,)


# Saves that are not the turn's.

async def test_stale_callback_is_a_lost_save(world: World) -> None:
    world.ds.stale.add(1)  # the second save's file is the first's: it lacks the second turn's marker
    await run(world.session(edits=[1, 1]), turns=2)

    assert world.seen() == [
        ("open", True, None), *OK_TURN,
        ("edit", True, None), ("flush", True, None), ("forcesave", True, None), ("callback", True, None),
        ("marker", False, "lost"),
    ]


async def test_a_save_without_the_turns_last_edit_is_lost(world: World) -> None:
    world.ds.partial.add(0)  # the save holds the turn's first edit, not its last
    await run(world.session(edits=[2]), turns=1)
    assert world.seen()[-2:] == [("callback", True, None), ("marker", False, "lost")]


async def test_late_earlier_callback_does_not_satisfy_a_later_turn(
    world: World, monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(session, "CALLBACK_WAIT_S", 1.0)
    world.ds.held[0] = 1  # the first save's callback comes once the second turn's forcesave is in
    world.ds.delays[1] = 0.2  # and before the second save's own
    await run(world.session(edits=[1, 1]), turns=2)

    assert world.seen() == [
        ("open", True, None),
        ("edit", True, None), ("flush", True, None), ("forcesave", True, None), ("callback", False, "timeout"),
        *OK_TURN,
    ]
    # Both came; the first turn's, late, was taken for neither turn, and its bytes are dropped too.
    assert [callback.userdata for callback in world.state.callbacks] == [
        command["userdata"] for command in world.ds.commands
    ]
    assert all(callback.data is None for callback in world.state.callbacks)


async def test_no_callback_within_the_wait_is_a_timeout(world: World, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(session, "CALLBACK_WAIT_S", 0.3)
    world.ds.held[0] = 99  # never comes
    started = time.monotonic()
    await run(world.session(), turns=1)

    assert world.seen()[-2:] == [("forcesave", True, None), ("callback", False, "timeout")]
    timed_out = world.log.events[-1]
    assert 0.3 <= timed_out.seconds < time.monotonic() - started


# Forcesave's answers.

async def test_forcesave_4_after_edits_is_lost_but_after_no_edit_is_fine(world: World) -> None:
    world.ds.errors[0] = 4  # the first turn's edits never reached the document server
    world.office.script["edit"] = [None, {"applied": False, "refused": "the slide changed"}]
    await run(world.session(edits=[1, 0, 1]), turns=3)

    assert world.seen() == [
        ("open", True, None),
        ("edit", True, None), ("flush", True, None), ("forcesave", False, "lost"),
        ("flush", True, None), ("forcesave", True, None),
        ("edit", False, "server"), ("flush", True, None), ("forcesave", True, None),
    ]


@pytest.mark.parametrize(("knob", "cause"), [("http_errors", "server"), ("slow", "timeout")])
async def test_a_forcesave_not_answered_waits_for_no_callback(
    world: World, monkeypatch: pytest.MonkeyPatch, knob: str, cause: str,
) -> None:
    monkeypatch.setattr(docserver, "FORCESAVE_TIMEOUT_S", 0.3)
    if knob == "slow":
        world.ds.slow.add(0)
    else:
        world.ds.http_errors[0] = 500
    await run(world.session(edits=[1, 1]), turns=2)

    assert world.seen() == [
        ("open", True, None), ("edit", True, None), ("flush", True, None), ("forcesave", False, cause), *OK_TURN,
    ]


async def test_a_flush_not_acknowledged_sends_no_forcesave(world: World) -> None:
    world.office.script["flush"] = [False]
    await run(world.session(), turns=1)
    assert world.seen() == [("open", True, None), ("edit", True, None), ("flush", False, "server")]
    assert world.ds.commands == []


# Failures along the way.

async def test_a_refused_request_is_the_tools_error_and_the_session_goes_on(world: World) -> None:
    world.office.script["edit"] = [RequestRefused('op 0: unknown op "set"', status=400)]
    await run(world.session(edits=[2, 1]), turns=2)

    assert world.seen() == [("open", True, None), ("error", False, "tool"), *OK_TURN, *OK_TURN]
    assert 'unknown op "set"' in (world.log.events[1].detail or "")


async def test_page_gone_reopens_once_and_a_second_loss_ends_the_session(world: World) -> None:
    gone = PageGone("no page is open", status=410, seconds=0.01)
    world.office.script["edit"] = [gone, None, PageGone("no page is open", status=410, seconds=0.01)]
    one = world.session(edits=[3])
    await run(one, turns=1)

    assert world.seen() == [
        ("open", True, None), ("edit", False, "server"), ("open", True, None), ("edit", True, None),
        ("edit", False, "server"),
    ]
    assert [call for call in world.office.calls if call.startswith("open")] == [f"open agent-{one.index}"] * 2
    assert world.ds.commands == []
    assert world.fleet.placed == (0,)


async def test_a_page_gone_at_the_flush_still_forcesaves_and_opens_again_for_the_next_turn(world: World) -> None:
    world.office.script["flush"] = [PageGone("the editor died", status=410, seconds=0.01)]
    world.ds.errors[0] = 4  # the document server saved it on its own as its last editor left
    await run(world.session(edits=[1, 1]), turns=2)

    assert world.seen() == [
        ("open", True, None), ("edit", True, None), ("flush", False, "server"), ("forcesave", True, None),
        ("open", True, None), *OK_TURN,
    ]


@pytest.mark.parametrize("error", [CallTimedOut("no answer within 90 s", seconds=90.0),
                                   ClientDown("the client dropped the call", seconds=90.0)])
async def test_an_open_that_timed_out_or_dropped_leaks_its_slot(world: World, error: Exception) -> None:
    world.office.script["open"] = [error]
    one = world.session()
    await run(one, turns=1)

    assert world.seen() == [("open", False, error.cause)]  # type: ignore[attr-defined]
    assert world.log.events[0].seconds == 90.0
    # The page may be opening still, and a close does not stop an open: closed by its id all the same, its slot
    # never handed out again.
    assert world.office.closed == [f"{one.doc_key}.agent-{one.index}"]
    assert (world.fleet.placed, world.fleet.leaked) == ((0,), (1,))
    world.fleet.place(1)
    with pytest.raises(FleetFull):
        world.fleet.place(1)


@pytest.mark.parametrize("error", [ClientBusy("6 documents open", status=503, seconds=0.01),
                                   RuntimeOfficeError("internal error", status=500, seconds=30.0)])
async def test_an_open_answered_with_a_failure_gives_its_slot_back(world: World, error: Exception) -> None:
    world.office.script["open"] = [error]
    await run(world.session(), turns=1)
    assert world.seen() == [("open", False, error.cause)]  # type: ignore[attr-defined]
    assert (world.fleet.placed, world.fleet.leaked) == ((0,), (0,))


async def test_a_page_whose_open_and_close_both_failed_leaks_one_slot(world: World) -> None:
    world.office.script["open"] = [CallTimedOut("no answer within 90 s", seconds=90.0)]
    world.office.script["close"] = [ClientDown("the client could not be reached", seconds=0.01)]
    await run(world.session(), turns=1)
    assert (world.fleet.placed, world.fleet.leaked) == ((0,), (1,))


async def test_a_page_that_did_not_close_leaks_its_slot(world: World) -> None:
    world.office.script["close"] = [None, CallTimedOut("no answer within 60 s", seconds=60.0)]
    await run(world.session(workload="chat"), turns=1)
    # The agent's page closed and its slot is back; the watcher's may live on.
    assert (world.fleet.placed, world.fleet.leaked) == ((0,), (1,))


# Decks.

async def test_a_deck_reads_its_slide_before_each_edit_and_its_outline_again_after_a_reopen(world: World) -> None:
    world.office.script["edit"] = [None, PageGone("the editor died", status=410, seconds=0.01)]
    one = world.session(kind="pptx", edits=[2, 1])
    await run(one, turns=2)

    opening = f"open agent-{one.index}"
    assert world.office.calls == [
        opening, "read outline", "read slide", "edit", "read slide", "edit",  # the second edit finds its page gone
        opening, "flush",
        "read outline", "read slide", "edit", "flush",
        "close",
    ]
    # Every edit carried the hash of the read just before it: none was refused.
    assert world.seen() == [
        ("open", True, None), ("edit", True, None), ("edit", False, "server"),
        ("open", True, None), ("flush", True, None), ("forcesave", True, None), ("callback", True, None),
        *OK_TURN,
    ]


async def test_a_deck_edit_that_failed_after_applying_may_be_the_one_saved(world: World) -> None:
    # Its text replaced the one before it: the save holds the edit whose answer was lost, not the last applied.
    world.office.script["edit"] = [None, AfterApplying(CallTimedOut("no answer within 300 s", seconds=300.0))]
    await run(world.session(kind="pptx", edits=[2]), turns=1)
    assert world.seen()[-6:] == [("edit", True, None), ("edit", False, "timeout"), *OK_TURN[1:]]


# Chat.

async def test_chat_opens_a_watching_page(world: World) -> None:
    one = world.session(workload="chat")
    await run(one, turns=1)

    assert world.seen() == [("open", True, None), ("open", True, None), *OK_TURN]
    [(agent_key, agent_session, agent), (watcher_key, watcher_session, watcher)] = world.office.opened
    assert agent_key == watcher_key == one.doc_key
    assert agent_session != watcher_session
    assert agent["editorConfig"]["user"]["id"] != watcher["editorConfig"]["user"]["id"]
    assert agent["editorConfig"]["user"]["name"] != watcher["editorConfig"]["user"]["name"]
    assert agent["editorConfig"]["customization"]["plugins"] is False
    assert "plugins" not in watcher["editorConfig"]["customization"]  # a person's editor has them
    # Only the agent's page edits; both close, the agent's first, and both slots are given back.
    assert world.office.calls.count("edit") == 1
    assert world.office.closed == [f"{one.doc_key}.{agent_session}", f"{one.doc_key}.{watcher_session}"]
    assert world.office.placed_at_close == [(2,), (2,)]
    assert world.fleet.placed == (0,)


# The log.

def test_event_log_writes_and_reads_events_jsonl(tmp_path: Path) -> None:
    path = tmp_path / "run" / "events.jsonl"
    log = EventLog(path)
    first = Event(3, "docx", "open", True, 1.5, None, 1700000000.0)
    second = Event(3, "docx", "marker", False, 0.02, "lost", 1700000009.5, detail="the save lacks m")
    log.add(first)
    log.add(second)

    assert log.events == [first, second]
    lines = path.read_text().splitlines()
    assert json.loads(lines[1]) == {"session": 3, "kind": "docx", "what": "marker", "ok": False, "seconds": 0.02,
                                    "cause": "lost", "at": 1700000009.5, "detail": "the save lacks m"}
    assert EventLog.read(path) == [first, second]
