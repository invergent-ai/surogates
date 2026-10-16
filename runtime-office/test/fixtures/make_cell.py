"""Builds the workbooks the spreadsheet tests open: run it with a Python that has openpyxl."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Protection
from openpyxl.workbook.protection import WorkbookProtection

here = Path(__file__).parent

book = Workbook()
sales = book.active
sales.title = "Sales"
for row in [["Region", "Q1", "Q2", "Total"], ["North", 120, 140, "=B2+C2"], ["East", 90, 110, "=B3+C3"], ["South", 70, 65, "=B4+C4"]]:
    sales.append(row)
for cell in sales[1]:
    cell.font = Font(bold=True)
sales["B2"].number_format = "#,##0"
book.create_sheet("Notes")["A1"] = "Figures in thousands of euros."
book.save(here / "sample.xlsx")

# A sheet protected in Excel, with one cell left unlocked for the user to fill in.
book = Workbook()
locked = book.active
locked.title = "Locked"
for row in [["Item", "Amount"], ["Rent", 1000], ["Heating", ""]]:
    locked.append(row)
locked["B3"].protection = Protection(locked=False)
locked.protection.sheet = True
book.create_sheet("Open")["A1"] = "free"
book.save(here / "protected.xlsx")

# A workbook whose structure is protected: no sheet may be added, renamed or deleted.
book = Workbook()
book.active.title = "Only"
book.security = WorkbookProtection(lockStructure=True)
book.save(here / "structure.xlsx")
