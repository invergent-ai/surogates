"""Builds sample.docx, the document the integration tests open, and protected.docx, the same one protected:
run it with a Python that has python-docx."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

document = Document()
document.add_heading("Quarterly report", level=1)
paragraph = document.add_paragraph("This report covers the quarter. The ")
paragraph.add_run("revenue target").bold = True
paragraph.add_run(" was met in all regions.")
document.add_heading("Regions", level=2)
document.add_paragraph("North grew fastest, then the east.")
document.add_paragraph("North", style="List Bullet")
document.add_paragraph("East", style="List Bullet")
document.add_paragraph("Hire two engineers", style="List Number")
document.add_paragraph("Open the Lisbon office", style="List Number")
table = document.add_table(rows=2, cols=2)
table.style = "Table Grid"
for (row, column), text in {(0, 0): "Region", (0, 1): "Growth", (1, 0): "North", (1, 1): "12%"}.items():
    table.cell(row, column).text = text
document.add_heading("Risks", level=2)
document.add_paragraph("Hiring could slip a quarter.")
document.save(Path(__file__).with_name("sample.docx"))

# As Word's Restrict Editing leaves a document that may only be read.
protection = OxmlElement("w:documentProtection")
protection.set(qn("w:edit"), "readOnly")
protection.set(qn("w:enforcement"), "1")
document.settings.element.find(qn("w:defaultTabStop")).addprevious(protection)
document.save(Path(__file__).with_name("protected.docx"))
