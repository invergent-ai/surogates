"""Builds sample.pptx, the deck the slide tests open: run it with a Python that has python-pptx."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

deck = Presentation()
deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
first = deck.slides.add_slide(deck.slide_layouts[1])
first.shapes.title.text = "Quarterly review"
first.placeholders[1].text = "Revenue met target\nHiring slipped"
second = deck.slides.add_slide(deck.slide_layouts[5])
second.shapes.title.text = "Next quarter"
box = second.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
box.text_frame.text = "Open the Lisbon office"
box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
second.notes_slide.notes_text_frame.text = "Mention the lease."
deck.save(Path(__file__).with_name("sample.pptx"))
